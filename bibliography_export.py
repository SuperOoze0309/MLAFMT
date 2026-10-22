"""Export a separately reviewed CSL bibliography as an in-memory Word file.

This is a renderer, not a citation formatter: the CSL processor supplies the
entry order and text. No paper, citation numbering or external URL is read.
"""

from __future__ import annotations

from dataclasses import dataclass
from html.parser import HTMLParser
from io import BytesIO
import math
import re
from urllib.parse import urlsplit

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt

MAX_ENTRIES = 200
MAX_HTML_BYTES = 2 * 1024 * 1024
MAX_ENTRY_BYTES = 32 * 1024
MAX_TEXT_BYTES = 2 * 1024 * 1024
HEADINGS = {"mla": "Works Cited", "apa": "References"}
_ID = re.compile(r"[A-Za-z0-9_-]{1,64}\Z")
_INVALID_TEXT = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f\ud800-\udfff\ufffe\uffff]")
_STYLE_KEYS = {"italic", "bold", "small_caps", "superscript", "subscript", "underline"}


class BibliographyExportError(ValueError):
    """A bounded, unsupported or malformed export request."""

    def __init__(self):
        super().__init__("invalid_options")


def _reject():
    raise BibliographyExportError()


def _text(value, byte_limit, *, nonempty=False):
    if not isinstance(value, str) or _INVALID_TEXT.search(value):
        _reject()
    if len(value.encode("utf-8")) > byte_limit or (nonempty and not value.strip()):
        _reject()
    return value


def _keys(value, required, optional=()):
    if not isinstance(value, dict) or set(value) - set(required) - set(optional):
        _reject()
    if not set(required).issubset(value):
        _reject()


def _number(value, low, high):
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        _reject()
    if not math.isfinite(value) or not low <= value <= high:
        _reject()
    return value


def _safe_url(value):
    _text(value, 8192, nonempty=True)
    if any(character.isspace() for character in value) or "\\" in value:
        _reject()
    try:
        parts = urlsplit(value)
        if (parts.scheme.lower() not in ("http", "https") or not parts.hostname
                or parts.username is not None or parts.password is not None):
            _reject()
        # Accessing port catches malformed authority such as :bad or :99999.
        _ = parts.port
    except (ValueError, UnicodeError):
        _reject()
    return value


@dataclass(frozen=True)
class Segment:
    text: str
    style: dict
    href: str | None = None


class _EntryHTML(HTMLParser):
    """Only the formatting vocabulary exported by our two pinned CSL styles."""

    _INLINE = {"i", "em", "b", "strong", "u", "sup", "sub", "small", "span", "a"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.stack = []
        self.segments = []
        self.root_seen = False
        self.root_closed = False

    def handle_starttag(self, tag, attributes):
        if len(self.stack) >= 32 or tag not in self._INLINE | {"div", "br"}:
            _reject()
        if len({key for key, _ in attributes}) != len(attributes):
            _reject()
        attributes = dict(attributes)
        if any(value is None for value in attributes.values()):
            _reject()
        if set(attributes) - {"class", "style", "href", "title"}:
            _reject()
        if "href" in attributes and tag != "a":
            _reject()
        if "title" in attributes:
            if tag != "a":
                _reject()
            _text(attributes["title"], MAX_ENTRY_BYTES)
        classes = attributes.get("class", "").split()
        allowed_classes = {"csl-entry"} if tag == "div" else {"nocase", "nodecor"} if tag == "span" else set()
        if set(classes) - allowed_classes:
            _reject()
        if tag == "div":
            if self.stack or self.root_seen or self.segments:
                _reject()
            self.root_seen = True
        elif self.root_closed:
            _reject()
        current = dict(self.stack[-1][1]) if self.stack else {}
        href = self.stack[-1][2] if self.stack else None
        if tag in {"i", "em"}:
            current["italic"] = True
        elif tag in {"b", "strong"}:
            current["bold"] = True
        elif tag == "u":
            current["underline"] = True
        elif tag == "small":
            # HTML small means reduced font size, not small caps. CSL does not
            # emit it for these styles; do not silently conflate the two.
            _reject()
        elif tag in {"sup", "sub"}:
            current["superscript"] = tag == "sup"
            current["subscript"] = tag == "sub"
        elif tag == "a":
            if href is not None or "href" not in attributes:
                _reject()
            href = _safe_url(attributes["href"])
        if "nodecor" in classes:
            current.update({key: False for key in _STYLE_KEYS})
        if "style" in attributes:
            self._css(current, attributes["style"])
        if tag == "br":
            if attributes:
                _reject()
            self.segments.append(Segment("\n", current, href))
        else:
            self.stack.append((tag, current, href))

    @staticmethod
    def _css(current, css):
        _text(css, 1000)
        seen = set()
        for declaration in css.split(";"):
            if not declaration.strip():
                continue
            pieces = declaration.split(":")
            if len(pieces) != 2:
                _reject()
            key, value = (piece.strip().lower() for piece in pieces)
            if key in seen:
                _reject()
            seen.add(key)
            if key == "font-style" and value in ("normal", "italic"):
                current["italic"] = value == "italic"
            elif key == "font-weight" and value in ("normal", "bold", "400", "700"):
                current["bold"] = value in ("bold", "700")
            elif key == "font-variant" and value in ("normal", "small-caps"):
                current["small_caps"] = value == "small-caps"
            elif key == "text-decoration" and value in ("none", "underline"):
                current["underline"] = value == "underline"
            elif key == "vertical-align" and value in ("baseline", "sup", "sub", "super"):
                current["superscript"] = value in ("sup", "super")
                current["subscript"] = value == "sub"
            else:
                _reject()

    def handle_endtag(self, tag):
        if not self.stack or self.stack[-1][0] != tag:
            _reject()
        self.stack.pop()
        if tag == "div":
            self.root_closed = True

    def handle_startendtag(self, tag, attributes):
        if tag != "br":
            _reject()
        self.handle_starttag(tag, attributes)

    def handle_data(self, data):
        _text(data, MAX_ENTRY_BYTES)
        if not self.stack and not data.strip():
            return
        if self.root_closed and data:
            _reject()
        style = dict(self.stack[-1][1]) if self.stack else {}
        href = self.stack[-1][2] if self.stack else None
        self.segments.append(Segment(data, style, href))

    def handle_comment(self, data):
        _reject()

    def handle_decl(self, declaration):
        _reject()

    def handle_pi(self, data):
        _reject()

    def unknown_decl(self, data):
        _reject()

    def parsed(self, html):
        self.feed(html)
        self.close()
        if self.stack or not any(segment.text.strip() for segment in self.segments):
            _reject()
        return self.segments


def validate_bibliography(data):
    """Validate all payloads before constructing any Word package."""
    _keys(data, {"style", "heading", "entries", "layout"}, {"html", "text"})
    style = data["style"]
    if not isinstance(style, str) or style not in HEADINGS or data["heading"] != HEADINGS[style]:
        _reject()
    _text(data["heading"], 80, nonempty=True)
    for optional in ("html", "text"):
        if optional in data:
            _text(data[optional], MAX_HTML_BYTES if optional == "html" else MAX_TEXT_BYTES)
    entries = data["entries"]
    if not isinstance(entries, list) or not 1 <= len(entries) <= MAX_ENTRIES:
        _reject()
    layout = data["layout"]
    _keys(layout, {"hangingindent", "linespacing", "entryspacing", "second-field-align", "maxoffset"})
    if not isinstance(layout["hangingindent"], bool):
        _reject()
    _number(layout["linespacing"], 0, 10)
    _number(layout["entryspacing"], 0, 20)
    _number(layout["maxoffset"], 0, 1000)
    if layout["second-field-align"] is not False:
        _reject()
    seen = set()
    parsed = []
    html_total = 0
    text_total = 0
    for entry in entries:
        _keys(entry, {"id", "html", "text"})
        identifier = entry["id"]
        if not isinstance(identifier, str) or not _ID.fullmatch(identifier) or identifier in seen:
            _reject()
        seen.add(identifier)
        html = _text(entry["html"], MAX_ENTRY_BYTES, nonempty=True)
        text = _text(entry["text"], MAX_ENTRY_BYTES, nonempty=True)
        html_total += len(html.encode("utf-8"))
        text_total += len(text.encode("utf-8"))
        if html_total > MAX_HTML_BYTES or text_total > MAX_TEXT_BYTES:
            _reject()
        parsed.append(_EntryHTML().parsed(html))
    return parsed


def _format_run(run, style):
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    # Apply East Asian font explicitly without touching the citation text.
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Times New Roman")
    for key, value in style.items():
        setattr(run.font, key, value)


def _add_segment(paragraph, segment):
    if segment.href is None:
        _format_run(paragraph.add_run(segment.text), segment.style)
        return
    from docx.text.run import Run

    relation = paragraph.part.relate_to(segment.href, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation)
    run_element = OxmlElement("w:r")
    hyperlink.append(run_element)
    paragraph._p.append(hyperlink)
    run = Run(run_element, paragraph)
    run.text = segment.text
    _format_run(run, segment.style)


def render_bibliography_docx(data):
    """Return DOCX bytes; never read or modify a draft or fetch any link."""
    parsed = validate_bibliography(data)
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Times New Roman", Pt(12)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 2
    document.core_properties.title = data["heading"]
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.subject = "Independent bibliography generated locally by FormCapy"
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.keep_with_next = True
    title.paragraph_format.space_after = Pt(0)
    _format_run(title.add_run(data["heading"]), {"bold": data["style"] == "apa"})
    layout = data["layout"]
    line_spacing = layout["linespacing"] or 2
    for segments in parsed:
        paragraph = document.add_paragraph()
        formatting = paragraph.paragraph_format
        formatting.line_spacing = line_spacing
        formatting.space_before = Pt(0)
        # CSL spacing values are blank-line multiples. Use 12pt text with the
        # CSL line-height multiplier, preserving the entry order verbatim.
        formatting.space_after = Pt(12 * line_spacing * layout["entryspacing"])
        if layout["hangingindent"]:
            formatting.left_indent = Inches(0.5)
            formatting.first_line_indent = Inches(-0.5)
        for segment in segments:
            _add_segment(paragraph, segment)
    result = BytesIO()
    document.save(result)
    return result.getvalue()
