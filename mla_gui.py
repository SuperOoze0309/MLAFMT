"""
mla_gui.py  (v2.4)
Desktop paper and course formatter.  中英西三语 / Chinese-English-Spanish UI.

Home screen = only what most people need:
    1. fill in name + title   2. drop the draft   3. click Convert
Everything else lives in  ⚙ Advanced / 高级选项  (secondary window).
"""

import json
import datetime
import locale
import os
import subprocess
import queue
import re
import shutil
import tempfile
import threading
from pathlib import Path
import sys
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from tkinter import font as tkfont

from mla_formatter import (
    MLG_DISCLAIMER,
    SUPPORTED_EXTS,
    __version__,
    convert_draft_to_mla,
    create_blank_template,
    date_warning,
    extract_last_name,
    format_summary,
    generate_filename,
    inspect_draft_paragraphs,
    today_mla,
    today_for_format,
)
from format_profiles import (
    BUILTIN_CUSTOM_PRESETS, DEFAULT_CUSTOM_PROFILE, FONT_CHOICES,
    normalize_custom_profile,
)

try:  # optional drag-and-drop
    from tkinterdnd2 import DND_FILES, TkinterDnD
    HAS_DND = True
except ImportError:
    HAS_DND = False

def _config_file(platform_name=None, environ=None, user_home=None):
    """Use each desktop's normal per-user application settings location."""
    platform_name = sys.platform if platform_name is None else platform_name
    environ = os.environ if environ is None else environ
    user_home = Path.home() if user_home is None else Path(user_home)
    if platform_name.startswith("win"):
        base = Path(environ.get("APPDATA") or user_home / "AppData" / "Roaming")
    elif platform_name == "darwin":
        base = user_home / "Library" / "Application Support"
    else:
        configured = Path(environ.get("XDG_CONFIG_HOME", ""))
        base = configured if configured.is_absolute() else user_home / ".config"
    return str(base / "MLAFMT" / "settings.json")


def _wheel_units(event):
    """Normalize Linux mouse buttons and macOS/Windows wheel deltas."""
    number = getattr(event, "num", None)
    if number in (4, 5):
        return -1 if number == 4 else 1
    delta = getattr(event, "delta", 0)
    if not delta:
        return 0
    return int(-delta / 120) if abs(delta) >= 120 else (-1 if delta > 0 else 1)


CONFIG_PATH = _config_file()
LEGACY_CONFIG_PATH = os.path.join(os.path.expanduser("~"), ".mlafmt.json")
PREVIOUS_CONFIG_PATH = (str(Path.home() / ".config" / "MLAFMT" / "settings.json")
                        if sys.platform == "darwin" else None)

# ──────────────────────────────────────────────────────────────────────
#  Strings
# ──────────────────────────────────────────────────────────────────────
T = {
    "zh": {
        "app": "MLA 一键排版",
        "tagline": "把你的英文作文草稿，一键变成符合 MLA 第 9 版格式的 Word 文档。",
        "what": "自动设置：Times New Roman 12 号 · 双倍行距 · 1 英寸页边距 · 右上角“姓 页码” · "
                "左上角个人信息 · 标题居中 · 首行缩进 · Works Cited 另起一页并按字母排序",
        "step1": "① 填写信息",
        "step2": "② 选择草稿",
        "step3": "③ 生成",
        "name": "你的姓名 *",
        "title": "作文标题 *",
        "instructor": "老师",
        "course": "课程",
        "date": "日期",
        "name_ph": "例：Jane Doe",
        "title_ph": "例：The Role of Memory in Hamlet",
        "instructor_ph": "例：Professor Smith",
        "course_ph": "例：ENG 101",
        "drop": "把草稿拖到这里，或点击选择文件\n支持 .docx / .txt / .md",
        "drop_nodnd": "点击这里选择草稿文件\n支持 .docx / .txt / .md",
        "convert": "转换为 MLA 格式",
        "blank": "还没写？生成空白 MLA 模板",
        "advanced": "⚙ 高级选项",
        "lang_btn": "English",
        "ready": "准备就绪：填好姓名和标题，选择草稿即可。",
        "loaded": "已选择：{f}",
        "need_file": "请先选择或拖入一个草稿文件。",
        "need_name": "请填写你的姓名。",
        "need_title": "请填写作文标题。",
        "date_tip": "日期建议用 MLA 格式（日 月 年），例如 “{d}”。\n\n仍然继续吗？",
        "date_tip_title": "检查日期",
        "unsupported": "只支持 .docx、.txt、.md 文件。\n（.doc / Pages / Google Docs 请先另存为 .docx）",
        "not_found": "找不到这个文件。",
        "save_as": "保存 MLA 文档为…",
        "save_blank": "保存空白模板为…",
        "same_file": "不能覆盖原草稿，请换一个文件名。",
        "working": "正在生成…",
        "done": "✓ 已保存：{f}（{b} 段正文，{w} 条参考文献{r}）",
        "removed": "，已去掉草稿里重复的 {n} 行抬头",
        "done_blank": "✓ 空白模板已保存：{f}",
        "empty": "草稿里没有读到正文，请检查文件。",
        "perm": "无法写入文件。如果它正在 Word 里打开，请先关闭再试，或换个文件夹。",
        "bad_doc": "无法读取这个文档，文件可能已损坏或格式不对。",
        "error": "出错了：{e}",
        "menu_file": "文件",
        "menu_open": "选择草稿…",
        "menu_blank": "生成空白模板…",
        "menu_quit": "退出",
        "menu_settings": "设置",
        "menu_help": "帮助",
        "menu_howto": "使用说明",
        "menu_about": "关于",
        # advanced window
        "adv_title": "高级选项",
        "adv_header": "页眉",
        "lastname": "页眉里的姓",
        "lastname_hint": "留空则自动取姓名最后一个词（会自动忽略 Jr. / III 等）。\n"
                         "复姓或特殊情况手动填写，例如 “Smith-Jones”。",
        "adv_txt": "纯文本（.txt / .md）草稿",
        "mode_blank": "空一行算一段（推荐）",
        "mode_line": "每一行都算一段",
        "md": "用 *星号* 包住的文字变斜体（书名、期刊名）",
        "adv_body": "正文处理",
        "strip": "去掉草稿里已有的姓名/标题抬头，避免重复",
        "headings": "自动识别小标题并加粗（实验性）",
        "bq": "以 > 开头的段落作为长引文缩进 0.5 英寸",
        "adv_after": "保存后",
        "open_after": "自动打开生成的文档",
        "adv_preview": "将应用的格式",
        "close": "完成",
        "howto": "1. 填写你的姓名和作文标题（老师、课程可选，日期默认今天）。\n"
                 "2. 把草稿 .docx 或 .txt 拖进窗口，或点击选择。\n"
                 "3. 点「转换为 MLA 格式」，选择保存位置。\n\n"
                 "小提示：\n"
                 "• 草稿最后写一行 “Works Cited”，下面每段一条文献，会自动另起一页、排序、悬挂缩进。\n"
                 "• Word 草稿里的斜体会保留；纯文本里用 *书名* 表示斜体。\n"
                 "• 其它设置在「⚙ 高级选项」。",
        "about": "MLA 一键排版 v{v}\n按 MLA Handbook 第 9 版排版。\n\n"
                 "本工具只处理正文文字：斜体/粗体/下划线会保留，图片、脚注、表格、列表等可能丢失。"
                 "提交前请自己再检查一遍。",
    },
    "en": {
        "app": "MLA Formatter",
        "tagline": "Turn your essay draft into a correctly formatted MLA (9th ed.) Word document in one click.",
        "what": "Sets up automatically: Times New Roman 12 pt · double spacing · 1-inch margins · "
                "\"LastName page#\" header · heading block · centered title · indents · "
                "Works Cited on a new page, alphabetized",
        "step1": "① Your info",
        "step2": "② Your draft",
        "step3": "③ Create",
        "name": "Your name *",
        "title": "Essay title *",
        "instructor": "Instructor",
        "course": "Course",
        "date": "Date",
        "name_ph": "e.g. Jane Doe",
        "title_ph": "e.g. The Role of Memory in Hamlet",
        "instructor_ph": "e.g. Professor Smith",
        "course_ph": "e.g. ENG 101",
        "drop": "Drag your draft here, or click to choose a file\n.docx / .txt / .md",
        "drop_nodnd": "Click here to choose your draft\n.docx / .txt / .md",
        "convert": "Convert to MLA",
        "blank": "No draft yet? Create a blank MLA template",
        "advanced": "⚙ Advanced",
        "lang_btn": "中文",
        "ready": "Ready: enter your name and title, then choose a draft.",
        "loaded": "Selected: {f}",
        "need_file": "Please choose or drop a draft file first.",
        "need_name": "Please enter your name.",
        "need_title": "Please enter the essay title.",
        "date_tip": "MLA dates are day month year, e.g. \"{d}\".\n\nContinue anyway?",
        "date_tip_title": "Check date",
        "unsupported": "Only .docx, .txt and .md files are supported.\n"
                       "(For .doc / Pages / Google Docs, save as .docx first.)",
        "not_found": "That file does not exist.",
        "save_as": "Save MLA document as…",
        "save_blank": "Save blank template as…",
        "same_file": "That would overwrite your draft — please pick a different name.",
        "working": "Working…",
        "done": "✓ Saved: {f} ({b} paragraphs, {w} Works Cited entries{r})",
        "removed": ", removed {n} duplicate heading line(s)",
        "done_blank": "✓ Blank template saved: {f}",
        "empty": "No body text was found in the draft. Please check the file.",
        "perm": "Can't write that file. If it is open in Word, close it and try again, "
                "or choose another folder.",
        "bad_doc": "Can't read that document — it may be damaged or not a real .docx.",
        "error": "Something went wrong: {e}",
        "menu_file": "File",
        "menu_open": "Choose draft…",
        "menu_blank": "Create blank template…",
        "menu_quit": "Quit",
        "menu_settings": "Settings",
        "menu_help": "Help",
        "menu_howto": "How to use",
        "menu_about": "About",
        "adv_title": "Advanced options",
        "adv_header": "Page header",
        "lastname": "Last name in header",
        "lastname_hint": "Blank = last word of your name (Jr., III etc. are skipped).\n"
                         "Fill in for special cases, e.g. \"Smith-Jones\".",
        "adv_txt": "Plain-text (.txt / .md) drafts",
        "mode_blank": "Blank line between paragraphs (recommended)",
        "mode_line": "Every line is a paragraph",
        "md": "*Asterisks* make italics (titles of books, journals)",
        "adv_body": "Body text",
        "strip": "Remove a name/title heading already in the draft (no duplicates)",
        "headings": "Detect section headings and bold them (experimental)",
        "bq": "Paragraphs starting with > become block quotes (0.5 in)",
        "adv_after": "After saving",
        "open_after": "Open the document automatically",
        "adv_preview": "Formatting that will be applied",
        "close": "Done",
        "howto": "1. Enter your name and essay title (instructor/course optional, date defaults to today).\n"
                 "2. Drag your .docx or .txt draft into the window, or click to choose it.\n"
                 "3. Click \"Convert to MLA\" and choose where to save.\n\n"
                 "Tips:\n"
                 "• End the draft with a line \"Works Cited\" and one source per paragraph — "
                 "it goes on a new page, sorted, with hanging indents.\n"
                 "• Italics in a Word draft are kept; in plain text write *Title*.\n"
                 "• Everything else is under \"⚙ Advanced\".",
        "about": "MLA Formatter v{v}\nFollows the MLA Handbook, 9th edition.\n\n" + MLG_DISCLAIMER,
    },
}


T["zh"].update({
    "hero": "让草稿，成为一篇好论文。", "eyebrow": "专注写作 · 排版交给我们",
    "tagline": "填写信息，选择草稿。几秒钟，得到整洁的 MLA Word 文档。",
    "step1": "01  论文信息", "step2": "02  导入草稿", "advanced": "高级选项",
    "name": "姓名 *", "title": "论文标题 *", "blank": "生成空白模板",
    "preview": "排版预览", "preview_hint": "布局示意 · 实际正文以生成文档为准",
    "preview_name": "Your Name", "preview_title": "Your Essay Title",
    "spec_title": "MLA 9 格式，自动就位", "spec_font": "Times New Roman · 12 号",
    "spec_spacing": "双倍行距 · 1 英寸页边距", "spec_header": "姓氏 + 自动页码",
    "spec_cited": "Works Cited · 排序与悬挂缩进",
    "local": "在本机处理，草稿无需上传", "draft_types": "Word / 文本 / Markdown 草稿",
    "word_type": "Word 文档", "text_type": "文本 / Markdown", "choose": "选择草稿",
    "remove": "移除文件", "selected": "已选择草稿", "file_size": "{size} KB · 点击更换",
    "limit": "支持 .docx / .txt / .md · 最大 20 MB", "too_large": "文件超过 20 MB，请选择较小的草稿。",
    "disclaimer": "保留斜体、粗体和下划线。图片、表格、脚注等可能丢失，提交前请检查。",
    "status_ready": "就绪 · 先填写姓名与标题，再选择草稿。",
    "done": "已保存 {f} · {b} 段正文 · {w} 条参考文献{r}",
    "done_blank": "空白模板已保存：{f}", "working": "正在生成文档…",
    "done_removed": "已保存 {f} · {b} 段正文 · {w} 条参考文献 · 去重 {r} 行抬头",
    "open_result": "打开文档", "view_folder": "打开文件夹", "menu_language": "语言 / Language",
    "reveal_result": "在 Finder 中显示", "empty_file": "这个文件是空的，请选择有内容的草稿。",
    "queued_file": "当前操作结束后将载入：{f}",
    "saved_next": "已保存 {f}；已载入下一份草稿：{next}",
    "one_draft": "已选择 {f}；每次处理一份草稿。",
    "format_preview": "MLA 格式预览", "missing_source": "草稿已被移动或删除，请重新选择。",
    "preview_header": "页眉：{name} 1", "preview_specs": "Times New Roman 12 号\n双倍行距 · 1 英寸页边距\n正文首行缩进 0.5 英寸\nWorks Cited 另起一页，悬挂缩进",
})
T["en"].update({
    "hero": "Your words. Beautifully in order.", "eyebrow": "FOCUS ON YOUR WORDS",
    "tagline": "Add your details and a draft. Get a clean MLA Word document in seconds.",
    "step1": "01  Paper details", "step2": "02  Add your draft", "advanced": "Advanced options",
    "name": "Your name *", "title": "Essay title *", "blank": "Create blank template",
    "preview": "Paper preview", "preview_hint": "Layout illustration · your document contains your draft",
    "preview_name": "Your Name", "preview_title": "Your Essay Title",
    "spec_title": "MLA 9, taken care of", "spec_font": "Times New Roman · 12 pt",
    "spec_spacing": "Double spacing · 1-inch margins", "spec_header": "Last name + automatic page number",
    "spec_cited": "Works Cited · sorted, hanging indents",
    "local": "Processed on your device. No upload needed.", "draft_types": "Word / text / Markdown drafts",
    "word_type": "Word documents", "text_type": "Text / Markdown", "choose": "Choose your draft",
    "remove": "Remove file", "selected": "Draft selected", "file_size": "{size} KB · click to replace",
    "limit": ".docx / .txt / .md · up to 20 MB", "too_large": "This file exceeds 20 MB. Please choose a smaller draft.",
    "disclaimer": "Italic, bold and underline are kept. Images, tables and footnotes may be lost. Always proofread.",
    "status_ready": "Ready · add your name and title, then choose a draft.",
    "done": "Saved {f} · {b} paragraphs · {w} sources{r}",
    "done_blank": "Blank template saved: {f}", "working": "Generating document…",
    "done_removed": "Saved {f} · {b} paragraphs · {w} sources · removed {r} heading lines",
    "open_result": "Open document", "view_folder": "Show in folder", "menu_language": "Language / 语言",
    "reveal_result": "Show in Finder", "empty_file": "This file is empty. Choose a draft with text.",
    "queued_file": "This draft will open after the current operation: {f}",
    "saved_next": "Saved {f}; next draft loaded: {next}",
    "one_draft": "Selected {f}; one draft is processed at a time.",
    "format_preview": "MLA format preview", "missing_source": "Your draft was moved or deleted. Please select it again.",
    "preview_header": "Header: {name} 1", "preview_specs": "Times New Roman 12 pt\nDouble spacing · 1-inch margins\nBody first-line indent: 0.5 inch\nWorks Cited: new page, hanging indents",
})

T["zh"].update({
    "app": "论文与课程文档排版", "tagline": "选择格式，填写信息，得到整洁的 Word 文档。",
    "convert": "生成 Word 文档", "save_as": "保存排版文档为…",
    "format": "文档格式", "format_mla": "MLA 9", "format_apa": "APA 7 · 学生论文",
    "format_custom": "中文预设 / 自定义模板", "affiliation": "学校 / 院系",
    "template": "课程模板", "edit_template": "编辑 / 新建模板", "template_title": "课程模板设置",
    "template_hint": "保存老师要求的版式；选择同名模板后保存会更新该模板。最多保存 20 个。",
    "builtin_prefix": "内置", "saved_prefix": "我的", "template_copy": "副本",
    "builtin_hint": "中文预设提供常用基础排版，请以学校要求为准；不生成目录、摘要或学校专用封面，也不改写参考文献。编辑后保存为个人副本。",
    "builtin_readonly": "内置预设不能删除；可编辑并保存为个人副本。",
    "typography": "文字排版", "typography_western": "西文排版", "typography_chinese": "中文论文排版",
    "latin_font_name": "西文与数字字体", "line_spacing_pt": "固定行距（磅，0 使用倍数）",
    "first_line_chars": "首行缩进（字符，0 使用厘米）",
    "chinese_controls_hint": "固定行距大于 0 时使用固定磅值；字符缩进大于 0 时使用字符数。中文排版分别设置中文与西文字体，并使用黑体标题。",
    "profile_label": "模板名称", "paper_size": "纸张", "font_name": "字体", "font_size": "字号（pt）",
    "line_spacing": "行距（倍）", "margin_cm": "四边页边距（cm）", "first_line_cm": "正文首行缩进（cm）",
    "alignment": "正文对齐", "align_left": "左对齐", "align_justify": "两端对齐",
    "header": "页码内容", "header_surname_page": "姓氏 + 页码", "header_page": "仅页码", "header_none": "不显示页码",
    "title_page": "单独封面页", "references_title": "参考文献页标题", "sort_references": "按字母顺序排列参考文献",
    "save_template": "保存并使用", "import_template": "导入 JSON", "export_template": "导出 JSON",
    "delete_template": "删除当前模板", "template_invalid": "模板设置有误：{e}",
    "template_limit": "最多保存 20 个模板。请删除一个模板，或使用已有模板的名称。",
    "template_saved": "课程模板已保存：{label}", "template_imported": "已导入课程模板：{label}",
    "template_session_only": "已使用模板 {label}，但无法保存设置；关闭应用后此修改可能丢失。",
    "template_io": "无法读写模板文件：{e}", "template_json": "FormCapy 模板（JSON）",
    "template_too_large": "模板文件超过 64 KB，请选择一个课程模板 JSON 文件。",
    "advanced_custom": "高级定制…", "advanced_custom_title": "高级定制",
    "advanced_custom_hint": "文字排版、缩进、行距和页码设置会随课程模板保存。",
    "page_number_position": "页码位置", "page_number_top_right": "右上角",
    "page_number_bottom_center": "页脚居中", "page_number_bottom_right": "页脚右侧",
    "show_cover_page_number": "在封面显示页码", "page_number_start": "起始页码",
    "page_number_hint": "隐藏封面页码时，起始数字用于正文第一页；没有独立封面时用于文档第一页。",
    "page_number_summary": "{position} · 从 {start} 开始", "cover_unnumbered": "封面不编号",
    "page_number_invalid": "起始页码必须是 1 到 9999 的整数。",
    "preserve_title": "保持原格式的段落", "preserve_choose": "选择保留原格式的段落…",
    "preserve_hint": "选中的 Word 段落保留字体、字号、行距和缩进；纸张、页边距与页码使用所选模板。",
    "preserve_scope": "选择仅用于当前草稿，不保存进课程模板。",
    "preserve_docx_only": "请先选择 Word（.docx）草稿，再选择要保留原格式的段落。",
    "preserve_count": "当前草稿：{n} 段保留原格式", "preserve_none": "当前草稿：全部段落应用所选格式",
    "preserve_column": "保留", "preserve_text": "原草稿段落", "preserve_images": "含图片",
    "preserve_toggle": "点击一行切换保留状态，也可用空格键；原草稿的段落编号可能有间隔。",
    "preserve_ranges": "按编号选择（如 2, 4-6）", "preserve_apply_ranges": "应用编号",
    "preserve_select_all": "全选", "preserve_clear": "清空选择", "preserve_apply": "使用所选段落",
    "preserve_invalid_ranges": "请输入列表中存在的段落编号或范围，例如 2, 4-6。",
    "preserve_empty": "这份 Word 草稿没有可选择的段落。", "preserve_loading_error": "无法读取原稿段落：{e}",
    "draft_changed": "草稿已更新，请重新打开段落选择并确认保留范围。",
    "format_preview": "排版设置", "spec_title": "本次应用的版式",
    "apa_hint": "APA 7 学生版：封面、页码和正文标题。请补齐学校、课程、老师；引文内容保持原文。",
    "custom_hint": "按所选课程模板排版；文献内容保持原文。",
    "mla_hint": "MLA 9：姓氏页码、首面信息块、Works Cited 新页。",
    "citation_hint": "引文和文献信息保持原文，请按课程要求核对。",
    "review_result": "请对照原稿检查：{items}", "review_omitted": "{n} 项内容未保留",
    "review_headings": "{n} 项不支持的标题按正文处理",
    "retained_result": "已保留：{items}", "retained_paragraphs": "{n} 段原格式",
    "retained_images": "正文图片 {n} 项", "retained_tables": "正文表格 {n} 个",
    "review_reference_order": "保留段落或对象之间仅对连续普通文献排序，请核对完整顺序",
    "menu_notices": "第三方许可证", "notices_missing": "尚未找到当前构建的许可证文件。请先使用构建脚本生成 THIRD-PARTY-NOTICES.txt。",
    "menu_source_kit": "保存对应源码…", "source_kit_saved": "已保存对应源码与重建说明：{path}",
    "unsupported_document": "文档超过处理范围，请减少内容或拆分文件后重试。",
    "disclaimer": "保留 DOCX 正文图片和表格；页面布局按当前格式设置。复杂文档请核对结果。",
    "howto": "1. 选择 MLA、APA 学生论文或自定义课程模板。\n"
             "2. 填写姓名与标题，导入 .docx / .txt / .md 草稿。\n"
             "3. 生成并保存 Word 文档。\n\n"
             "自定义模板可保存、导入和导出老师要求的版式。APA / 自定义模式可填写学校 / 院系。\n"
             "高级选项可选择 Word 中保持原格式的段落。\n"
             "在草稿末尾写 Works Cited / References / Bibliography，每段一条文献。\n"
             "本工具调整版式，不会把引文或参考文献内容自动转换成另一规范。",
    "about": "FormCapy v{v}\n支持 MLA 9、APA 7 学生论文和自定义课程模板。\n\n"
             "保留 DOCX 正文图片和表格；页面布局按当前格式设置。\n"
             "选中段落保留字体、字号与段落格式；原分页位置可能改变。复杂文档请核对结果与引用。",
    "chinese_course": "中文课程论文（通用）", "chinese_thesis": "中文毕业论文（基础版）",
    "list_separator": "；", "preview_affiliation": "学校 / 院系",
    "preview_instructor_default": "Professor Smith", "preview_course_default": "ENG 101",
    "spec_apa_font": "Times New Roman · 12 号",
    "spec_apa_spacing": "双倍行距 · Letter · 1 英寸页边距",
    "spec_apa_titlepage": "独立封面 · 右上角自动页码",
    "spec_apa_references": "References 新页 · 悬挂缩进 0.5 英寸",
    "spacing_exact_line": "固定 {pt:g} 磅行距",
    "spacing_multiple_line": "{spacing:g} 倍行距",
    "custom_spacing_margin": "{spacing} · 页边距 {margin:g} cm",
    "custom_first_line": "首行缩进 {indent} · {alignment}",
    "indent_chars": "{n:g} 字符",
    "title_page_suffix": " · 独立封面",
    "reference_new_page": "{title} · 另起一页",
    "reference_sorted": " · 字母排序",
    "reference_original_gui": " · 保留顺序",
})
T["en"].update({
    "app": "Paper and Course Formatter", "tagline": "Choose a format, add your details, and get a clean Word document.",
    "done": "✓ Saved: {f} ({b} paragraphs, {w} reference entries{r})",
    "convert": "Create Word document", "save_as": "Save formatted document as…",
    "format": "Document format", "format_mla": "MLA 9", "format_apa": "APA 7 · Student paper",
    "format_custom": "Chinese presets / Custom template", "affiliation": "School / department",
    "template": "Course template", "edit_template": "Edit / create template", "template_title": "Course template settings",
    "template_hint": "Save your instructor's layout requirements. Saving an existing name updates it. Up to 20 templates.",
    "builtin_prefix": "Built-in", "saved_prefix": "My templates", "template_copy": "copy",
    "builtin_hint": "Chinese presets provide a basic layout. Follow your school's requirements; no contents, abstract or school-specific cover is generated, and reference text is unchanged. Edit and save a personal copy.",
    "builtin_readonly": "Built-in presets cannot be deleted. Edit one and save a personal copy.",
    "typography": "Text layout", "typography_western": "Western", "typography_chinese": "Chinese paper",
    "latin_font_name": "Latin / numeral font", "line_spacing_pt": "Exact line spacing (pt; 0 uses multiple)",
    "first_line_chars": "First-line indent (characters; 0 uses cm)",
    "chinese_controls_hint": "A positive point value uses exact spacing; a positive character value uses character indentation. Chinese layout uses separate Chinese/Latin fonts and Heiti headings.",
    "profile_label": "Template name", "paper_size": "Paper size", "font_name": "Font", "font_size": "Font size (pt)",
    "line_spacing": "Line spacing (multiple)", "margin_cm": "All page margins (cm)", "first_line_cm": "Body first-line indent (cm)",
    "alignment": "Body alignment", "align_left": "Left", "align_justify": "Justified",
    "header": "Page-number content", "header_surname_page": "Surname + page number", "header_page": "Page number only", "header_none": "No page numbers",
    "title_page": "Separate title page", "references_title": "Reference page title", "sort_references": "Sort reference entries alphabetically",
    "save_template": "Save and use", "import_template": "Import JSON", "export_template": "Export JSON",
    "delete_template": "Delete current template", "template_invalid": "Check the template settings: {e}",
    "template_limit": "You can save up to 20 templates. Delete one or use an existing template name.",
    "template_saved": "Course template saved: {label}", "template_imported": "Course template imported: {label}",
    "template_session_only": "Using {label}, but settings could not be saved. This change may be lost when the app closes.",
    "template_io": "Could not read or write the template: {e}", "template_json": "FormCapy template (JSON)",
    "template_too_large": "This template exceeds 64 KB. Choose a course template JSON file.",
    "advanced_custom": "Advanced customization…", "advanced_custom_title": "Advanced customization",
    "advanced_custom_hint": "Text layout, indentation, spacing and page-number settings are saved with the course template.",
    "page_number_position": "Page-number position", "page_number_top_right": "Top right",
    "page_number_bottom_center": "Bottom center", "page_number_bottom_right": "Bottom right",
    "show_cover_page_number": "Show a page number on the cover", "page_number_start": "Starting page number",
    "page_number_hint": "When the cover number is hidden, this starts the first body page. Without a separate cover, it starts the first document page.",
    "page_number_summary": "{position} · starts at {start}", "cover_unnumbered": "Cover unnumbered",
    "page_number_invalid": "Starting page number must be an integer from 1 to 9999.",
    "preserve_title": "Keep selected paragraphs unchanged", "preserve_choose": "Choose paragraphs to keep unchanged…",
    "preserve_hint": "Selected Word paragraphs keep their font, size, line spacing and indents. Paper size, margins and page numbers follow the selected template.",
    "preserve_scope": "This selection belongs to the current draft and is not saved in a course template.",
    "preserve_docx_only": "Choose a Word (.docx) draft before selecting paragraphs to keep unchanged.",
    "preserve_count": "Current draft: keep {n} paragraphs unchanged", "preserve_none": "Current draft: apply the selected format to all paragraphs",
    "preserve_column": "Keep", "preserve_text": "Original draft paragraph", "preserve_images": "Includes an image",
    "preserve_toggle": "Click a row or press Space to toggle it. Original paragraph numbers may have gaps.",
    "preserve_ranges": "Select numbers (e.g. 2, 4-6)", "preserve_apply_ranges": "Apply numbers",
    "preserve_select_all": "Select all", "preserve_clear": "Clear selection", "preserve_apply": "Use selected paragraphs",
    "preserve_invalid_ranges": "Enter paragraph numbers or ranges from the list, such as 2, 4-6.",
    "preserve_empty": "This Word draft has no paragraphs to select.", "preserve_loading_error": "Could not read the draft paragraphs: {e}",
    "draft_changed": "The draft has changed. Reopen paragraph selection and confirm what to keep.",
    "format_preview": "Formatting settings", "spec_title": "Formatting for this document",
    "apa_hint": "APA 7 student layout: title page, page numbers and repeated paper title. Add your affiliation, course and instructor. Citation text is unchanged.",
    "custom_hint": "Apply the selected course template. Reference text is unchanged.",
    "mla_hint": "MLA 9: surname/page header, first-page details, and a new Works Cited page.",
    "citation_hint": "Citation and reference text is unchanged. Check it against your course requirements.",
    "review_result": "Compare with your original draft: {items}", "review_omitted": "{n} content elements were not retained",
    "review_headings": "{n} unsupported headings were kept as body text",
    "retained_result": "Retained: {items}", "retained_paragraphs": "{n} paragraphs with their original formatting",
    "retained_images": "{n} body image elements", "retained_tables": "{n} body tables",
    "review_reference_order": "Only consecutive plain references between retained paragraphs or objects were sorted; check the full order",
    "menu_notices": "Third-party licenses", "notices_missing": "No license file was found for this build. Generate THIRD-PARTY-NOTICES.txt using the build script first.",
    "menu_source_kit": "Save corresponding source…", "source_kit_saved": "Saved corresponding source and rebuild instructions: {path}",
    "unsupported_document": "This document exceeds the processing limits. Reduce its size or split it into smaller files.",
    "disclaimer": "DOCX body images and tables are retained; page layout follows the selected format. Review complex documents.",
    "howto": "1. Choose MLA, APA student paper, or a custom course template.\n"
             "2. Add your name and title, then choose a .docx / .txt / .md draft.\n"
             "3. Create and save your Word document.\n\n"
             "Save, import, or export custom templates for your instructor's requirements. APA and custom formats include a school / department field.\n"
             "Use Advanced options to choose Word paragraphs that keep their original formatting.\n"
             "End your draft with Works Cited / References / Bibliography and one source per paragraph.\n"
             "The tool adjusts layout; it does not convert citation or reference text between citation styles.",
    "about": "FormCapy v{v}\nMLA 9, APA 7 student papers, and custom course templates.\n\n"
             "DOCX body images and tables are retained; page layout follows the selected format.\n"
             "Selected paragraphs keep their font, size and paragraph formatting; page breaks may move. Review complex documents and citations.",
    "chinese_course": "Chinese course paper · General", "chinese_thesis": "Chinese graduation paper · Basic",
    "list_separator": "; ", "preview_affiliation": "School / Department",
    "preview_instructor_default": "Professor Smith", "preview_course_default": "ENG 101",
    "spec_apa_font": "Times New Roman · 12 pt",
    "spec_apa_spacing": "Double spacing · Letter · 1-inch margins",
    "spec_apa_titlepage": "Title page · automatic page-number header",
    "spec_apa_references": "References: new page · 0.5-inch hanging indents",
    "spacing_exact_line": "Exactly {pt:g} pt",
    "spacing_multiple_line": "{spacing:g}× spacing",
    "custom_spacing_margin": "{spacing} · {margin:g} cm margins",
    "custom_first_line": "First-line indent: {indent} · {alignment}",
    "indent_chars": "{n:g} characters",
    "title_page_suffix": " · title page",
    "reference_new_page": "{title} · new page",
    "reference_sorted": " · alphabetical",
    "reference_original_gui": " · original order",
})

T["es"] = {
    **T["en"],
    "app": "Formato de trabajos y cursos",
    "tagline": "Elige un formato, añade tus datos y obtén un documento Word limpio.",
    "what": "Configura automáticamente: Times New Roman 12 pt · interlineado doble · márgenes de 1 pulgada · "
            "encabezado «apellido página» · bloque de datos · título centrado · sangrías · "
            "Works Cited en página nueva y por orden alfabético",
    "step1": "① Tus datos",
    "step2": "② Tu borrador",
    "step3": "③ Crear",
    "name": "Tu nombre *",
    "title": "Título del trabajo *",
    "instructor": "Profesor/a",
    "course": "Curso",
    "date": "Fecha",
    "name_ph": "p. ej., Jane Doe",
    "title_ph": "p. ej., The Role of Memory in Hamlet",
    "instructor_ph": "p. ej., Professor Smith",
    "course_ph": "p. ej., ENG 101",
    "drop": "Arrastra tu borrador aquí o haz clic para elegir un archivo\n.docx / .txt / .md",
    "drop_nodnd": "Haz clic aquí para elegir tu borrador\n.docx / .txt / .md",
    "convert": "Convertir a MLA",
    "blank": "¿Aún no tienes borrador? Crear una plantilla MLA vacía",
    "advanced": "⚙ Opciones avanzadas",
    "lang_btn": "Español",
    "ready": "Listo: escribe tu nombre y el título y elige un borrador.",
    "loaded": "Seleccionado: {f}",
    "need_file": "Elige o arrastra primero un borrador.",
    "need_name": "Escribe tu nombre.",
    "need_title": "Escribe el título del trabajo.",
    "date_tip": "Las fechas MLA usan el orden día mes año, p. ej., «{d}».\n\n¿Continuar de todos modos?",
    "date_tip_title": "Comprobar fecha",
    "unsupported": "Solo se admiten archivos .docx, .txt y .md.\n"
                   "(Para .doc / Pages / Google Docs, guárdalos antes como .docx).",
    "not_found": "Ese archivo no existe.",
    "save_as": "Guardar el documento MLA como…",
    "save_blank": "Guardar la plantilla vacía como…",
    "same_file": "Eso sobrescribiría tu borrador; elige otro nombre.",
    "working": "Generando…",
    "done": "✓ Guardado: {f} ({b} párrafos, {w} referencias{r})",
    "removed": ", se eliminaron {n} líneas de encabezado duplicadas",
    "done_blank": "✓ Plantilla vacía guardada: {f}",
    "empty": "No se encontró texto en el borrador. Revisa el archivo.",
    "perm": "No se puede escribir ese archivo. Si está abierto en Word, ciérralo e inténtalo de nuevo, "
            "o elige otra carpeta.",
    "bad_doc": "No se puede leer ese documento; puede estar dañado o no ser un .docx real.",
    "error": "Algo salió mal: {e}",
    "menu_file": "Archivo",
    "menu_open": "Elegir borrador…",
    "menu_blank": "Crear plantilla vacía…",
    "menu_quit": "Salir",
    "menu_settings": "Ajustes",
    "menu_help": "Ayuda",
    "menu_howto": "Cómo se usa",
    "menu_about": "Acerca de",
    "adv_title": "Opciones avanzadas",
    "adv_header": "Encabezado de página",
    "lastname": "Apellido en el encabezado",
    "lastname_hint": "Vacío = última palabra de tu nombre (se omiten Jr., III, etc.).\n"
                     "Rellénalo en casos especiales, p. ej., «Smith-Jones».",
    "adv_txt": "Borradores de texto (.txt / .md)",
    "mode_blank": "Línea en blanco entre párrafos (recomendado)",
    "mode_line": "Cada línea es un párrafo",
    "md": "*Los asteriscos* crean cursiva (títulos de libros, revistas)",
    "adv_body": "Texto del cuerpo",
    "strip": "Quitar el encabezado de nombre/título que ya tenga el borrador (sin duplicados)",
    "headings": "Detectar títulos de sección y ponerlos en negrita (experimental)",
    "bq": "Los párrafos que empiezan por > se convierten en citas en bloque (0,5 pulgadas)",
    "adv_after": "Después de guardar",
    "open_after": "Abrir el documento automáticamente",
    "adv_preview": "Formato que se aplicará",
    "close": "Listo",
    "howto": "1. Escribe tu nombre y el título del trabajo (profesor/a y curso son opcionales; la fecha es hoy por defecto).\n"
             "2. Arrastra tu borrador .docx o .txt a la ventana, o haz clic para elegirlo.\n"
             "3. Pulsa «Convertir a MLA» y elige dónde guardarlo.\n\n"
             "Consejos:\n"
             "• Termina el borrador con una línea «Works Cited» y una fuente por párrafo: irá a una página nueva, ordenada y con sangría francesa.\n"
             "• La cursiva de un borrador Word se conserva; en texto plano escribe *Título*.\n"
             "• Todo lo demás está en «⚙ Opciones avanzadas».",
    "about": "Formato de trabajos FormCapy v{v}\nSigue el MLA Handbook, 9.ª edición.\n\n" + MLG_DISCLAIMER,
}
T["es"].update({
    "hero": "Tus palabras, bien ordenadas.",
    "eyebrow": "CÉNTRATE EN ESCRIBIR",
    "tagline": "Añade tus datos y un borrador. Obtén un documento Word MLA limpio en segundos.",
    "step1": "01  Datos del trabajo",
    "step2": "02  Añade tu borrador",
    "advanced": "Opciones avanzadas",
    "name": "Tu nombre *",
    "title": "Título del trabajo *",
    "blank": "Crear plantilla vacía",
    "preview": "Vista del trabajo",
    "preview_hint": "Ilustración del diseño · el documento contendrá tu borrador",
    "preview_name": "Tu nombre",
    "preview_title": "Título de tu trabajo",
    "spec_font": "Times New Roman · 12 pt",
    "spec_spacing": "Interlineado doble · márgenes de 1 pulgada",
    "spec_header": "Apellido + número de página automático",
    "spec_cited": "Works Cited · ordenado y con sangría francesa",
    "local": "Se procesa en tu dispositivo. No hace falta subir nada.",
    "draft_types": "Borradores Word / texto / Markdown",
    "word_type": "Documentos Word",
    "text_type": "Texto / Markdown",
    "choose": "Elige tu borrador",
    "remove": "Quitar archivo",
    "selected": "Borrador seleccionado",
    "file_size": "{size} KB · haz clic para cambiar",
    "limit": ".docx / .txt / .md · hasta 20 MB",
    "too_large": "El archivo supera los 20 MB. Elige un borrador más pequeño.",
    "status_ready": "Listo · añade tu nombre y el título y elige un borrador.",
    "done": "Guardado {f} · {b} párrafos · {w} referencias{r}",
    "done_blank": "Plantilla vacía guardada: {f}",
    "working": "Generando documento…",
    "done_removed": "Guardado {f} · {b} párrafos · {w} referencias · {r} líneas de encabezado eliminadas",
    "open_result": "Abrir documento",
    "view_folder": "Mostrar en la carpeta",
    "menu_language": "Idioma / Language",
    "reveal_result": "Mostrar en Finder",
    "empty_file": "El archivo está vacío. Elige un borrador con texto.",
    "queued_file": "Este borrador se abrirá al terminar la operación actual: {f}",
    "saved_next": "Guardado {f}; siguiente borrador cargado: {next}",
    "one_draft": "Seleccionado {f}; se procesa un borrador a la vez.",
    "missing_source": "Tu borrador se movió o se eliminó. Vuelve a seleccionarlo.",
    "preview_header": "Encabezado: {name} 1",
    "preview_specs": "Times New Roman 12 pt\nInterlineado doble · márgenes de 1 pulgada\nSangría de primera línea: 0,5 pulgadas\nWorks Cited: página nueva, sangría francesa",
    "convert": "Crear documento Word",
    "save_as": "Guardar documento con formato como…",
    "format": "Formato del documento",
    "format_mla": "MLA 9",
    "format_apa": "APA 7 · Trabajo de estudiante",
    "format_custom": "Preajustes chinos / Plantilla personalizada",
    "affiliation": "Escuela / departamento",
    "template": "Plantilla de curso",
    "edit_template": "Editar / crear plantilla",
    "template_title": "Ajustes de la plantilla de curso",
    "template_hint": "Guarda el diseño que pida tu profesor/a. Guardar con un nombre existente lo actualiza. Máximo: 20 plantillas.",
    "builtin_prefix": "Integrada",
    "saved_prefix": "Mis plantillas",
    "template_copy": "copia",
    "builtin_hint": "Los preajustes chinos ofrecen un diseño básico. Sigue los requisitos de tu centro; no se genera índice, resumen ni portada específica, y el texto de las referencias no cambia. Edita y guarda una copia personal.",
    "builtin_readonly": "Los preajustes integrados no se pueden eliminar. Edita uno y guarda una copia personal.",
    "typography": "Composición del texto",
    "typography_western": "Occidental",
    "typography_chinese": "Trabajo en chino",
    "latin_font_name": "Fuente latina / numérica",
    "line_spacing_pt": "Interlineado exacto (pt; 0 usa múltiplo)",
    "first_line_chars": "Sangría de primera línea (caracteres; 0 usa cm)",
    "chinese_controls_hint": "Un valor positivo en puntos usa interlineado exacto; un valor positivo en caracteres usa sangría por caracteres. La composición china usa fuentes china/latina separadas y títulos en Heiti.",
    "profile_label": "Nombre de la plantilla",
    "paper_size": "Tamaño del papel",
    "font_name": "Fuente",
    "font_size": "Tamaño de fuente (pt)",
    "line_spacing": "Interlineado (múltiplo)",
    "margin_cm": "Márgenes de página (cm)",
    "first_line_cm": "Sangría de primera línea del cuerpo (cm)",
    "alignment": "Alineación del cuerpo",
    "align_left": "Izquierda",
    "align_justify": "Justificado",
    "header": "Contenido del número de página",
    "header_surname_page": "Apellido + número de página",
    "header_page": "Solo número de página",
    "header_none": "Sin números de página",
    "title_page": "Portada independiente",
    "references_title": "Título de la página de referencias",
    "sort_references": "Ordenar alfabéticamente las referencias",
    "save_template": "Guardar y usar",
    "import_template": "Importar JSON",
    "export_template": "Exportar JSON",
    "delete_template": "Eliminar plantilla actual",
    "template_invalid": "Revisa los ajustes de la plantilla: {e}",
    "template_limit": "Puedes guardar hasta 20 plantillas. Elimina una o usa el nombre de una existente.",
    "template_saved": "Plantilla de curso guardada: {label}",
    "template_imported": "Plantilla de curso importada: {label}",
    "template_session_only": "Se usa {label}, pero no se pudieron guardar los ajustes. Este cambio puede perderse al cerrar la aplicación.",
    "template_io": "No se pudo leer o escribir la plantilla: {e}",
    "template_json": "Plantilla de FormCapy (JSON)",
    "template_too_large": "La plantilla supera los 64 KB. Elige un archivo JSON de plantilla de curso.",
    "advanced_custom": "Personalización avanzada…",
    "advanced_custom_title": "Personalización avanzada",
    "advanced_custom_hint": "La composición, la sangría, el interlineado y los números de página se guardan con la plantilla de curso.",
    "page_number_position": "Posición del número de página",
    "page_number_top_right": "Arriba a la derecha",
    "page_number_bottom_center": "Pie centrado",
    "page_number_bottom_right": "Pie a la derecha",
    "show_cover_page_number": "Mostrar número en la portada",
    "page_number_start": "Número de página inicial",
    "page_number_hint": "Si ocultas el número de la portada, este número inicia la primera página del cuerpo. Sin portada independiente, inicia la primera página del documento.",
    "page_number_summary": "{position} · empieza en {start}",
    "cover_unnumbered": "Portada sin número",
    "page_number_invalid": "El número de página inicial debe ser un entero entre 1 y 9999.",
    "preserve_title": "Mantener sin cambios los párrafos seleccionados",
    "preserve_choose": "Elegir párrafos que mantener sin cambios…",
    "preserve_hint": "Los párrafos Word seleccionados conservan su fuente, tamaño, interlineado y sangrías. El papel, los márgenes y los números de página siguen la plantilla elegida.",
    "preserve_scope": "Esta selección pertenece al borrador actual y no se guarda en una plantilla de curso.",
    "preserve_docx_only": "Elige un borrador Word (.docx) antes de seleccionar párrafos que mantener sin cambios.",
    "preserve_count": "Borrador actual: mantener {n} párrafos sin cambios",
    "preserve_none": "Borrador actual: aplicar el formato elegido a todos los párrafos",
    "preserve_column": "Mantener",
    "preserve_text": "Párrafo original del borrador",
    "preserve_images": "Incluye una imagen",
    "preserve_toggle": "Haz clic en una fila o pulsa Espacio para cambiarla. Los números de párrafo originales pueden tener huecos.",
    "preserve_ranges": "Seleccionar números (p. ej., 2, 4-6)",
    "preserve_apply_ranges": "Aplicar números",
    "preserve_select_all": "Seleccionar todo",
    "preserve_clear": "Borrar selección",
    "preserve_apply": "Usar párrafos seleccionados",
    "preserve_invalid_ranges": "Introduce números o rangos de la lista, como 2, 4-6.",
    "preserve_empty": "Este borrador Word no tiene párrafos que seleccionar.",
    "preserve_loading_error": "No se pudieron leer los párrafos del borrador: {e}",
    "draft_changed": "El borrador ha cambiado. Vuelve a abrir la selección de párrafos y confirma qué conservar.",
    "format_preview": "Ajustes de formato",
    "spec_title": "Formato de este documento",
    "apa_hint": "Diseño APA 7 de estudiante: portada, números de página y título repetido. Añade tu institución, curso y profesor/a. El texto de las citas no cambia.",
    "custom_hint": "Aplica la plantilla de curso elegida. El texto de las referencias no cambia.",
    "mla_hint": "MLA 9: encabezado apellido/página, datos en la primera página y Works Cited en página nueva.",
    "citation_hint": "El texto de las citas y referencias no cambia. Compruébalo con los requisitos de tu curso.",
    "review_result": "Compara con tu borrador original: {items}",
    "review_omitted": "{n} elementos de contenido no se conservaron",
    "review_headings": "{n} títulos no admitidos se conservaron como texto del cuerpo",
    "retained_result": "Conservado: {items}",
    "retained_paragraphs": "{n} párrafos con su formato original",
    "retained_images": "{n} imágenes del cuerpo",
    "retained_tables": "{n} tablas del cuerpo",
    "review_reference_order": "Solo se ordenaron las referencias normales consecutivas entre párrafos u objetos conservados; revisa el orden completo",
    "menu_notices": "Licencias de terceros",
    "notices_missing": "No se encontró el archivo de licencias de esta compilación. Genera THIRD-PARTY-NOTICES.txt con el script de compilación.",
    "menu_source_kit": "Guardar el código fuente correspondiente…",
    "source_kit_saved": "Código fuente correspondiente e instrucciones de recompilación guardados: {path}",
    "unsupported_document": "El documento supera los límites de procesamiento. Reduce el contenido o divídelo en archivos más pequeños.",
    "disclaimer": "Se conservan las imágenes y tablas del cuerpo DOCX; el diseño de página sigue el formato elegido. Revisa los documentos complejos.",
    "howto": "1. Elige MLA, APA de estudiante o una plantilla de curso personalizada.\n"
             "2. Añade tu nombre y el título, y elige un borrador .docx / .txt / .md.\n"
             "3. Crea y guarda tu documento Word.\n\n"
             "Guarda, importa o exporta plantillas para los requisitos de tu profesor/a. APA y las plantillas personalizadas incluyen un campo de escuela / departamento.\n"
             "Usa Opciones avanzadas para elegir párrafos Word que conserven su formato original.\n"
             "Termina el borrador con Works Cited / References / Bibliography y una fuente por párrafo.\n"
             "La herramienta ajusta el diseño; no convierte el texto de las citas o referencias entre estilos.",
    "about": "Formato de trabajos FormCapy v{v}\nMLA 9, trabajos APA 7 de estudiante y plantillas de curso personalizadas.\n\n"
             "Se conservan las imágenes y tablas del cuerpo DOCX; el diseño de página sigue el formato elegido.\n"
             "Los párrafos seleccionados conservan fuente, tamaño y formato de párrafo; los saltos de página pueden cambiar. Revisa documentos complejos y citas.",
    "chinese_course": "Trabajo de curso en chino (general)",
    "chinese_thesis": "Trabajo de fin de estudios en chino (básico)",
    "list_separator": "; ",
    "preview_affiliation": "Escuela / Departamento",
    "preview_instructor_default": "Profesor/a Ejemplo",
    "preview_course_default": "CURSO 101",
    "spec_apa_font": "Times New Roman · 12 pt",
    "spec_apa_spacing": "Interlineado doble · Letter · márgenes de 1 pulgada",
    "spec_apa_titlepage": "Portada · encabezado automático de página",
    "spec_apa_references": "References: página nueva · sangría francesa de 0,5 pulgadas",
    "spacing_exact_line": "Interlineado exacto de {pt:g} pt",
    "spacing_multiple_line": "Interlineado de {spacing:g}×",
    "custom_spacing_margin": "{spacing} · márgenes de {margin:g} cm",
    "custom_first_line": "Sangría de primera línea: {indent} · {alignment}",
    "indent_chars": "{n:g} caracteres",
    "title_page_suffix": " · portada independiente",
    "reference_new_page": "{title} · página nueva",
    "reference_sorted": " · orden alfabético",
    "reference_original_gui": " · orden original",
})
ACCENT = "#4446a6"
ACCENT_DARK = "#303278"
MUTED = "#686a77"
BG = "#f7f6f2"
CARD = "#ffffff"
BORDER = "#e5e3dd"
TEXT = "#262839"
SOFT = "#f0f0fa"
MAX_FILE_SIZE = 20 * 1024 * 1024


def _default_lang():
    try:
        loc = (locale.getlocale()[0] or "") + os.environ.get("LANG", "")
        if sys.platform.startswith("win"):
            import ctypes
            language_id = ctypes.windll.kernel32.GetUserDefaultUILanguage() & 0x3FF
            if language_id == 0x04:
                return "zh"
            if language_id == 0x0A:
                return "es"
        lowered = loc.lower()
        if "zh" in lowered or "chinese" in lowered:
            return "zh"
        if "spanish" in lowered or re.search(r"(^|[^a-z])es([_-]|$)", lowered):
            return "es"
        return "en"
    except Exception:
        return "en"


def _load_config():
    for path in (CONFIG_PATH, PREVIOUS_CONFIG_PATH, LEGACY_CONFIG_PATH):
        if path is None:
            continue
        try:
            with open(path, encoding="utf-8") as f:
                cfg = json.load(f)
            if isinstance(cfg, dict):
                return cfg
        except (OSError, ValueError):
            continue
    return {}


def _save_config(cfg):
    """Save atomically in the platform's application-settings directory."""
    tmp = None
    try:
        folder = os.path.dirname(CONFIG_PATH)
        os.makedirs(folder, exist_ok=True)
        with tempfile.NamedTemporaryFile(mode="w", encoding="utf-8", dir=folder,
                                         suffix=".tmp", delete=False) as f:
            tmp = f.name
            json.dump(cfg, f, ensure_ascii=False, indent=2)
        os.replace(tmp, CONFIG_PATH)
        return True
    except OSError:
        return False
    finally:
        if tmp and os.path.exists(tmp):
            os.unlink(tmp)


def _same_file(source, output):
    try:
        return os.path.samefile(source, output)
    except OSError:
        return os.path.normcase(os.path.realpath(source)) == os.path.normcase(os.path.realpath(output))


def _launch_files(arguments):
    """Keep document arguments; ignore Finder's legacy process serial number."""
    files = []
    literal = False
    for argument in arguments:
        if argument == "--" and not literal:
            literal = True
        elif literal or not argument.startswith("-"):
            files.append(argument)
    return files


class MLAGui:
    def __init__(self, remember=True, initial_files=()):
        self.remember = remember
        try:
            self.root = TkinterDnD.Tk() if HAS_DND else tk.Tk()
            self.has_dnd = HAS_DND
        except tk.TclError:
            self.root = tk.Tk()
            self.has_dnd = False
        self.root.configure(bg=BG)
        self._brand_image = None
        mascot = Path(getattr(sys, "_MEIPASS", Path(__file__).parent)) / "static" / "capybara-64.png"
        if mascot.is_file():
            try:
                self._brand_image = tk.PhotoImage(master=self.root, file=str(mascot)).subsample(2, 2)
                if sys.platform.startswith("linux"):
                    self.root.iconphoto(True, self._brand_image)
            except tk.TclError:
                pass
        if sys.platform.startswith("win"):
            icon = Path(getattr(sys, "_MEIPASS", Path(__file__).parent / "build" / "windows")) / "MLAFMT.ico"
            if icon.is_file():
                try:
                    self.root.iconbitmap(default=str(icon))
                except tk.TclError:
                    pass
        self.root.geometry("1080x820")
        self.root.minsize(880, 620)
        self.root.protocol("WM_DELETE_WINDOW", self._on_close)
        cfg = _load_config() if remember else {}
        saved_lang = cfg.get("lang")
        self.lang = saved_lang if isinstance(saved_lang, str) and saved_lang in T else _default_lang()
        self._input_file = None
        self._last_output = None
        self._review_counts = (0, 0)
        self._reference_sort_skipped = False
        self._retained_counts = (0, 0, 0)
        self._adv_win = None
        self._template_win = None
        self._preserve_win = None
        self._preserved_paragraphs = set()
        self._preservation_fingerprint = None
        self._preservation_source = None
        self._preserve_count_labels = []
        self._busy = False
        self._progress_after = None
        self._progress_tick = 0
        self._choosing_output = False
        self._closing = False
        self._pending_file = None
        self._results = queue.Queue()
        self._status_state = ("status_ready", "muted", {})
        self.v = {k: tk.StringVar(self.root) for k in
                  ("name", "title", "instructor", "course", "date", "last_name")}
        for key in ("name", "instructor", "course", "last_name"):
            value = cfg.get(key, "")
            self.v[key].set(value if isinstance(value, str) else "")
        affiliation = cfg.get("affiliation", "")
        self.affiliation = tk.StringVar(self.root, value=affiliation if isinstance(affiliation, str) else "")
        saved_format = cfg.get("format_id", "mla")
        self.format_id = tk.StringVar(self.root, value=saved_format if saved_format in ("mla", "apa", "custom") else "mla")
        self.custom_profiles = []
        saved_profiles = cfg.get("custom_profiles", [])
        if isinstance(saved_profiles, list):
            for profile in saved_profiles[:20]:
                if not isinstance(profile, dict):
                    continue
                try:
                    normalized = normalize_custom_profile(profile)
                except (TypeError, ValueError):
                    continue
                if not any(p["label"] == normalized["label"] for p in self.custom_profiles):
                    self.custom_profiles.append(normalized)
        if not self.custom_profiles:
            self.custom_profiles.append(dict(DEFAULT_CUSTOM_PROFILE))
        selected = cfg.get("selected_profile")
        if not isinstance(selected, str) or selected not in [p["label"] for p in self.custom_profiles]:
            selected = self.custom_profiles[0]["label"]
        self.selected_profile = tk.StringVar(self.root, value=selected)
        builtin = cfg.get("selected_builtin", "")
        self.selected_builtin = tk.StringVar(self.root, value=builtin if builtin in
            {preset["id"] for preset in BUILTIN_CUSTOM_PRESETS} else "")
        self.v["date"].set(self._default_date())
        mode = cfg.get("txt_mode", "blank_line")
        self.txt_mode = tk.StringVar(self.root, value=mode if mode in
                                    ("blank_line", "line_by_line") else "blank_line")
        self.opt = {k: tk.BooleanVar(self.root, value=cfg[k] if isinstance(cfg.get(k), bool) else d)
                    for k, d in (("md", True), ("strip", True), ("headings", False),
                                 ("bq", False), ("open_after", False))}
        self._style()
        self._build()
        for var in self.v.values():
            var.trace_add("write", self._draw_preview)
        self.affiliation.trace_add("write", self._draw_preview)
        self.root.bind("<Control-o>", lambda _e: self._on_browse())
        self.root.bind("<Control-Return>", lambda _e: self._on_convert())
        if sys.platform == "darwin":
            self.root.bind("<Command-o>", lambda _e: self._on_browse())
            self.root.bind("<Command-Return>", lambda _e: self._on_convert())
            self.root.bind("<Command-q>", lambda _e: self._on_close())
            self.root.bind("<Command-comma>", lambda _e: self._open_advanced())
            # Finder's Open With and Dock drops arrive as Apple Events, even
            # after launch. Tk handles these without PyInstaller argv emulation.
            self.root.createcommand("::tk::mac::OpenDocument", self._open_documents)
            self.root.createcommand("::tk::mac::ReopenApplication", self._restore_window)
            self.root.createcommand("::tk::mac::Quit", self._on_close)
            self.root.createcommand("::tk::mac::ShowPreferences", self._open_advanced)
        if initial_files:
            self.root.after_idle(lambda: self._open_documents(*initial_files))
        self.root.after(80, self._poll_results)

    def t(self, key, **kw):
        value = T[self.lang].get(key, T["en"].get(key, key))
        return value.format(**kw) if kw else value

    def _style(self):
        st = ttk.Style(self.root)
        st.theme_use("clam")
        family = ("Microsoft YaHei UI" if sys.platform.startswith("win") else
                  tkfont.nametofont("TkDefaultFont", root=self.root).actual("family"))
        self.font = (family, 10)
        st.configure(".", font=self.font, background=BG, foreground=TEXT)
        st.configure("Card.TFrame", background=CARD)
        st.configure("Card.TLabel", background=CARD)
        st.configure("Muted.TLabel", background=CARD, foreground=MUTED, font=(self.font[0], 9))
        st.configure("Step.TLabel", background=CARD, foreground=ACCENT, font=(self.font[0], 10, "bold"))
        for style in ("Card.TCheckbutton", "Card.TRadiobutton"):
            st.configure(style, background=CARD, padding=(0, 4))
        st.configure("Accent.TButton", foreground=CARD, background=ACCENT,
                     font=(self.font[0], 11, "bold"), padding=(14, 11), borderwidth=0)
        st.map("Accent.TButton", background=[("disabled", "#a5a5c5"), ("active", ACCENT_DARK)])
        st.configure("Ghost.TButton", background=CARD, bordercolor=BORDER, padding=(10, 7))
        st.map("Ghost.TButton", background=[("active", SOFT)])
        st.configure("Lang.TRadiobutton", background=CARD, padding=(10, 6), indicatoron=False)
        st.map("Lang.TRadiobutton", background=[("selected", ACCENT), ("active", SOFT)],
               foreground=[("selected", CARD)])
        st.configure("TEntry", padding=(8, 7), fieldbackground=CARD, bordercolor=BORDER)
        st.map("TEntry", bordercolor=[("focus", ACCENT)])

    def _card(self, parent, padding=18):
        outer = tk.Frame(parent, bg=BORDER)
        outer.pack(fill=tk.X, pady=(0, 12))
        inner = ttk.Frame(outer, style="Card.TFrame", padding=padding)
        inner.pack(fill=tk.BOTH, expand=True, padx=1, pady=1)
        return inner

    def _build(self):
        self._stop_progress_animation()
        for widget in self.root.winfo_children():
            widget.destroy()
        self.root.title(f"FormCapy · {self.t('app')} · v{__version__}")
        self._build_menu()
        canvas = tk.Canvas(self.root, background=BG, highlightthickness=0)
        bar = ttk.Scrollbar(self.root, orient="vertical", command=canvas.yview)
        bar.pack(side=tk.RIGHT, fill=tk.Y)
        canvas.pack(fill=tk.BOTH, expand=True)
        canvas.configure(yscrollcommand=bar.set)
        main = ttk.Frame(canvas, padding=(28, 20, 28, 16))
        item = canvas.create_window((0, 0), window=main, anchor="nw")
        canvas.bind("<Configure>", lambda e: canvas.itemconfigure(item, width=e.width))
        main.bind("<Configure>", lambda _e: canvas.configure(scrollregion=canvas.bbox("all")))
        def wheel(event):
            if event.widget.winfo_toplevel() == self.root:
                amount = _wheel_units(event)
                if amount:
                    canvas.yview_scroll(amount, "units")
        for sequence in ("<MouseWheel>", "<Button-4>", "<Button-5>"):
            self.root.bind_all(sequence, wheel)

        top = ttk.Frame(main)
        top.pack(fill=tk.X, pady=(0, 16))
        if self._brand_image is not None:
            ttk.Label(top, image=self._brand_image).pack(side=tk.LEFT)
        ttk.Label(top, text="FormCapy", font=(self.font[0], 16, "bold")).pack(side=tk.LEFT, padx=(10, 8))
        ttk.Label(top, text=f"v{__version__}", foreground=MUTED, font=(self.font[0], 9)).pack(side=tk.LEFT)
        self._lang_var = tk.StringVar(self.root, value=self.lang)
        self._language_buttons = []
        for value, label in (("es", "Español"), ("en", "English"), ("zh", "中文")):
            button = ttk.Radiobutton(top, text=label, value=value, variable=self._lang_var,
                                     style="Lang.TRadiobutton", command=lambda: self._set_lang(self._lang_var.get()))
            button.pack(side=tk.RIGHT)
            self._language_buttons.append(button)
        ttk.Label(main, text=self.t("eyebrow"), foreground=ACCENT,
                  font=(self.font[0], 9, "bold")).pack(anchor=tk.W)
        ttk.Label(main, text=self.t("hero"), font=(self.font[0], 23, "bold")).pack(anchor=tk.W, pady=(4, 4))
        ttk.Label(main, text=self.t("tagline"), foreground=MUTED, wraplength=920).pack(anchor=tk.W, pady=(0, 18))
        columns = ttk.Frame(main)
        columns.pack(fill=tk.BOTH, expand=True)
        columns.columnconfigure(0, weight=3, minsize=440)
        columns.columnconfigure(1, weight=2, minsize=300)
        left = ttk.Frame(columns)
        left.grid(row=0, column=0, sticky="nsew", padx=(0, 20))
        right = ttk.Frame(columns)
        right.grid(row=0, column=1, sticky="nsew")
        formats = self._card(left)
        ttk.Label(formats, text=self.t("format"), style="Step.TLabel").pack(anchor=tk.W, pady=(0, 6))
        self._format_label = tk.StringVar(self.root, value=self.t("format_" + self.format_id.get()))
        self._format_combo = ttk.Combobox(formats, textvariable=self._format_label,
            values=[self.t("format_" + key) for key in ("mla", "apa", "custom")], state="readonly")
        self._format_combo.pack(fill=tk.X)
        self._format_combo.bind("<<ComboboxSelected>>", self._on_format_selected)
        ttk.Label(formats, text=self.t(self.format_id.get() + "_hint"), style="Muted.TLabel", wraplength=420).pack(anchor=tk.W, pady=(8, 0))
        self._profile_combo = self._edit_profile_btn = None
        if self.format_id.get() == "custom":
            ttk.Label(formats, text=self.t("template"), style="Card.TLabel").pack(anchor=tk.W, pady=(10, 4))
            choices = self._template_choices()
            selected_choice = next((label for label, builtin, saved in choices
                if (builtin and builtin == self.selected_builtin.get()) or
                   (not self.selected_builtin.get() and not builtin and saved == self.selected_profile.get())), choices[0][0])
            self._profile_choice = tk.StringVar(self.root, value=selected_choice)
            self._profile_combo = ttk.Combobox(formats, textvariable=self._profile_choice,
                values=[label for label, _builtin, _saved in choices], state="readonly")
            self._profile_combo.pack(fill=tk.X)
            self._profile_combo.bind("<<ComboboxSelected>>", self._on_profile_selected)
            if self.selected_builtin.get() or self._custom_profile().get("typography") == "chinese":
                ttk.Label(formats, text=self.t("builtin_hint"), style="Muted.TLabel", wraplength=420).pack(anchor=tk.W, pady=(8, 0))
            self._edit_profile_btn = ttk.Button(formats, text=self.t("edit_template"), style="Ghost.TButton", command=self._edit_template)
            self._edit_profile_btn.pack(anchor=tk.W, pady=(8, 0))
        info = self._card(left)
        ttk.Label(info, text=self.t("step1"), style="Step.TLabel").grid(row=0, column=0, columnspan=2, sticky=tk.W, pady=(0, 12))
        self._entries = {}
        for row, pair in ((1, ("name", "date")), (3, ("title",)), (5, ("instructor", "course"))):
            for col, key in enumerate(pair):
                span = 2 if key == "title" else 1
                pad = (0, 12) if col == 0 and span == 1 else (0, 0)
                ttk.Label(info, text=self.t(key), style="Card.TLabel").grid(row=row, column=col, columnspan=span, sticky="w", padx=pad)
                entry = ttk.Entry(info, textvariable=self.v[key], width=20)
                entry.grid(row=row + 1, column=col, columnspan=span, sticky="ew", padx=pad, pady=(4, 10))
                self._entries[key] = entry
        if self.format_id.get() in ("apa", "custom"):
            ttk.Label(info, text=self.t("affiliation"), style="Card.TLabel").grid(row=7, column=0, columnspan=2, sticky="w")
            entry = ttk.Entry(info, textvariable=self.affiliation)
            entry.grid(row=8, column=0, columnspan=2, sticky="ew", pady=(4, 10))
            self._entries["affiliation"] = entry
        for col in (0, 1):
            info.columnconfigure(col, weight=1)
        draft = self._card(left)
        row = ttk.Frame(draft, style="Card.TFrame")
        row.pack(fill=tk.X, pady=(0, 10))
        ttk.Label(row, text=self.t("step2"), style="Step.TLabel").pack(side=tk.LEFT)
        self._remove_btn = ttk.Button(row, text=self.t("remove"), style="Ghost.TButton", command=self._clear_file)
        self._drop = tk.Label(draft, bg=SOFT, fg=ACCENT, cursor="hand2", height=3,
                              highlightthickness=1, highlightbackground="#d8d8ef",
                              font=self.font, takefocus=True, wraplength=410)
        self._drop.pack(fill=tk.X)
        for event in ("<Button-1>", "<Return>", "<space>"):
            self._drop.bind(event, lambda _e: self._on_browse())
        if self.has_dnd:
            self._drop.drop_target_register(DND_FILES)
            self._drop.dnd_bind("<<Drop>>", self._on_drop)
        self._refresh_drop()
        ttk.Label(draft, text=self.t("limit"), style="Muted.TLabel").pack(anchor=tk.W, pady=(8, 0))
        self._btn = ttk.Button(left, text=self.t("convert"), style="Accent.TButton", command=self._on_convert)
        self._btn.pack(fill=tk.X, pady=(0, 10))
        actions = ttk.Frame(left)
        actions.pack(fill=tk.X)
        self._blank_btn = ttk.Button(actions, text=self.t("blank"), style="Ghost.TButton", command=self._on_blank)
        self._blank_btn.pack(side=tk.LEFT)
        self._advanced_btn = ttk.Button(actions, text=self.t("advanced"), style="Ghost.TButton", command=self._open_advanced)
        self._advanced_btn.pack(side=tk.RIGHT)
        self._progress_canvas = tk.Canvas(left, background=BG, height=40, highlightthickness=0)
        self._progress_canvas.bind("<Configure>", lambda _event: self._draw_processing_progress())
        self._progress_canvas.bind("<Destroy>", lambda _event: self._stop_progress_animation())
        self._status = ttk.Label(left, foreground=MUTED, wraplength=390, font=(self.font[0], 9))
        self._status.pack(fill=tk.X, pady=(12, 0))
        self._result_row = ttk.Frame(left)
        self._result_row.pack(fill=tk.X, pady=(8, 0))
        self._render_result()
        self._review_note = ttk.Label(left, foreground="#936221", wraplength=420, font=(self.font[0], 9))
        self._retained_note = ttk.Label(left, foreground="#24725b", wraplength=420, font=(self.font[0], 9))
        self._render_review()
        preview = self._card(right, padding=16)
        ttk.Label(preview, text=self.t("preview"), style="Step.TLabel").pack(anchor=tk.W)
        self._paper = tk.Canvas(preview, background=CARD, height=340, highlightthickness=0)
        self._paper.pack(fill=tk.X, pady=(12, 8))
        self._paper.bind("<Configure>", self._draw_preview)
        ttk.Label(preview, text=self.t("preview_hint"), style="Muted.TLabel", wraplength=300).pack(anchor=tk.W)
        specs = self._card(right, padding=16)
        ttk.Label(specs, text=self.t("spec_title"), style="Step.TLabel").pack(anchor=tk.W, pady=(0, 8))
        for line in self._summary_lines():
            ttk.Label(specs, text="✓  " + line, style="Muted.TLabel", wraplength=300).pack(anchor=tk.W, pady=3)
        ttk.Label(right, text=self.t("disclaimer"), foreground=MUTED, wraplength=330,
                  font=(self.font[0], 9)).pack(anchor=tk.W, pady=(0, 8))
        ttk.Label(main, text=self.t("local"), foreground=MUTED, font=(self.font[0], 9)).pack(anchor=tk.W, pady=(12, 0))
        self._render_status()
        self._draw_preview()
        self._set_busy(self._busy)

    def _template_choices(self):
        names = {"zh_course": "chinese_course", "zh_thesis": "chinese_thesis"}
        choices = [(f"{self.t('builtin_prefix')} · {self.t(names[preset['id']])}",
                    preset["id"], "") for preset in BUILTIN_CUSTOM_PRESETS]
        choices.extend((f"{self.t('saved_prefix')} · {profile['label']}", "", profile["label"])
                       for profile in self.custom_profiles)
        return choices

    def _custom_profile(self):
        for preset in BUILTIN_CUSTOM_PRESETS:
            if preset["id"] == self.selected_builtin.get():
                return normalize_custom_profile(dict(preset["profile"]))
        for profile in self.custom_profiles:
            if profile["label"] == self.selected_profile.get():
                return dict(profile)
        return dict(self.custom_profiles[0])

    def _default_date(self):
        if self.format_id.get() == "custom" and self._custom_profile().get("typography") == "chinese":
            today = datetime.date.today()
            return f"{today.year}年{today.month}月{today.day}日"
        return today_for_format(self.format_id.get())

    def _on_format_selected(self, _event=None):
        selected = next((key for key in ("mla", "apa", "custom")
                         if self.t("format_" + key) == self._format_label.get()), self.format_id.get())
        if self._busy or self._choosing_output:
            self._format_label.set(self.t("format_" + self.format_id.get()))
            return
        old_format = self.format_id.get()
        if selected == old_format:
            return
        automatic_date = self.v["date"].get() == self._default_date()
        self.format_id.set(selected)
        if automatic_date:
            self.v["date"].set(self._default_date())
        self._adv_win = None
        self._build()
        self._save_prefs()

    def _on_profile_selected(self, _event=None):
        if self._busy or self._choosing_output:
            return
        automatic_date = self.v["date"].get() == self._default_date()
        for label, builtin, saved in self._template_choices():
            if label == self._profile_choice.get():
                self.selected_builtin.set(builtin)
                if not builtin:
                    self.selected_profile.set(saved)
                break
        if automatic_date:
            self.v["date"].set(self._default_date())
        self._adv_win = None
        self._build()
        self._save_prefs()

    def _summary_lines(self):
        kind = self.format_id.get()
        if kind == "mla":
            return [self.t(key) for key in ("spec_font", "spec_spacing", "spec_header", "spec_cited")]
        if kind == "apa":
            return [self.t(key) for key in
                    ("spec_apa_font", "spec_apa_spacing", "spec_apa_titlepage", "spec_apa_references")]
        p = self._custom_profile()
        font = f"{p['font_name']} · {p['font_size']:g} pt · {p['paper_size']}"
        if p.get("typography") == "chinese":
            font += f" · {p['latin_font_name']}"
        page_numbers = self.t("header_" + p["header"])
        if p["header"] != "none":
            page_numbers += " · " + self.t("page_number_summary", position=self.t("page_number_" + p.get("page_number_position", "top_right")),
                                           start=p.get("page_number_start", 1))
            if p["title_page"] and not p.get("show_cover_page_number", True):
                page_numbers += " · " + self.t("cover_unnumbered")
        spacing = (self.t("spacing_exact_line", pt=p["line_spacing_pt"])
                   if p.get("line_spacing_pt", 0)
                   else self.t("spacing_multiple_line", spacing=p["line_spacing"]))
        indent = (self.t("indent_chars", n=p["first_line_chars"])
                  if p.get("first_line_chars", 0) else f"{p['first_line_cm']:g} cm")
        return [
            font,
            self.t("custom_spacing_margin", spacing=spacing, margin=p["margin_cm"]),
            self.t("custom_first_line", indent=indent, alignment=self.t("align_" + p["alignment"])),
            page_numbers + (self.t("title_page_suffix") if p["title_page"] else ""),
            self.t("reference_new_page", title=p["references_title"])
            + (self.t("reference_sorted") if p["sort_references"] else self.t("reference_original_gui")),
        ]

    def _draw_preview(self, *_):
        if not hasattr(self, "_paper") or not self._paper.winfo_exists():
            return
        c = self._paper
        c.delete("all")
        width = max(c.winfo_width(), 280)
        kind = self.format_id.get()
        profile = self._custom_profile()
        page_ratio = 210 / 297 if kind == "custom" and profile["paper_size"] == "A4" else 8.5 / 11
        page_width = min(width - 12, 329 * page_ratio)
        x0 = (width - page_width) / 2
        c.create_rectangle(x0 + 3, 4, x0 + page_width + 3, 333, fill="#eeede8", outline="")
        c.create_rectangle(x0, 0, x0 + page_width, 329, fill="#fffefa", outline=BORDER)
        margin = page_width * profile["margin_cm"] / (21 if profile["paper_size"] == "A4" else 21.59) if kind == "custom" else 32
        left, right = x0 + margin, x0 + page_width - margin
        name = self.v["name"].get().strip() or self.t("preview_name")
        last = extract_last_name(name, self.v["last_name"].get().strip() or None)
        header = profile["header"] if kind == "custom" else ("page" if kind == "apa" else "surname_page")
        font = profile["font_name"] if kind == "custom" else "Times New Roman"
        text_size = max(6, min(16, round(profile["font_size"] * 0.75))) if kind == "custom" else 9
        chinese = kind == "custom" and profile.get("typography") == "chinese"
        title_font = ("黑体", 12, "bold") if chinese else (font, text_size + 1, "bold")
        page_position = profile.get("page_number_position", "top_right") if kind == "custom" else "top_right"
        page_start = profile.get("page_number_start", 1) if kind == "custom" else 1
        cover_hidden = kind == "custom" and profile["title_page"] and not profile.get("show_cover_page_number", True)
        if header != "none" and not cover_hidden:
            x = width / 2 if page_position == "bottom_center" else right
            y = 308 if page_position.startswith("bottom_") else 24
            c.create_text(x, y, text=f"{last[:28]} {page_start}" if header == "surname_page" else str(page_start),
                          anchor="center" if page_position == "bottom_center" else "e", font=(font, text_size), fill=TEXT)
        title = self.v["title"].get().strip() or self.t("preview_title")
        if kind == "apa" or (kind == "custom" and profile["title_page"]):
            title_item = c.create_text(width / 2, 86, text=title[:100], width=right-left, anchor="n",
                                      font=title_font, fill=TEXT)
            top = max(128, c.bbox(title_item)[3] + 24)
            details = [name]
            if kind == "apa":
                details.append(self.affiliation.get().strip() or self.t("preview_affiliation"))
            elif self.affiliation.get().strip():
                details.append(self.affiliation.get().strip())
            details.extend(self.v[key].get().strip() for key in ("course", "instructor", "date"))
            for index, value in enumerate(value for value in details if value):
                c.create_text(width / 2, top + index * 21, text=value[:52], anchor="n", width=right-left,
                              font=(font, text_size), fill=TEXT)
            return
        details = [name]
        if kind == "custom" and self.affiliation.get().strip():
            details.append(self.affiliation.get().strip())
        if chinese:
            details.extend(self.v[key].get().strip() for key in ("instructor", "course", "date"))
            details = [value for value in details if value]
            title_item = c.create_text(width / 2, 52, text=title[:100], width=right-left,
                                      anchor="n", font=title_font, fill=TEXT)
            detail_top = c.bbox(title_item)[3] + 20
            for index, value in enumerate(details):
                c.create_text(width / 2, detail_top + index * 20, text=value[:52], anchor="n",
                              width=right-left, font=(font, text_size), fill=TEXT)
            body_top = detail_top + len(details) * 20 + 16
        else:
            details.extend((self.v["instructor"].get().strip() or self.t("preview_instructor_default"),
                            self.v["course"].get().strip() or self.t("preview_course_default"),
                            self.v["date"].get().strip() or today_mla()))
            for index, value in enumerate(details):
                c.create_text(left, 52 + index * 20, text=value[:52], anchor="w", width=right-left,
                              font=(font, text_size), fill=TEXT)
            title_top = 52 + len(details) * 20
            title_font = (font, text_size + 1, "bold") if kind == "custom" else (font, text_size + 1)
            title_item = c.create_text(width / 2, title_top, text=title[:100], width=right-left, anchor="n", font=title_font, fill=TEXT)
            body_top = max(title_top + 43, c.bbox(title_item)[3] + 18)
        line_step = (profile["line_spacing_pt"] * text_size / profile["font_size"]
                     if kind == "custom" and profile.get("line_spacing_pt", 0)
                     else profile["line_spacing"] * 8.5 if kind == "custom" else 17)
        line_step = max(1, line_step)
        body_bottom = 288 if header != "none" and page_position.startswith("bottom_") else 312
        for index in range(max(0, min(12, int((body_bottom - body_top) / line_step) + 1))):
            y = body_top + index * line_step
            if kind == "custom" and profile.get("first_line_chars", 0):
                indent_width = tkfont.Font(root=self.root, font=(font, text_size)).measure("汉") * profile["first_line_chars"]
            else:
                indent_width = page_width * profile["first_line_cm"] / (21 if profile["paper_size"] == "A4" else 21.59) if kind == "custom" else 12
            indent = indent_width if index in (0, 5) else 0
            end = right - (38 if index in (4, 7) else 0)
            c.create_line(left + indent, y, end, y, fill="#d6d3cb", width=2)

    def _build_menu(self):
        mb = tk.Menu(self.root)
        file = tk.Menu(mb, tearoff=0)
        file.add_command(label=self.t("menu_open"), accelerator="Cmd+O" if sys.platform == "darwin" else "Ctrl+O", command=self._on_browse)
        file.add_command(label=self.t("menu_blank"), command=self._on_blank)
        file.add_separator()
        file.add_command(label=self.t("menu_quit"), command=self._on_close)
        mb.add_cascade(label=self.t("menu_file"), menu=file)
        settings = tk.Menu(mb, tearoff=0)
        settings.add_command(label=self.t("advanced"), command=self._open_advanced)
        language = tk.Menu(settings, tearoff=0)
        for lang, label in (("es", "Español"), ("en", "English"), ("zh", "中文")):
            language.add_command(label=label, command=lambda lg=lang: self._set_lang(lg))
        settings.add_cascade(label=self.t("menu_language"), menu=language)
        mb.add_cascade(label=self.t("menu_settings"), menu=settings)
        help_menu = tk.Menu(mb, tearoff=0)
        help_menu.add_command(label=self.t("menu_howto"), command=lambda: messagebox.showinfo(self.t("menu_howto"), self.t("howto"), parent=self.root))
        help_menu.add_command(label=self.t("menu_notices"), command=self._show_third_party_notices)
        help_menu.add_command(label=self.t("menu_source_kit"), command=self._save_source_kit,
                              state=tk.NORMAL if self._bundled_file("MLAFMT-SOURCE-KIT.zip") else tk.DISABLED)
        help_menu.add_command(label=self.t("menu_about"), command=lambda: messagebox.showinfo(self.t("menu_about"), self.t("about", v=__version__), parent=self.root))
        mb.add_cascade(label=self.t("menu_help"), menu=help_menu)
        self.root.config(menu=mb)

    def _set_status(self, key, kind="muted", **kw):
        self._status_state = (key, kind, kw)
        self._render_status()

    def _render_status(self):
        key, kind, kw = self._status_state
        self._status.configure(text=self.t(key, **kw), foreground={"muted": MUTED, "ok": "#24725b", "err": "#b84040"}[kind])

    def _set_lang(self, lang):
        if self._busy or lang not in T or lang == self.lang:
            return
        was_open = self._adv_win is not None and self._adv_win.winfo_exists()
        self.lang = lang
        self._adv_win = None
        self._build()
        if was_open:
            self._open_advanced()
        self._save_prefs()

    def _save_prefs(self):
        if not self.remember:
            return True
        cfg = {"lang": self.lang, "txt_mode": self.txt_mode.get(), "format_id": self.format_id.get(),
               "custom_profiles": self.custom_profiles, "selected_profile": self.selected_profile.get(),
               "selected_builtin": self.selected_builtin.get()}
        cfg.update({key: self.v[key].get().strip() for key in ("name", "instructor", "course", "last_name")})
        cfg["affiliation"] = self.affiliation.get().strip()
        cfg.update({key: value.get() for key, value in self.opt.items()})
        return _save_config(cfg)

    def _preservation_section(self, parent):
        card = self._card(parent)
        ttk.Label(card, text=self.t("preserve_title"), style="Step.TLabel").pack(anchor=tk.W, pady=(0, 6))
        ttk.Label(card, text=self.t("preserve_hint"), style="Muted.TLabel", wraplength=510).pack(anchor=tk.W)
        count = ttk.Label(card, style="Muted.TLabel", wraplength=510)
        count.pack(anchor=tk.W, pady=(8, 6))
        button = ttk.Button(card, text=self.t("preserve_choose"), style="Ghost.TButton",
                            command=lambda: self._select_preserved_paragraphs(parent.winfo_toplevel()))
        button.pack(anchor=tk.W)
        ttk.Label(card, text=self.t("preserve_docx_only"), style="Muted.TLabel", wraplength=510).pack(anchor=tk.W, pady=(6, 0))
        ttk.Label(card, text=self.t("preserve_scope"), style="Muted.TLabel", wraplength=510).pack(anchor=tk.W, pady=(4, 0))
        self._preserve_count_labels.append((count, button))
        self._refresh_preservation_controls()
        return card

    def _refresh_preservation_controls(self):
        controls = []
        valid = bool(self._input_file and Path(self._input_file).suffix.lower() == ".docx")
        count = len(self._preserved_paragraphs) if self._preservation_source == self._input_file else 0
        for label, button in self._preserve_count_labels:
            if not label.winfo_exists() or not button.winfo_exists():
                continue
            label.configure(text=self.t("preserve_count", n=count) if count else self.t("preserve_none"))
            button.state(["!disabled"] if valid and not self._busy and not self._choosing_output else ["disabled"])
            controls.append((label, button))
        self._preserve_count_labels = controls

    def _select_preserved_paragraphs(self, parent=None):
        if self._busy or self._choosing_output:
            return
        parent = parent or self.root
        if self._preserve_win is not None and self._preserve_win.winfo_exists():
            self._preserve_win.lift()
            return
        source = self._input_file
        if not source or Path(source).suffix.lower() != ".docx":
            messagebox.showinfo(self.t("preserve_title"), self.t("preserve_docx_only"), parent=parent)
            return
        try:
            inspection = inspect_draft_paragraphs(source)
            fingerprint = inspection["source_fingerprint"]
            paragraphs = inspection["paragraphs"]
        except Exception as error:
            messagebox.showerror(self.t("preserve_title"), self.t("preserve_loading_error", e=str(error)), parent=parent)
            return
        if not paragraphs:
            messagebox.showinfo(self.t("preserve_title"), self.t("preserve_empty"), parent=parent)
            return
        same_draft = self._preservation_source == source and self._preservation_fingerprint == fingerprint
        available = {item["id"]: item for item in paragraphs}
        selected = set(self._preserved_paragraphs) & set(available) if same_draft else set()
        previous_grab = self.root.grab_current()
        w = tk.Toplevel(parent)
        self._preserve_win = w
        w.title(self.t("preserve_title"))
        w.configure(bg=BG)
        w.transient(parent)
        w.geometry("820x620")
        w.minsize(660, 450)
        w.grab_set()
        body = ttk.Frame(w, padding=18)
        body.pack(fill=tk.BOTH, expand=True)
        ttk.Label(body, text=self.t("preserve_hint"), foreground=MUTED, wraplength=760).pack(anchor=tk.W, pady=(0, 6))
        ttk.Label(body, text=self.t("preserve_toggle"), foreground=MUTED, wraplength=760).pack(anchor=tk.W, pady=(0, 10))
        if self._preserved_paragraphs and self._preservation_source == source and not same_draft:
            ttk.Label(body, text=self.t("draft_changed"), foreground="#936221", wraplength=760).pack(anchor=tk.W, pady=(0, 8))
        listing = ttk.Frame(body)
        listing.pack(fill=tk.BOTH, expand=True)
        tree = ttk.Treeview(listing, columns=("keep", "number", "text"), show="headings", selectmode="browse")
        tree.heading("keep", text=self.t("preserve_column"))
        tree.heading("number", text="#")
        tree.heading("text", text=self.t("preserve_text"))
        tree.column("keep", width=52, stretch=False, anchor="center")
        tree.column("number", width=60, stretch=False, anchor="center")
        tree.column("text", width=590, stretch=True)
        scrollbar = ttk.Scrollbar(listing, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=scrollbar.set)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        tree.pack(fill=tk.BOTH, expand=True)
        previews = {}
        for identifier, item in available.items():
            text = " ".join(item["text"].split())[:220]
            if item.get("has_images"):
                text += " · " + self.t("preserve_images")
            previews[identifier] = text
            tree.insert("", tk.END, iid=str(identifier), values=("☑" if identifier in selected else "☐", identifier, text))
        ranges = tk.StringVar(w, value=", ".join(str(identifier) for identifier in sorted(selected)))
        selection_count = tk.StringVar(w)

        def refresh_selection():
            ranges.set(", ".join(str(identifier) for identifier in sorted(selected)))
            selection_count.set(self.t("preserve_count", n=len(selected)))
            for identifier in available:
                tree.item(str(identifier), values=("☑" if identifier in selected else "☐", identifier, previews[identifier]))

        def toggle(identifier):
            if not identifier:
                return
            identifier = int(identifier)
            if identifier in selected:
                selected.remove(identifier)
            else:
                selected.add(identifier)
            refresh_selection()

        def click(event):
            identifier = tree.identify_row(event.y)
            if identifier:
                tree.focus(identifier)
                tree.selection_set(identifier)
                toggle(identifier)

        tree.bind("<Button-1>", click)
        tree.bind("<space>", lambda _event: toggle(tree.focus()))
        numbers = ttk.Frame(body)
        numbers.pack(fill=tk.X, pady=(12, 6))
        ttk.Label(numbers, text=self.t("preserve_ranges")).pack(anchor=tk.W)
        range_entry = ttk.Entry(numbers, textvariable=ranges)
        range_entry.pack(side=tk.LEFT, fill=tk.X, expand=True, pady=(4, 0))

        def apply_ranges():
            try:
                text = ranges.get().strip()
                if len(text) > 20000:
                    raise ValueError()
                numbers = set()
                for token in re.split(r"[,，;；\s]+", text) if text else []:
                    match = re.fullmatch(r"([0-9]+)(?:[-–—]([0-9]+))?", token)
                    if not match:
                        raise ValueError()
                    start = int(match.group(1))
                    end = int(match.group(2) or match.group(1))
                    if start < 1 or start > end or end > max(available):
                        raise ValueError()
                    found = {identifier for identifier in available if start <= identifier <= end}
                    if not found:
                        raise ValueError()
                    numbers.update(found)
                selected.clear()
                selected.update(numbers)
                refresh_selection()
                return True
            except (TypeError, ValueError):
                messagebox.showerror(self.t("preserve_title"), self.t("preserve_invalid_ranges"), parent=w)
                range_entry.focus_set()
                return False

        ttk.Button(numbers, text=self.t("preserve_apply_ranges"), style="Ghost.TButton", command=apply_ranges).pack(side=tk.RIGHT, padx=(8, 0), pady=(4, 0))
        actions = ttk.Frame(body)
        actions.pack(fill=tk.X, pady=(6, 10))

        def set_selection(all_paragraphs):
            selected.clear()
            if all_paragraphs:
                selected.update(available)
            refresh_selection()

        ttk.Button(actions, text=self.t("preserve_select_all"), style="Ghost.TButton", command=lambda: set_selection(True)).pack(side=tk.LEFT)
        ttk.Button(actions, text=self.t("preserve_clear"), style="Ghost.TButton", command=lambda: set_selection(False)).pack(side=tk.LEFT, padx=(8, 0))
        ttk.Label(actions, textvariable=selection_count, foreground=MUTED).pack(side=tk.RIGHT)

        def close():
            self._preserve_win = None
            w.destroy()
            if previous_grab is not None and previous_grab.winfo_exists():
                previous_grab.grab_set()
            self._refresh_preservation_controls()

        def use_selected():
            if self._input_file != source:
                messagebox.showwarning(self.t("preserve_title"), self.t("draft_changed"), parent=w)
                close()
                return
            if not apply_ranges():
                return
            self._preserved_paragraphs = set(selected)
            self._preservation_fingerprint = fingerprint
            self._preservation_source = source
            close()

        ttk.Label(body, text=self.t("preserve_scope"), foreground=MUTED, wraplength=760).pack(anchor=tk.W, pady=(0, 10))
        ttk.Button(body, text=self.t("preserve_apply"), style="Accent.TButton", command=use_selected).pack(fill=tk.X)
        refresh_selection()
        w.protocol("WM_DELETE_WINDOW", close)
        w.bind("<Escape>", lambda _event: close())

    def _edit_template(self):
        if self._busy or self._choosing_output:
            return
        if self._template_win is not None and self._template_win.winfo_exists():
            self._template_win.lift()
            return
        w = tk.Toplevel(self.root)
        self._template_win = w
        w.title(self.t("template_title"))
        w.configure(bg=BG)
        w.transient(self.root)
        w.geometry("680x760")
        w.minsize(580, 480)
        # The form is edited separately; generation always uses the saved copy.
        w.grab_set()
        canvas = tk.Canvas(w, bg=BG, highlightthickness=0)
        bar = ttk.Scrollbar(w, orient="vertical", command=canvas.yview)
        bar.pack(side=tk.RIGHT, fill=tk.Y)
        canvas.pack(fill=tk.BOTH, expand=True)
        canvas.configure(yscrollcommand=bar.set)
        body = ttk.Frame(canvas, padding=18)
        item = canvas.create_window(0, 0, window=body, anchor="nw")
        canvas.bind("<Configure>", lambda e: canvas.itemconfigure(item, width=e.width))
        body.bind("<Configure>", lambda _e: canvas.configure(scrollregion=canvas.bbox("all")))
        for sequence in ("<MouseWheel>", "<Button-4>", "<Button-5>"):
            w.bind(sequence, lambda e: canvas.yview_scroll(_wheel_units(e), "units"))
        ttk.Label(body, text=self.t("template_hint"), foreground=MUTED, wraplength=600).pack(anchor=tk.W, pady=(0, 12))
        card = self._card(body)
        for column in (0, 1):
            card.columnconfigure(column, weight=1)
        current = {"page_number_position": "top_right", "show_cover_page_number": True,
                   "page_number_start": 1, "typography": "western", "latin_font_name": "Times New Roman",
                   "line_spacing_pt": 0, "first_line_chars": 0, **self._custom_profile()}
        if self.selected_builtin.get():
            # Built-ins remain immutable and never overwrite a saved namesake.
            base_label = current["label"]
            taken = {profile["label"] for profile in self.custom_profiles}
            number = 1
            while True:
                suffix = f" {self.t('template_copy')}" + (f" {number}" if number > 1 else "")
                current["label"] = base_label[:60 - len(suffix)] + suffix
                if current["label"] not in taken:
                    break
                number += 1
        variables = {key: (tk.BooleanVar(w, value=value) if isinstance(value, bool)
                           else tk.StringVar(w, value=f"{value:g}" if isinstance(value, (int, float)) else value))
                     for key, value in current.items()}
        choice_maps = {
            "alignment": {self.t("align_" + key): key for key in ("left", "justify")},
            "header": {self.t("header_" + key): key for key in ("surname_page", "page", "none")},
            "page_number_position": {self.t("page_number_" + key): key
                                     for key in ("top_right", "bottom_center", "bottom_right")},
            "typography": {self.t("typography_" + key): key for key in ("western", "chinese")},
        }
        display_vars = {key: tk.StringVar(w, value=next(label for label, value in choices.items() if value == current[key]))
                        for key, choices in choice_maps.items()}
        fields = (("label", "profile_label"), ("paper_size", "paper_size"), ("font_name", "font_name"),
                  ("font_size", "font_size"), ("line_spacing", "line_spacing"), ("margin_cm", "margin_cm"),
                  ("first_line_cm", "first_line_cm"), ("alignment", "alignment"), ("header", "header"),
                  ("references_title", "references_title"))
        ranges = {"font_size": (8, 24, 0.5), "line_spacing": (1, 3, 0.25),
                  "margin_cm": (0.5, 5.08, 0.1), "first_line_cm": (0, 2.54, 0.1),
                  "line_spacing_pt": (0, 60, 1), "first_line_chars": (0, 4, 0.5)}
        main_controls = {}
        for row, (key, label) in enumerate(fields):
            ttk.Label(card, text=self.t(label), style="Card.TLabel").grid(row=row, column=0, sticky="w", padx=(0, 16), pady=6)
            if key in choice_maps:
                entry = ttk.Combobox(card, textvariable=display_vars[key], values=list(choice_maps[key]), state="readonly")
            elif key in ("paper_size", "font_name"):
                entry = ttk.Combobox(card, textvariable=variables[key], values=("Letter", "A4") if key == "paper_size" else FONT_CHOICES, state="readonly")
            elif key in ranges:
                minimum, maximum, increment = ranges[key]
                entry = ttk.Spinbox(card, textvariable=variables[key], from_=minimum, to=maximum, increment=increment, width=18)
            else:
                entry = ttk.Entry(card, textvariable=variables[key], width=26)
            entry.grid(row=row, column=1, sticky="ew", pady=6)
            main_controls[key] = entry
        for row, key in enumerate(("title_page", "sort_references"), start=len(fields)):
            ttk.Checkbutton(card, text=self.t(key), variable=variables[key], style="Card.TCheckbutton").grid(row=row, column=0, columnspan=2, sticky="w", pady=5)

        def read_profile():
            values = {key: variable.get() for key, variable in variables.items()}
            for key, choices in choice_maps.items():
                values[key] = choices[display_vars[key].get()]
            for key in ranges:
                values[key] = float(values[key])
            values["page_number_start"] = int(values["page_number_start"])
            return normalize_custom_profile(values)

        def invalid(error):
            messagebox.showerror(self.t("template_title"), self.t("template_invalid", e=str(error)), parent=w)

        def close():
            self._template_win = None
            w.destroy()

        advanced_window = None

        def sync_main_controls(*_):
            for alternate, primary in (("line_spacing_pt", "line_spacing"), ("first_line_chars", "first_line_cm")):
                try:
                    active = float(variables[alternate].get()) > 0
                except (TypeError, ValueError, tk.TclError):
                    active = False
                main_controls[primary].state(["disabled"] if active else ["!disabled"])

        for key in ("line_spacing_pt", "first_line_chars"):
            variables[key].trace_add("write", sync_main_controls)
        sync_main_controls()

        def advanced_custom():
            nonlocal advanced_window
            if advanced_window is not None and advanced_window.winfo_exists():
                advanced_window.lift()
                return
            advanced = tk.Toplevel(w)
            advanced_window = advanced
            advanced.title(self.t("advanced_custom_title"))
            advanced.configure(bg=BG)
            advanced.transient(w)
            advanced.geometry("620x660")
            advanced.minsize(540, 420)
            advanced.grab_set()
            original_pages = {key: variables[key].get() for key in
                              ("show_cover_page_number", "page_number_start")}
            original_position = display_vars["page_number_position"].get()
            original_text = {key: variables[key].get() for key in
                             ("latin_font_name", "line_spacing_pt", "first_line_chars")}
            original_typography = display_vars["typography"].get()
            advanced_canvas = tk.Canvas(advanced, bg=BG, highlightthickness=0)
            advanced_bar = ttk.Scrollbar(advanced, orient="vertical", command=advanced_canvas.yview)
            advanced_bar.pack(side=tk.RIGHT, fill=tk.Y)
            advanced_canvas.pack(fill=tk.BOTH, expand=True)
            advanced_canvas.configure(yscrollcommand=advanced_bar.set)
            content = ttk.Frame(advanced_canvas, padding=18)
            advanced_item = advanced_canvas.create_window(0, 0, window=content, anchor="nw")
            advanced_canvas.bind("<Configure>", lambda event: advanced_canvas.itemconfigure(advanced_item, width=event.width))
            content.bind("<Configure>", lambda _event: advanced_canvas.configure(scrollregion=advanced_canvas.bbox("all")))
            for sequence in ("<MouseWheel>", "<Button-4>", "<Button-5>"):
                advanced.bind(sequence, lambda event: advanced_canvas.yview_scroll(_wheel_units(event), "units"))
            ttk.Label(content, text=self.t("advanced_custom_hint"), foreground=MUTED,
                      wraplength=550).pack(anchor=tk.W, pady=(0, 12))
            text_options = self._card(content)
            text_options.columnconfigure(1, weight=1)
            ttk.Label(text_options, text=self.t("typography"), style="Card.TLabel").grid(row=0, column=0, sticky="w", padx=(0, 16), pady=6)
            typography = ttk.Combobox(text_options, textvariable=display_vars["typography"],
                values=list(choice_maps["typography"]), state="readonly")
            typography.grid(row=0, column=1, sticky="ew", pady=6)
            ttk.Label(text_options, text=self.t("latin_font_name"), style="Card.TLabel").grid(row=1, column=0, sticky="w", padx=(0, 16), pady=6)
            latin_font = ttk.Combobox(text_options, textvariable=variables["latin_font_name"],
                values=("Times New Roman", "Arial", "Calibri"), state="readonly")
            latin_font.grid(row=1, column=1, sticky="ew", pady=6)
            for row, key in enumerate(("line_spacing_pt", "first_line_chars"), start=2):
                ttk.Label(text_options, text=self.t(key), style="Card.TLabel", wraplength=320).grid(row=row, column=0, sticky="w", padx=(0, 16), pady=6)
                minimum, maximum, increment = ranges[key]
                ttk.Spinbox(text_options, textvariable=variables[key], from_=minimum, to=maximum,
                    increment=increment).grid(row=row, column=1, sticky="ew", pady=6)
            ttk.Label(content, text=self.t("chinese_controls_hint"), foreground=MUTED,
                      wraplength=550).pack(anchor=tk.W, pady=(0, 14))
            options = self._card(content)
            options.columnconfigure(1, weight=1)
            ttk.Label(options, text=self.t("page_number_position"), style="Card.TLabel").grid(row=0, column=0, sticky="w", padx=(0, 16), pady=6)
            position = ttk.Combobox(options, textvariable=display_vars["page_number_position"],
                                    values=list(choice_maps["page_number_position"]), state="readonly")
            position.grid(row=0, column=1, sticky="ew", pady=6)
            ttk.Label(options, text=self.t("page_number_start"), style="Card.TLabel").grid(row=1, column=0, sticky="w", padx=(0, 16), pady=6)
            start = ttk.Spinbox(options, textvariable=variables["page_number_start"], from_=1, to=9999, increment=1)
            start.grid(row=1, column=1, sticky="ew", pady=6)
            cover = ttk.Checkbutton(options, text=self.t("show_cover_page_number"),
                                    variable=variables["show_cover_page_number"], style="Card.TCheckbutton")
            cover.grid(row=2, column=0, columnspan=2, sticky="w", pady=6)
            ttk.Label(content, text=self.t("page_number_hint"), foreground=MUTED,
                      wraplength=550).pack(anchor=tk.W, pady=(0, 14))
            self._preservation_section(content)

            def update_enabled(*_):
                enabled = choice_maps["header"].get(display_vars["header"].get()) != "none"
                position.state(["!disabled"] if enabled else ["disabled"])
                start.state(["!disabled"] if enabled else ["disabled"])
                cover.state(["!disabled"] if enabled and variables["title_page"].get() else ["disabled"])
                chinese = choice_maps["typography"].get(display_vars["typography"].get()) == "chinese"
                latin_font.state(["!disabled"] if chinese else ["disabled"])

            traces = [(display_vars["header"], display_vars["header"].trace_add("write", update_enabled)),
                      (variables["title_page"], variables["title_page"].trace_add("write", update_enabled)),
                      (display_vars["typography"], display_vars["typography"].trace_add("write", update_enabled))]
            update_enabled()

            def close_advanced(save=True):
                nonlocal advanced_window
                if save:
                    try:
                        first_page = int(variables["page_number_start"].get())
                        if not 1 <= first_page <= 9999:
                            raise ValueError(self.t("page_number_invalid"))
                        read_profile()
                    except (TypeError, ValueError, tk.TclError) as error:
                        messagebox.showerror(self.t("advanced_custom_title"), self.t("template_invalid", e=str(error)), parent=advanced)
                        start.focus_set()
                        return
                else:
                    for key, value in original_pages.items():
                        variables[key].set(value)
                    display_vars["page_number_position"].set(original_position)
                    for key, value in original_text.items():
                        variables[key].set(value)
                    display_vars["typography"].set(original_typography)
                for variable, trace in traces:
                    variable.trace_remove("write", trace)
                advanced_window = None
                advanced.destroy()
                if w.winfo_exists():
                    w.grab_set()
                    w.lift()

            ttk.Button(content, text=self.t("close"), style="Accent.TButton", command=close_advanced).pack(fill=tk.X, side=tk.BOTTOM)
            advanced.protocol("WM_DELETE_WINDOW", lambda: close_advanced(False))
            advanced.bind("<Escape>", lambda _e: close_advanced(False))

        def use_profile(profile, status_key="template_saved"):
            automatic_date = self.v["date"].get() == self._default_date()
            profiles = [p for p in self.custom_profiles if p["label"] != profile["label"]]
            if len(profiles) >= 20:
                messagebox.showerror(self.t("template_title"), self.t("template_limit"), parent=w)
                return
            profiles.append(profile)
            self.custom_profiles = profiles
            self.selected_builtin.set("")
            self.selected_profile.set(profile["label"])
            self.format_id.set("custom")
            if automatic_date:
                self.v["date"].set(self._default_date())
            saved = self._save_prefs()
            close()
            self._adv_win = None
            self._build()
            self._set_status(status_key if saved else "template_session_only", "ok" if saved else "muted", label=profile["label"])

        def save():
            try:
                profile = read_profile()
            except (TypeError, ValueError, tk.TclError) as error:
                invalid(error)
                return
            use_profile(profile)

        def import_profile():
            path = filedialog.askopenfilename(parent=w, title=self.t("import_template"),
                filetypes=[(self.t("template_json"), "*.json")])
            if not path:
                return
            try:
                with open(path, "rb") as source:
                    raw = source.read(64 * 1024 + 1)
                if len(raw) > 64 * 1024:
                    raise ValueError(self.t("template_too_large"))
                data = json.loads(raw.decode("utf-8-sig"))
                if isinstance(data, dict) and "custom_profile" in data:
                    data = data["custom_profile"]
                if not isinstance(data, dict):
                    raise ValueError("Expected a template object")
                profile = normalize_custom_profile(data)
            except (OSError, UnicodeError, TypeError, ValueError) as error:
                invalid(error)
                return
            use_profile(profile, "template_imported")

        def export_profile():
            try:
                profile = read_profile()
            except (TypeError, ValueError, tk.TclError) as error:
                invalid(error)
                return
            safe_name = "".join("_" if char in '\\/:*?"<>|\x00' or ord(char) < 32 else char for char in profile["label"]).strip(" .") or "Course template"
            path = filedialog.asksaveasfilename(parent=w, title=self.t("export_template"), defaultextension=".json",
                initialfile=safe_name + ".json", filetypes=[(self.t("template_json"), "*.json")])
            if not path:
                return
            temporary = None
            try:
                with tempfile.NamedTemporaryFile(mode="w", encoding="utf-8", dir=os.path.dirname(os.path.abspath(path)),
                                                 suffix=".tmp", delete=False) as target:
                    temporary = target.name
                    json.dump(profile, target, ensure_ascii=False, indent=2)
                    target.write("\n")
                os.replace(temporary, path)
            except OSError as error:
                messagebox.showerror(self.t("template_title"), self.t("template_io", e=str(error)), parent=w)
            finally:
                if temporary and os.path.exists(temporary):
                    try:
                        os.unlink(temporary)
                    except OSError:
                        pass

        def delete_profile():
            if self.selected_builtin.get():
                return
            self.custom_profiles = [p for p in self.custom_profiles if p["label"] != self.selected_profile.get()]
            if not self.custom_profiles:
                self.custom_profiles = [dict(DEFAULT_CUSTOM_PROFILE)]
            self.selected_profile.set(self.custom_profiles[0]["label"])
            self._save_prefs()
            close()
            self._build()

        row = ttk.Frame(body)
        row.pack(fill=tk.X, pady=(0, 10))
        ttk.Button(body, text=self.t("advanced_custom"), style="Ghost.TButton", command=advanced_custom).pack(anchor=tk.W, pady=(0, 12), before=row)
        ttk.Button(row, text=self.t("import_template"), style="Ghost.TButton", command=import_profile).pack(side=tk.LEFT)
        ttk.Button(row, text=self.t("export_template"), style="Ghost.TButton", command=export_profile).pack(side=tk.LEFT, padx=8)
        delete_button = ttk.Button(body, text=self.t("delete_template"), style="Ghost.TButton", command=delete_profile)
        delete_button.pack(anchor=tk.W, pady=(0, 12))
        if self.selected_builtin.get():
            delete_button.state(["disabled"])
            ttk.Label(body, text=self.t("builtin_readonly"), foreground=MUTED, wraplength=600).pack(anchor=tk.W, pady=(0, 12))
        ttk.Button(body, text=self.t("save_template"), style="Accent.TButton", command=save).pack(fill=tk.X)
        w.protocol("WM_DELETE_WINDOW", close)
        w.bind("<Escape>", lambda _e: close())

    def _open_advanced(self):
        if self._busy or self._choosing_output:
            return
        if self._adv_win is not None and self._adv_win.winfo_exists():
            self._adv_win.lift()
            return
        w = tk.Toplevel(self.root)
        self._adv_win = w
        w.title(self.t("adv_title"))
        w.configure(bg=BG)
        w.transient(self.root)
        w.geometry("620x690")
        w.minsize(560, 420)
        canvas = tk.Canvas(w, bg=BG, highlightthickness=0)
        bar = ttk.Scrollbar(w, orient="vertical", command=canvas.yview)
        bar.pack(side=tk.RIGHT, fill=tk.Y)
        canvas.pack(fill=tk.BOTH, expand=True)
        canvas.configure(yscrollcommand=bar.set)
        body = ttk.Frame(canvas, padding=18)
        item = canvas.create_window(0, 0, window=body, anchor="nw")
        canvas.bind("<Configure>", lambda e: canvas.itemconfigure(item, width=e.width))
        body.bind("<Configure>", lambda _e: canvas.configure(scrollregion=canvas.bbox("all")))
        for sequence in ("<MouseWheel>", "<Button-4>", "<Button-5>"):
            w.bind(sequence, lambda e: canvas.yview_scroll(_wheel_units(e), "units"))
        def section(key):
            card = self._card(body)
            ttk.Label(card, text=self.t(key), style="Step.TLabel").pack(anchor=tk.W, pady=(0, 6))
            return card
        card = section("adv_header")
        ttk.Label(card, text=self.t("lastname"), style="Card.TLabel").pack(anchor=tk.W)
        lastname_entry = ttk.Entry(card, textvariable=self.v["last_name"])
        lastname_entry.pack(fill=tk.X, pady=4)
        if self.format_id.get() == "apa" or (self.format_id.get() == "custom" and self._custom_profile()["header"] != "surname_page"):
            lastname_entry.state(["disabled"])
        ttk.Label(card, text=self.t("lastname_hint"), style="Muted.TLabel", wraplength=510).pack(anchor=tk.W)
        card = section("adv_txt")
        for key, mode in (("mode_blank", "blank_line"), ("mode_line", "line_by_line")):
            ttk.Radiobutton(card, text=self.t(key), variable=self.txt_mode, value=mode, style="Card.TRadiobutton").pack(anchor=tk.W)
        ttk.Checkbutton(card, text=self.t("md"), variable=self.opt["md"], style="Card.TCheckbutton").pack(anchor=tk.W)
        card = section("adv_body")
        for key in ("strip", "headings", "bq"):
            ttk.Checkbutton(card, text=self.t(key), variable=self.opt[key], style="Card.TCheckbutton").pack(anchor=tk.W)
        self._preservation_section(body)
        card = section("adv_after")
        ttk.Checkbutton(card, text=self.t("open_after"), variable=self.opt["open_after"], style="Card.TCheckbutton").pack(anchor=tk.W)
        card = section("adv_preview")
        summary = tk.StringVar(w)
        def refresh_summary(*_):
            lines = self._summary_lines()
            if self.format_id.get() == "mla" or (self.format_id.get() == "custom" and self._custom_profile()["header"] == "surname_page"):
                lines = [self.t("preview_header", name=extract_last_name(self.v["name"].get(), self.v["last_name"].get() or None))] + lines
            summary.set("\n".join(lines))
        refresh_summary()
        lastname_entry.bind("<KeyRelease>", refresh_summary)
        ttk.Label(card, textvariable=summary, style="Muted.TLabel", wraplength=510).pack(anchor=tk.W)
        ttk.Label(card, text=self.t("citation_hint"), style="Muted.TLabel", wraplength=510).pack(anchor=tk.W, pady=(8, 0))
        def close():
            self._save_prefs()
            self._adv_win = None
            w.destroy()
        ttk.Button(body, text=self.t("close"), style="Accent.TButton", command=close).pack(fill=tk.X)
        w.protocol("WM_DELETE_WINDOW", close)
        w.bind("<Escape>", lambda _e: close())

    def _refresh_drop(self):
        if self._input_file:
            try:
                size = os.path.getsize(self._input_file) / 1024
            except OSError:
                size = 0
            self._drop.config(text=os.path.basename(self._input_file) + "\n" + self.t("file_size", size=f"{size:.1f}"), fg=TEXT)
            self._remove_btn.pack(side=tk.RIGHT)
        else:
            self._drop.config(text=self.t("drop" if self.has_dnd else "drop_nodnd"), fg=ACCENT)
            self._remove_btn.pack_forget()

    def _clear_file(self):
        if self._busy:
            return
        self._input_file = None
        self._preserved_paragraphs.clear()
        self._preservation_fingerprint = None
        self._preservation_source = None
        self._refresh_preservation_controls()
        self._refresh_drop()
        self._set_status("status_ready")

    def _set_file(self, path):
        if self._busy:
            return False
        if not os.path.isfile(path):
            self._set_status("not_found", "err")
            return False
        if Path(path).suffix.lower() not in SUPPORTED_EXTS:
            self._set_status("unsupported", "err")
            return False
        try:
            size = os.path.getsize(path)
        except OSError as error:
            self._report_error(error)
            return False
        if size > MAX_FILE_SIZE:
            self._set_status("too_large", "err")
            return False
        if size == 0:
            self._set_status("empty_file", "err")
            return False
        self._input_file = os.path.abspath(path)
        self._preserved_paragraphs.clear()
        self._preservation_fingerprint = None
        self._preservation_source = None
        self._refresh_preservation_controls()
        self._refresh_drop()
        self._set_status("loaded", f=os.path.basename(path))
        return True

    def _restore_window(self):
        if self._choosing_output or self._closing:
            return
        if self.root.grab_current() is not None:
            return
        self.root.deiconify()
        self.root.lift()
        self.root.focus_force()

    def _open_documents(self, *paths):
        """Accept native file-open events without changing an active conversion."""
        if not paths:
            return
        if self._busy or self._choosing_output:
            self._pending_file = paths[0]
            self._set_status("queued_file", f=os.path.basename(paths[0]))
            return
        self._restore_window()
        if self._set_file(paths[0]):
            if len(paths) > 1:
                self._set_status("one_draft", f=os.path.basename(paths[0]))
            if self.root.grab_current() is None:
                self._entries["name" if not self.v["name"].get().strip() else "title"].focus_set()

    def _load_pending_file(self, output=None):
        pending, self._pending_file = self._pending_file, None
        if not self._closing and pending and self._set_file(pending) and output:
            self._set_status("saved_next", f=os.path.basename(output), next=os.path.basename(pending))

    def _on_drop(self, event):
        paths = self.root.tk.splitlist(event.data)
        if paths:
            self._set_file(paths[0])

    def _on_browse(self):
        if self._busy or self._choosing_output:
            return
        path = filedialog.askopenfilename(parent=self.root, title=self.t("choose"),
            filetypes=[(self.t("draft_types"), "*.docx *.txt *.md"), (self.t("word_type"), "*.docx"), (self.t("text_type"), "*.txt *.md")])
        if path:
            self._set_file(path)

    def _meta(self):
        return {k: self.v[k].get().strip() for k in ("name", "title", "instructor", "course", "date")} | {
            "last_name": self.v["last_name"].get().strip() or None,
            "format_id": self.format_id.get(), "custom_profile": self._custom_profile(), "affiliation": self.affiliation.get().strip(),
            "preserve_paragraphs": sorted(self._preserved_paragraphs) if self._preservation_source == self._input_file else [],
            "source_fingerprint": self._preservation_fingerprint if self._preservation_source == self._input_file else None,
        }

    def _check(self, meta):
        for key, error in (("name", "need_name"), ("title", "need_title")):
            if not meta[key]:
                self._set_status(error, "err")
                self._entries[key].focus_set()
                return False
        if date_warning(meta["date"], format_id=meta["format_id"]):
            return messagebox.askyesno(self.t("date_tip_title"), self.t("date_tip", d=today_mla()), parent=self.root)
        return True

    def _ask_save(self, meta, title_key):
        return filedialog.asksaveasfilename(parent=self.root, defaultextension=".docx", filetypes=[(self.t("word_type"), "*.docx")],
            initialfile=generate_filename(meta["name"], meta["last_name"], format_id=meta["format_id"], custom_profile=meta["custom_profile"]), initialdir=os.path.dirname(self._input_file) if self._input_file else None, title=self.t(title_key))

    def _open_file(self, path):
        try:
            if sys.platform.startswith("win"):
                os.startfile(path)
            else:
                subprocess.Popen(["open" if sys.platform == "darwin" else "xdg-open", path])
        except OSError as e:
            self._set_status("error", "err", e=str(e))

    def _reveal_result(self):
        if not self._last_output:
            return
        try:
            if sys.platform == "darwin":
                subprocess.Popen(["open", "-R", self._last_output])
            else:
                self._open_file(os.path.dirname(self._last_output))
        except OSError as error:
            self._set_status("error", "err", e=str(error))

    def _render_result(self):
        for w in self._result_row.winfo_children():
            w.destroy()
        if self._last_output:
            ttk.Button(self._result_row, text=self.t("open_result"), style="Ghost.TButton", command=lambda: self._open_file(self._last_output)).pack(side=tk.LEFT)
            ttk.Button(self._result_row, text=self.t("reveal_result" if sys.platform == "darwin" else "view_folder"), style="Ghost.TButton", command=self._reveal_result).pack(side=tk.LEFT, padx=8)

    @staticmethod
    def _bundled_file(name):
        resource_root = Path(getattr(sys, "_MEIPASS", Path(__file__).parent))
        candidates = [resource_root / name, Path(sys.executable).parent / name,
                      Path(__file__).parent / "build" / "windows" / name]
        return next((path for path in candidates if path.is_file()), None)

    def _save_source_kit(self):
        source = self._bundled_file("MLAFMT-SOURCE-KIT.zip")
        if source is None:
            return
        destination = filedialog.asksaveasfilename(parent=self.root, title=self.t("menu_source_kit"),
                                                   initialfile=source.name, defaultextension=".zip",
                                                   filetypes=[("ZIP", "*.zip")])
        if not destination:
            return
        try:
            if Path(destination).resolve() != source.resolve():
                shutil.copyfile(source, destination)
            messagebox.showinfo(self.t("menu_source_kit"), self.t("source_kit_saved", path=destination), parent=self.root)
        except OSError as error:
            messagebox.showerror(self.t("menu_source_kit"), str(error), parent=self.root)

    def _show_third_party_notices(self):
        try:
            filename = self._bundled_file("THIRD-PARTY-NOTICES.txt")
            if filename is None:
                messagebox.showinfo(self.t("menu_notices"), self.t("notices_missing"), parent=self.root)
                return
            notices = filename.read_text(encoding="utf-8", errors="replace")
        except OSError as error:
            messagebox.showerror(self.t("menu_notices"), str(error), parent=self.root)
            return
        window = tk.Toplevel(self.root)
        window.title(self.t("menu_notices"))
        window.geometry("820x620")
        container = ttk.Frame(window, padding=12)
        container.pack(fill=tk.BOTH, expand=True)
        scrollbar = ttk.Scrollbar(container)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        content = tk.Text(container, wrap=tk.WORD, yscrollcommand=scrollbar.set)
        content.pack(fill=tk.BOTH, expand=True)
        scrollbar.configure(command=content.yview)
        content.insert("1.0", notices)
        content.configure(state=tk.DISABLED)

    def _render_review(self):
        omitted, headings = self._review_counts
        items = []
        if omitted:
            items.append(self.t("review_omitted", n=omitted))
        if headings:
            items.append(self.t("review_headings", n=headings))
        if self._reference_sort_skipped:
            items.append(self.t("review_reference_order"))
        if items:
            self._review_note.configure(text=self.t("review_result", items=self.t("list_separator").join(items)))
            self._review_note.pack(fill=tk.X, pady=(8, 0))
        else:
            self._review_note.pack_forget()
        retained = [self.t(key, n=count) for key, count in zip(
            ("retained_paragraphs", "retained_images", "retained_tables"), self._retained_counts) if count]
        if retained:
            self._retained_note.configure(text=self.t("retained_result", items=self.t("list_separator").join(retained)))
            self._retained_note.pack(fill=tk.X, pady=(8, 0))
        else:
            self._retained_note.pack_forget()

    def _report_error(self, e):
        if isinstance(e, ValueError) and str(e) in ("draft_changed", "too_large", "empty_file", "unsupported_document"):
            key, kw = str(e), {}
        elif isinstance(e, PermissionError):
            key, kw = "perm", {}
        elif isinstance(e, FileNotFoundError):
            key, kw = "missing_source", {}
        elif "package not found" in str(e).lower() or "zip" in str(e).lower() or "badzip" in type(e).__name__.lower():
            key, kw = "bad_doc", {}
        else:
            key, kw = "error", {"e": str(e)}
        self._set_status(key, "err", **kw)
        messagebox.showerror(self.t("app"), self.t(key, **kw), parent=self.root)

    def _stop_progress_animation(self):
        pending = getattr(self, "_progress_after", None)
        self._progress_after = None
        if pending is not None:
            try:
                self.root.after_cancel(pending)
            except tk.TclError:
                pass

    def _draw_processing_progress(self):
        canvas = getattr(self, "_progress_canvas", None)
        if canvas is None:
            return
        try:
            if not canvas.winfo_exists():
                return
            width = max(48, canvas.winfo_width())
        except tk.TclError:
            return
        tick = self._progress_tick % 64
        forward = tick < 32
        position = tick / 32 if forward else (64 - tick) / 32
        origin = round(max(0, width - 48) * position)
        canvas.delete("processing")
        canvas.create_line(0, 34, width, 34, fill="#4d8cda", width=2, tags="processing")
        # A repeating position and four leg poses indicate activity, not completion.
        legs = (
            ((5, 11, 2, 4), (4, 15, 3, 1), (17, 11, 2, 4), (17, 15, 3, 1)),
            ((6, 11, 2, 3), (7, 14, 3, 1), (16, 11, 2, 3), (15, 14, 3, 1)),
            ((5, 11, 2, 4), (5, 15, 3, 1), (17, 11, 2, 4), (16, 15, 3, 1)),
            ((4, 11, 2, 3), (3, 14, 3, 1), (18, 11, 2, 3), (19, 14, 3, 1)),
        )[(tick // 2) % 4]
        groups = (
            ("#754d33", legs),
            ("#9b6846", ((3, 5, 14, 7), (2, 7, 2, 4), (5, 4, 9, 2), (14, 4, 6, 7), (14, 2, 3, 3))),
            ("#ba8b5c", ((17, 7, 6, 4), (17, 10, 5, 2), (5, 5, 8, 1))),
            ("#754d33", ((4, 10, 13, 2), (15, 3, 1, 1))),
            ("#302820", ((18, 5, 1, 1), (22, 7, 1, 2))),
        )
        for color, rectangles in groups:
            for x, y, rect_width, rect_height in rectangles:
                left = x if forward else 24 - x - rect_width
                canvas.create_rectangle(origin + left * 2, 2 + y * 2,
                                        origin + (left + rect_width) * 2, 2 + (y + rect_height) * 2,
                                        fill=color, outline="", tags="processing")

    def _animate_progress(self):
        self._progress_after = None
        if self._closing or not self._busy or not self._progress_canvas.winfo_exists():
            return
        self._draw_processing_progress()
        self._progress_tick = (self._progress_tick + 1) % 64
        self._progress_after = self.root.after(70, self._animate_progress)

    def _set_busy(self, busy):
        was_busy = self._busy
        self._busy = busy
        if busy:
            if not was_busy:
                self._progress_tick = 0
            self._progress_canvas.pack(fill=tk.X, pady=(10, 0), before=self._status)
            if self._progress_after is None:
                self._animate_progress()
        else:
            self._stop_progress_animation()
            self._progress_canvas.pack_forget()
        self._refresh_preservation_controls()
        for button in (self._btn, self._blank_btn, self._advanced_btn, self._remove_btn, self._format_combo,
                       self._profile_combo, self._edit_profile_btn, *self._language_buttons):
            if button is None:
                continue
            button.state(["disabled"] if busy else ["!disabled"])

    def _on_blank(self):
        self._generate(True)

    def _on_convert(self):
        self._generate(False)

    def _generate(self, blank):
        if self._busy or self._choosing_output:
            return
        if self._preserve_win is not None and self._preserve_win.winfo_exists():
            self._preserve_win.lift()
            return
        if not blank and not self._input_file:
            self._set_status("need_file", "err")
            self._on_browse()
            if not self._input_file:
                return
        meta, source = self._meta(), self._input_file
        options = dict(txt_paragraph_mode=self.txt_mode.get(), enable_heading_detection=self.opt["headings"].get(), enable_block_quote=self.opt["bq"].get(), markdown_emphasis=self.opt["md"].get(), strip_existing_heading=self.opt["strip"].get())
        options.update(preserve_paragraphs=meta["preserve_paragraphs"], source_fingerprint=meta["source_fingerprint"])
        # Native save/date dialogs run their own event loop. Queue Finder's
        # document events until the dialog ends, keeping this draft fixed.
        self._choosing_output = True
        try:
            if not self._check(meta):
                return
            out = self._ask_save(meta, "save_blank" if blank else "save_as")
            if not out:
                return
            if source and _same_file(source, out):
                self._set_status("same_file", "err")
                return
            self._set_busy(True)
        finally:
            self._choosing_output = False
            if not self._busy:
                self._load_pending_file()
        kwargs = dict(name=meta["name"], instructor=meta["instructor"], course=meta["course"], date_str=meta["date"], title=meta["title"], last_name_override=meta["last_name"],
                      affiliation=meta["affiliation"], format_id=meta["format_id"], custom_profile=meta["custom_profile"])
        self._save_prefs()
        self._set_status("working")
        def worker():
            try:
                stats = create_blank_template(out, **kwargs) if blank else convert_draft_to_mla(source, out, **kwargs, **options)
                self._results.put((out, blank, stats, None))
            except Exception as e:
                self._results.put((out, blank, None, e))
        threading.Thread(target=worker, name="MLAFMT-convert").start()

    def _poll_results(self):
        try:
            out, blank, stats, error = self._results.get_nowait()
        except queue.Empty:
            pass
        else:
            self._set_busy(False)
            if error:
                self._report_error(error)
            else:
                self._last_output = os.path.abspath(out)
                self._review_counts = ((stats or {}).get("omitted_elements", 0), (stats or {}).get("unsupported_headings", 0))
                self._reference_sort_skipped = bool((stats or {}).get("reference_sort_skipped", 0))
                self._retained_counts = tuple((stats or {}).get(key, 0) for key in
                                              ("preserved_paragraphs", "retained_images", "retained_tables"))
                if blank:
                    self._set_status("done_blank", "ok", f=os.path.basename(out))
                else:
                    removed = stats.get("removed_heading_lines", 0)
                    self._set_status("done_removed" if removed else "done", "ok", f=os.path.basename(out), b=stats["body"], w=stats["works_cited"], r=removed if removed else "")
                    if stats["body"] == 0 and not (stats.get("retained_images", 0) or stats.get("retained_tables", 0)):
                        messagebox.showwarning(self.t("app"), self.t("empty"), parent=self.root)
                self._render_result()
                self._render_review()
                if self.opt["open_after"].get():
                    self._open_file(out)
            self._load_pending_file(None if error else out)
        self.root.after(80, self._poll_results)

    def _on_close(self):
        if self._closing:
            return
        self._closing = True
        self._stop_progress_animation()
        self._save_prefs()
        self.root.destroy()

    def run(self):
        self.root.mainloop()


def _smoke_test(require_dnd=False):
    """Exercise the installed GUI and actual document engine in a clean folder."""
    from docx import Document
    from docx.oxml.ns import qn

    app = MLAGui(remember=False)
    try:
        if require_dnd and not app.has_dnd:
            raise RuntimeError("The packaged tkDnD library could not load")
        app.root.withdraw()
        for language in ("zh", "en"):
            app._set_lang(language)
            app.root.update_idletasks()
        with tempfile.TemporaryDirectory(prefix="mlafmt-smoke-") as folder:
            folder = Path(folder)
            blank = folder / "blank.docx"
            create_blank_template(blank, name="Jane Doe", title="Smoke test")
            assert Document(blank).paragraphs[0].text == "Jane Doe"
            draft = folder / "Draft with spaces, 中文.txt"
            draft.write_text("A paragraph about *Hamlet*.\n\nWorks Cited\n\n"
                             "Zed. Last source.\n\nAlpha. First source.", encoding="utf-8")
            result = folder / "MLA result, 中文.docx"
            stats = convert_draft_to_mla(draft, result, name="Jane Doe", title="Smoke test",
                                        instructor="", course="", date_str="")
            document = Document(result)
            paragraphs = document.paragraphs
            heading = next(p for p in paragraphs if p.text == "Works Cited")
            assert heading.paragraph_format.page_break_before
            assert stats["body"] == 1 and stats["works_cited"] == 2
            assert paragraphs[-2].text.startswith("Alpha.")
            assert paragraphs[-1].paragraph_format.left_indent.inches == 0.5
            assert paragraphs[-1].paragraph_format.first_line_indent.inches == -0.5
            section = document.sections[0]
            assert section.page_width.inches == 8.5 and section.page_height.inches == 11
            assert section.left_margin.inches == section.right_margin.inches == 1
            normal = document.styles["Normal"]
            assert normal.font.name == "Times New Roman" and normal.font.size.pt == 12
            assert paragraphs[5].paragraph_format.line_spacing == 2
            assert paragraphs[-1].paragraph_format.line_spacing == 2
            assert any("PAGE" in (node.text or "") for node in
                       section.header._element.iter(qn("w:instrText")))
    finally:
        app.root.destroy()


if __name__ == "__main__":
    if "--smoke-test" in sys.argv:
        _smoke_test(require_dnd="--require-dnd" in sys.argv)
    else:
        MLAGui(initial_files=_launch_files(sys.argv[1:])).run()
