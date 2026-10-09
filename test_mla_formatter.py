"""
test_mla_formatter.py - regression tests for the MLA engine.
Run:  python -m pytest test_mla_formatter.py -q      (or)   python test_mla_formatter.py
"""
import io
import os
import tempfile
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches

import mla_formatter as m

TMP = tempfile.mkdtemp(prefix="mlafmt_test_")
META = dict(name="Jane Doe", instructor="Professor Smith",
            course="ENG 101", date_str="8 October 2026", title="On Reading")


def _txt(content, name="d.txt", enc="utf-8"):
    p = os.path.join(TMP, name)
    with open(p, "w", encoding=enc, newline="") as f:
        f.write(content)
    return p


def _convert(path, **kw):
    out = os.path.join(TMP, "out.docx")
    args = dict(META, **kw)
    stats = m.convert_draft_to_mla(path, out, **args)
    return Document(out), stats


def _texts(doc):
    return [p.text for p in doc.paragraphs]


# ---------------- names / dates ----------------

def test_last_name_basic():
    assert m.extract_last_name("Alice Johnson") == "Johnson"


def test_last_name_suffix_ignored():
    assert m.extract_last_name("Mary Jane Smith Jr.") == "Smith"
    assert m.extract_last_name("Martin Luther King, Jr.") == "King"
    assert m.extract_last_name("Henry Ford III") == "Ford"


def test_last_name_override_and_single():
    assert m.extract_last_name("Mary Smith Jr.", "Smith-Jones") == "Smith-Jones"
    assert m.extract_last_name("Cher") == "Cher"
    assert m.extract_last_name("") == "LastName"


def test_today_mla():
    import datetime
    assert m.today_mla(datetime.date(2026, 10, 8)) == "8 October 2026"
    assert m.date_warning("8 October 2026") is None
    assert m.date_warning("10/08/2026") is not None


def test_filename_sanitized():
    assert m.generate_filename("Ana de la Cruz") == "Cruz_MLA_Essay.docx"
    assert m.generate_filename("X", "Van Der Berg") == "Van_Der_Berg_MLA_Essay.docx"


# ---------------- layout ----------------

def test_page_setup_and_header():
    doc, _ = _convert(_txt("Body."))
    s = doc.sections[0]
    assert s.left_margin == Inches(1) and s.page_width == Inches(8.5)
    assert s.header_distance == Inches(0.5)
    hdr = s.header.paragraphs[0]
    assert hdr.text.startswith("Doe ")
    xml = hdr._p.xml
    assert "PAGE" in xml and 'fldCharType="separate"' in xml


def test_no_update_fields_prompt():
    doc, _ = _convert(_txt("Body."))
    assert "updateFields" not in doc.settings.element.xml


def test_heading_block_and_title():
    doc, _ = _convert(_txt("Body text."))
    t = _texts(doc)
    assert t[:5] == ["Jane Doe", "Professor Smith", "ENG 101", "8 October 2026", "On Reading"]
    assert doc.paragraphs[4].alignment == 1  # centered
    assert doc.paragraphs[5].paragraph_format.first_line_indent == Inches(0.5)


def test_blank_date_defaults_to_today():
    out = os.path.join(TMP, "x.docx")
    m.convert_draft_to_mla(_txt("Body."), out, **dict(META, date_str=""))
    assert Document(out).paragraphs[3].text == m.today_mla()


def test_fonts_have_no_theme_override():
    doc, _ = _convert(_txt("Body."))
    xml = doc.styles.element.xml
    assert "minorHAnsi" not in xml.split("<w:style ")[0]   # docDefaults clean
    assert doc.styles["Normal"].font.name == "Times New Roman"


# ---------------- txt parsing ----------------

def test_hard_wrapped_lines_are_joined():
    doc, _ = _convert(_txt("This is a hard\nwrapped paragraph\nfrom an editor.\n\nSecond."))
    assert "This is a hard wrapped paragraph from an editor." in _texts(doc)
    assert not any("\n" in t for t in _texts(doc))


def test_crlf_bom_and_line_by_line():
    p = _txt("﻿One\r\nTwo\r\n", name="crlf.txt")
    doc, st = _convert(p, txt_paragraph_mode="line_by_line")
    assert _texts(doc)[5:] == ["One", "Two"] and st["body"] == 2


def test_gbk_encoded_txt():
    p = _txt("中文段落。\n\nSecond.", name="gbk.txt", enc="gbk")
    doc, _ = _convert(p)
    assert "中文段落。" in _texts(doc)


def test_markdown_italics_in_txt():
    doc, _ = _convert(_txt("In *Hamlet* the prince waits. See file_name_here."))
    p = doc.paragraphs[5]
    runs = [(r.text, r.italic) for r in p.runs]
    assert ("Hamlet", True) in runs
    assert p.text == "In Hamlet the prince waits. See file_name_here."


def test_markdown_off():
    doc, _ = _convert(_txt("Keep *stars*."), markdown_emphasis=False)
    assert doc.paragraphs[5].text == "Keep *stars*."


def test_empty_file():
    doc, st = _convert(_txt(""))
    assert st["body"] == 0 and len(doc.paragraphs) == 5


# ---------------- docx parsing ----------------

def test_docx_keeps_italics_and_strips_tabs():
    src = Document()
    p = src.add_paragraph("\tSee ")
    p.add_run("Hamlet").italic = True
    p.add_run(" here.")
    path = os.path.join(TMP, "src.docx")
    src.save(path)
    doc, _ = _convert(path)
    body = doc.paragraphs[5]
    assert body.text == "See Hamlet here."
    assert [r.italic for r in body.runs] == [None, True, None]


def test_docx_keeps_hyperlink_text_in_order():
    src = Document()
    p = src.add_paragraph("Read ")
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), p.part.relate_to("https://example.com", RT.HYPERLINK,
                                             is_external=True))
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    properties.append(OxmlElement("w:i"))
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = "Hamlet"
    run.append(text)
    hyperlink.append(run)
    p._p.append(hyperlink)
    p.add_run(" today.")
    path = os.path.join(TMP, "hyperlink.docx")
    src.save(path)

    doc, _ = _convert(path)
    assert doc.paragraphs[5].text == "Read Hamlet today."
    assert any(r.text == "Hamlet" and r.italic for r in doc.paragraphs[5].runs)


def test_docx_keeps_inherited_emphasis_and_explicit_off():
    src = Document()
    parent = src.styles.add_style("Essay Emphasis", WD_STYLE_TYPE.CHARACTER)
    parent.font.italic = True
    child = src.styles.add_style("Inherited Emphasis", WD_STYLE_TYPE.CHARACTER)
    child.base_style = parent
    paragraph_style = src.styles.add_style("Essay Body", WD_STYLE_TYPE.PARAGRAPH)
    paragraph_style.font.bold = True
    p = src.add_paragraph()
    p.style = paragraph_style
    run = p.add_run("Hamlet")
    run.style = child
    run = p.add_run(" stays plain.")
    run.bold = False
    run.italic = False
    path = os.path.join(TMP, "styled.docx")
    src.save(path)

    doc, _ = _convert(path)
    runs = doc.paragraphs[5].runs
    assert runs[0].text == "Hamlet" and runs[0].italic and runs[0].bold
    assert runs[1].text == " stays plain." and not runs[1].italic and not runs[1].bold


def test_existing_heading_block_removed():
    src = Document()
    for t in ["Jane Doe", "Professor Smith", "ENG 101", "8 October 2026", "On Reading",
              "Real first paragraph."]:
        src.add_paragraph(t)
    path = os.path.join(TMP, "dup.docx")
    src.save(path)
    doc, st = _convert(path)
    assert st["removed_heading_lines"] == 5
    assert _texts(doc).count("Jane Doe") == 1
    assert _texts(doc)[5] == "Real first paragraph."


def test_existing_heading_kept_when_disabled():
    doc, st = _convert(_txt("Jane Doe\n\nBody."), strip_existing_heading=False)
    assert _texts(doc).count("Jane Doe") == 2


def test_single_matching_title_or_name_is_kept_as_body():
    for first in (META["name"], META["title"], META["title"] + "."):
        doc, stats = _convert(_txt(first + "\n\nBody."))
        assert _texts(doc)[5:] == [first, "Body."]
        assert stats["removed_heading_lines"] == 0


def test_existing_heading_with_running_header_and_italic_title():
    text = "\n\n".join(["Doe 1", META["name"], META["instructor"], META["course"],
                          META["date_str"], "On Hamlet", "Body."])
    doc, stats = _convert(_txt(text), title="On *Hamlet*")
    assert stats["removed_heading_lines"] == 6
    assert _texts(doc)[5:] == ["Body."]
    assert any(r.text == "Hamlet" and r.italic for r in doc.paragraphs[4].runs)


# ---------------- Works Cited ----------------

WC_DRAFT = ("Body.\n\nWorks Cited:\n\n"
            "\"The Zebra Essay.\" Web, 2020.\n\n"
            "Smith, John. *Some Book*. Publisher, 2020.\n\n"
            "The Atlantic Staff. Article. 2019.\n\n"
            "Brown, Alice. Third Source. 2019.")


def test_works_cited_new_page_no_blank_line():
    doc, st = _convert(_txt(WC_DRAFT))
    paras = doc.paragraphs
    idx = _texts(doc).index("Works Cited")
    assert paras[idx].paragraph_format.page_break_before is True
    assert paras[idx - 1].text == "Body."          # no empty spacer paragraph
    assert st["works_cited"] == 4


def test_works_cited_sorted_ignoring_quotes_and_articles():
    doc, _ = _convert(_txt(WC_DRAFT))
    t = _texts(doc)
    entries = t[t.index("Works Cited") + 1:]
    assert [e.split()[0] for e in entries] == ["The", "Brown,", "Smith,", "\"The"]
    # Atlantic < Brown < Smith < Zebra


def test_works_cited_hanging_indent_and_italic():
    doc, _ = _convert(_txt(WC_DRAFT))
    t = _texts(doc)
    p = doc.paragraphs[t.index("Works Cited") + 3]   # Smith
    assert p.paragraph_format.first_line_indent == Inches(-0.5)
    assert any(r.italic and r.text == "Some Book" for r in p.runs)


def test_works_cited_variants():
    for h in ("Work Cited", "WORKS CITED.", "Bibliography", "References"):
        doc, st = _convert(_txt(f"Body.\n\n{h}\n\nA, B. C. 2020."))
        assert st["works_cited"] == 1 and "Works Cited" in _texts(doc), h
    doc, _ = _convert(_txt("Body.\n\nWorks Consulted\n\nA, B. C."))
    assert "Works Consulted" in _texts(doc)


# ---------------- options ----------------

def test_block_quote_multiline():
    doc, _ = _convert(_txt("Intro.\n\n> line one of the quote\n> line two\n\nAfter."),
                      enable_block_quote=True)
    bq = doc.paragraphs[6]
    assert bq.text == "line one of the quote line two"
    assert bq.paragraph_format.left_indent == Inches(0.5)
    assert bq.paragraph_format.first_line_indent == Inches(0)


def test_headings_opt_in():
    doc, _ = _convert(_txt("Introduction\n\nBody."))
    assert not any(r.bold for r in doc.paragraphs[5].runs)
    doc, _ = _convert(_txt("Introduction\n\nBody."), enable_heading_detection=True)
    assert all(r.bold for r in doc.paragraphs[5].runs)


def test_stream_output_and_blank_template():
    buf = io.BytesIO()
    m.create_blank_template(buf, name="Jane Doe", title="T")
    doc = Document(io.BytesIO(buf.getvalue()))
    assert doc.paragraphs[3].text == m.today_mla()
    assert doc.paragraphs[-1].text.startswith("[Start writing")


def test_conversion_rejects_replacing_original_draft():
    source = Path(TMP) / "original.docx"
    doc = Document()
    doc.add_paragraph("Original text.")
    doc.save(source)
    original = source.read_bytes()
    try:
        m.convert_draft_to_mla(source, source, **META)
    except ValueError as error:
        assert "different file" in str(error)
    else:
        raise AssertionError("expected source-overwrite protection")
    assert source.read_bytes() == original


def test_conversion_rejects_hardlink_to_original_draft():
    source = Path(_txt("Original text.", name="original-hardlink.txt"))
    target = Path(TMP) / "original-hardlink-alias.docx"
    try:
        os.link(source, target)
    except OSError:  # Some filesystems do not support hard links.
        return
    try:
        m.convert_draft_to_mla(source, target, **META)
    except ValueError:
        pass
    else:
        raise AssertionError("expected source-alias protection")
    assert source.read_text(encoding="utf-8") == "Original text."


def test_failed_save_preserves_existing_output_and_removes_temporary_file():
    source = _txt("Body.", name="save-failure.txt")
    target = Path(TMP) / "preserved-output.docx"
    original = b"existing document bytes"
    target.write_bytes(original)
    files_before = set(os.listdir(TMP))

    def fail_after_partial_write(path):
        Path(path).write_bytes(b"incomplete output")
        raise OSError("simulated save failure")

    with patch("docx.document.Document.save", side_effect=fail_after_partial_write):
        try:
            m.convert_draft_to_mla(source, target, **META)
        except OSError:
            pass
        else:
            raise AssertionError("expected simulated save failure")
    assert target.read_bytes() == original
    assert set(os.listdir(TMP)) == files_before


def test_completed_save_replaces_existing_output():
    target = Path(TMP) / "replace-output.docx"
    target.write_bytes(b"old output")
    m.convert_draft_to_mla(_txt("New body."), target, **META)
    assert Document(target).paragraphs[5].text == "New body."


def test_unsupported_extension():
    try:
        m.convert_draft_to_mla(_txt("x", name="a.pdf"), os.path.join(TMP, "o.docx"), **META)
    except ValueError:
        return
    raise AssertionError("expected ValueError")


def test_cli(capsys=None):
    out = os.path.join(TMP, "cli.docx")
    assert m._cli([_txt("Body."), "-o", out, "--name", "A B", "--title", "T"]) == 0
    assert os.path.exists(out)


if __name__ == "__main__":
    import inspect
    fns = [(n, f) for n, f in sorted(globals().items()) if n.startswith("test_") and inspect.isfunction(f)]
    failed = 0
    for n, f in fns:
        try:
            f()
            print("PASS", n)
        except Exception as e:  # noqa
            failed += 1
            print("FAIL", n, "->", repr(e))
    print(f"\n{len(fns) - failed}/{len(fns)} passed")
    raise SystemExit(1 if failed else 0)
