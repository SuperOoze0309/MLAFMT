"""Desktop regressions using real Tk widgets and isolated preferences.

Run with ``python -m unittest test_mla_gui -v``. Only inability to initialize
Tk skips GUI tests, so Linux CI without a display can still run config tests.
"""

import gc
import json
import os
from pathlib import Path
import tempfile
import threading
import time
from types import SimpleNamespace
import unittest
from unittest.mock import patch

from docx import Document

import mla_gui as gui


class ConfigTests(unittest.TestCase):
    def _patch(self, context):
        value = context.__enter__()
        self.addCleanup(context.__exit__, None, None, None)
        return value

    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="mlafmt_gui_config_")
        self.addCleanup(self.temporary.cleanup)
        self.config = Path(self.temporary.name) / "settings.json"
        self.legacy = Path(self.temporary.name) / "legacy.json"
        self._patch(patch.object(gui, "CONFIG_PATH", str(self.config)))
        self._patch(patch.object(gui, "LEGACY_CONFIG_PATH", str(self.legacy)))

    def test_corrupt_config_falls_back_to_legacy(self):
        self.config.write_text("{invalid json", encoding="utf-8")
        self.legacy.write_text(json.dumps({"lang": "zh"}), encoding="utf-8")
        self.assertEqual(gui._load_config(), {"lang": "zh"})

    def test_non_object_or_missing_config_uses_defaults(self):
        for value in ([], None, "unexpected", 12):
            with self.subTest(value=value):
                self.config.write_text(json.dumps(value), encoding="utf-8")
                self.assertEqual(gui._load_config(), {})
        self.config.unlink()
        self.assertEqual(gui._load_config(), {})


class DesktopTests(unittest.TestCase):
    def _patch(self, context):
        value = context.__enter__()
        self.addCleanup(context.__exit__, None, None, None)
        return value

    @classmethod
    def setUpClass(cls):
        # Probe Tk itself, separately from constructing the application. A Tcl
        # error inside MLAGui is an application failure and must fail the test.
        try:
            root = gui.tk.Tk()
        except gui.tk.TclError as error:
            raise unittest.SkipTest(f"Tk cannot initialize on this host: {error}")
        root.withdraw()
        # Tcl interpreters are process-wide native resources. Keep one alive
        # for the suite, rather than letting GC finalize a previous interpreter
        # while the next Tk instance is loading its native scripts on Windows.
        cls.tk_root = root
        cls.addClassCleanup(root.destroy)

    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="mlafmt_gui_")
        self.addCleanup(self.temporary.cleanup)
        self.folder = Path(self.temporary.name)
        self.config = self.folder / "isolated-settings.json"
        self._patch(patch.object(gui, "CONFIG_PATH", str(self.config)))
        self._patch(patch.object(gui, "LEGACY_CONFIG_PATH", str(self.folder / "legacy.json")))
        self.save_config = self._patch(patch.object(gui, "_save_config", return_value=True))
        self._patch(patch.object(gui, "_default_lang", return_value="en"))
        self._patch(patch.object(gui, "HAS_DND", False))
        self.showerror = self._patch(patch.object(gui.messagebox, "showerror"))
        self._patch(patch.object(gui.messagebox, "showwarning"))
        self._patch(patch.object(gui.messagebox, "askyesno", return_value=True))
        self.apps = []
        self.instances = []
        self.interpreter_commands = set(self.tk_root._tclCommands or ())
        self.worker_threads = []
        real_thread = threading.Thread

        def track_thread(*args, **kwargs):
            thread = real_thread(*args, **kwargs)
            self.worker_threads.append(thread)
            return thread

        self._patch(patch.object(gui.threading, "Thread", side_effect=track_thread))
        self.callback_errors = []
        self.addCleanup(self._cleanup_apps)
        self.app = self._new_app()
        self.draft = self.folder / "draft with spaces, 中文.txt"
        self.draft.write_text("First paragraph.\n\nSecond paragraph.", encoding="utf-8")

    def _new_app(self, remember=False):
        # Track the root even if application construction fails, for cleanup.
        root = gui.tk.Toplevel(self.tk_root)
        root.withdraw()
        self.apps.append(root)
        with patch.object(gui.tk, "Tk", return_value=root):
            app = gui.MLAGui(remember=remember)
        self.instances.append(app)
        root.report_callback_exception = lambda *error: self.callback_errors.append(error)
        root.update_idletasks()
        return app

    def _cleanup_apps(self):
        # Wait for real worker termination, not just receipt of its queue item.
        # All Tcl objects and trace callbacks are then disposed on this thread.
        for thread in self.worker_threads:
            thread.join(timeout=5)
            self.assertFalse(thread.is_alive(), "Background worker survived the test")
        for app in self.instances:
            for variable in app.v.values():
                for mode, callback in variable.trace_info():
                    variable.trace_remove(mode, callback)
        for root in self.apps:
            for callback in root.tk.splitlist(root.tk.call("after", "info")):
                script, _kind = root.tk.call("after", "info", callback)
                command = root.tk.splitlist(script)[0]
                if command in (root._tclCommands or ()):
                    root.after_cancel(callback)
            root.unbind_all("<MouseWheel>")
            root.destroy()
        # bind_all registers callbacks on the interpreter's root. Remove only
        # commands introduced by this test, after destroying its own windows.
        for command in set(self.tk_root._tclCommands or ()) - self.interpreter_commands:
            self.tk_root.deletecommand(command)
        self.app = None
        self.instances.clear()
        self.apps.clear()
        gc.collect()
        self.assertEqual(self.callback_errors, [], "Tk callback raised an application exception")

    def _fill(self):
        values = {"name": "Jane Smith Jr.", "title": "On *Hamlet* 与记忆",
                  "instructor": "Professor Doe", "course": "ENG 101",
                  "date": "8 October 2026", "last_name": "Smith-Jones"}
        for key, value in values.items():
            self.app.v[key].set(value)
        return values

    def _switch(self, language):
        button = next(button for button in self.app._language_buttons
                      if button.cget("value") == language)
        button.invoke()
        self.app.root.update_idletasks()
        self.assertEqual(self.app.lang, language)
        self.assertEqual(self.callback_errors, [])

    def _wait_for_idle(self, timeout=5):
        deadline = time.monotonic() + timeout
        while self.app._busy and time.monotonic() < deadline:
            self.app.root.update()
            time.sleep(0.01)
        self.assertFalse(self.app._busy, "Background conversion did not complete")
        self.assertEqual(self.callback_errors, [])

    def test_language_buttons_preserve_all_inputs_file_and_advanced_options(self):
        expected = self._fill()
        self.app._on_drop(SimpleNamespace(data="{" + str(self.draft) + "}"))
        self.app._open_advanced()
        self.app.txt_mode.set("line_by_line")
        options = {"md": False, "strip": False, "headings": True,
                   "bq": True, "open_after": True}
        for key, value in options.items():
            self.app.opt[key].set(value)
        for language in ("zh", "en"):
            with self.subTest(language=language):
                self._switch(language)
                self.assertEqual({k: v.get() for k, v in self.app.v.items()}, expected)
                self.assertEqual(self.app._input_file, str(self.draft.resolve()))
                self.assertEqual(self.app.txt_mode.get(), "line_by_line")
                self.assertEqual({k: v.get() for k, v in self.app.opt.items()}, options)
                self.assertTrue(self.app._adv_win.winfo_exists())
                self.assertEqual(self.app._adv_win.title(), gui.T[language]["adv_title"])
                self.assertEqual(self.app._status.cget("text"),
                                 gui.T[language]["loaded"].format(f=self.draft.name))
                self.assertIn(self.draft.name, self.app._drop.cget("text"))
                self.assertIn(gui.T[language]["file_size"].split("{size}")[1],
                              self.app._drop.cget("text"))
        self.save_config.assert_not_called()
        self.assertFalse(self.config.exists())

    def test_error_status_is_translated_after_language_switch(self):
        self.app._report_error(PermissionError("locked"))
        for language in ("zh", "en"):
            self._switch(language)
            self.assertEqual(self.app._status.cget("text"), gui.T[language]["perm"])
        self.showerror.assert_called_once()


    def test_unicode_entry_edits_survive_language_rebuild(self):
        values = {"name": "欧阳 慧 Jane", "title": "《哈姆雷特》中的记忆 — café e\u0301",
                  "instructor": "王老师", "course": "文学 101"}
        for key, value in values.items():
            entry = self.app._entries[key]
            entry.delete(0, gui.tk.END)
            entry.insert(0, value)
        for language in ("zh", "en"):
            self._switch(language)
            for key, value in values.items():
                self.assertEqual(self.app._entries[key].get(), value)


    def test_last_name_and_advanced_preferences_are_saved_and_restored(self):
        self._fill()
        self.app.v["last_name"].set("欧阳-Smith")
        self.app.txt_mode.set("line_by_line")
        self.app.opt["headings"].set(True)
        # Saving is mocked and all config paths point inside the test folder.
        self.app.remember = True
        self._switch("zh")
        saved = self.save_config.call_args.args[0]
        self.assertEqual(saved["last_name"], "欧阳-Smith")
        self.assertEqual(saved["lang"], "zh")
        self.assertEqual(saved["txt_mode"], "line_by_line")
        self.assertTrue(saved["headings"])
        self.config.write_text(json.dumps(saved, ensure_ascii=False), encoding="utf-8")
        restored = self._new_app(remember=True)
        self.assertEqual(restored.lang, "zh")
        self.assertEqual(restored.v["last_name"].get(), "欧阳-Smith")
        self.assertEqual(restored.txt_mode.get(), "line_by_line")
        self.assertTrue(restored.opt["headings"].get())

    def test_success_status_and_result_actions_are_translated(self):
        out = self.folder / "essay final, 稿.docx"
        self.app._set_busy(True)
        self.app._results.put((str(out), False,
                               {"body": 2, "works_cited": 1, "removed_heading_lines": 0}, None))
        self.app._poll_results()
        for language in ("zh", "en"):
            self._switch(language)
            expected = gui.T[language]["done"].format(f=out.name, b=2, w=1, r="")
            self.assertEqual(self.app._status.cget("text"), expected)
            labels = [child.cget("text") for child in self.app._result_row.winfo_children()]
            self.assertEqual(labels, [gui.T[language]["open_result"], gui.T[language]["view_folder"]])
            self.assertEqual(self.app._last_output, str(out.resolve()))

    def test_blank_success_status_is_translated(self):
        out = self.folder / "blank.docx"
        self.app._set_busy(True)
        self.app._results.put((str(out), True, None, None))
        self.app._poll_results()
        self._switch("zh")
        self.assertEqual(self.app._status.cget("text"),
                         gui.T["zh"]["done_blank"].format(f=out.name))


    def test_removed_heading_count_is_translated_with_success_status(self):
        out = self.folder / "clean-heading.docx"
        self.app._set_busy(True)
        self.app._results.put((str(out), False,
                               {"body": 2, "works_cited": 1, "removed_heading_lines": 5}, None))
        self.app._poll_results()
        for language in ("zh", "en"):
            self._switch(language)
            expected = gui.T[language]["done_removed"].format(f=out.name, b=2, w=1, r=5)
            self.assertEqual(self.app._status.cget("text"), expected)

    def test_corrupt_config_and_invalid_language_do_not_break_startup(self):
        for contents in ("{broken", json.dumps({"lang": "xx"}),
                         json.dumps({"lang": []}), json.dumps({"lang": {}}),
                         json.dumps({"lang": None, "name": [], "md": "false",
                                     "txt_mode": ["line_by_line"]})):
            with self.subTest(contents=contents):
                self.config.write_text(contents, encoding="utf-8")
                app = self._new_app(remember=True)
                self.assertEqual(app.lang, "en")
                self.assertEqual(app.txt_mode.get(), "blank_line")
                self.assertIsInstance(app.v["name"].get(), str)

    def test_template_and_conversion_reject_original_or_hardlink_output(self):
        self._fill()
        self.app._set_file(str(self.draft))
        targets = [self.draft]
        alias = self.folder / "draft-alias.docx"
        try:
            os.link(self.draft, alias)
        except OSError:
            pass
        else:
            targets.append(alias)
        original = self.draft.read_bytes()
        with patch.object(gui, "create_blank_template") as blank, \
                patch.object(gui, "convert_draft_to_mla") as convert:
            for action in (self.app._on_blank, self.app._on_convert):
                for target in targets:
                    with self.subTest(action=action.__name__, target=target.name), \
                            patch.object(gui.filedialog, "asksaveasfilename", return_value=str(target)):
                        action()
                        self.assertFalse(self.app._busy)
                        self.assertEqual(self.app._status_state[0], "same_file")
            blank.assert_not_called()
            convert.assert_not_called()
        self.assertEqual(self.draft.read_bytes(), original)

    def test_conversion_is_async_ignores_repeat_click_and_updates_result(self):
        self._fill()
        self.app._set_file(str(self.draft))
        self.app.txt_mode.set("line_by_line")
        self.app.opt["bq"].set(True)
        out = self.folder / "async output.docx"
        entered, release = threading.Event(), threading.Event()
        calling_threads = []

        def convert(*args, **kwargs):
            calling_threads.append(threading.get_ident())
            entered.set()
            if not release.wait(5):
                raise TimeoutError("test did not release worker")
            return {"body": 2, "works_cited": 0, "removed_heading_lines": 0}

        with patch.object(gui.filedialog, "asksaveasfilename", return_value=str(out)) as save, \
                patch.object(gui, "convert_draft_to_mla", side_effect=convert) as conversion:
            try:
                self.app._btn.invoke()
                self.assertTrue(entered.wait(2))
                self.assertTrue(self.app._busy)
                self.assertEqual(self.app._status_state[0], "working")
                for button in (self.app._btn, self.app._blank_btn, self.app._advanced_btn,
                               *self.app._language_buttons):
                    self.assertTrue(button.instate(["disabled"]))
                self.app._on_convert()
                self.app._on_blank()
                self.app._btn.invoke()
                save.assert_called_once()
                conversion.assert_called_once()
                self.assertNotEqual(calling_threads, [threading.get_ident()])
                self.assertEqual(conversion.call_args.kwargs["txt_paragraph_mode"], "line_by_line")
                self.assertTrue(conversion.call_args.kwargs["enable_block_quote"])
            finally:
                release.set()
            self._wait_for_idle()
        self.assertEqual(self.app._last_output, str(out.resolve()))
        self.assertEqual(self.app._status_state[0], "done")
        self.assertFalse(self.app._btn.instate(["disabled"]))

    def test_background_error_reenables_controls_and_can_be_translated(self):
        self._fill()
        self.app._set_file(str(self.draft))
        with patch.object(gui.filedialog, "asksaveasfilename", return_value=str(self.folder / "failed.docx")), \
                patch.object(gui, "convert_draft_to_mla", side_effect=PermissionError("locked")):
            self.app._on_convert()
            self._wait_for_idle()
        self.assertIsNone(self.app._last_output)
        self.assertFalse(self.app._btn.instate(["disabled"]))
        self.assertEqual(self.app._status_state[0], "perm")
        self._switch("zh")
        self.assertEqual(self.app._status.cget("text"), gui.T["zh"]["perm"])

    def test_real_template_and_conversion_write_valid_documents(self):
        self._fill()
        for blank in (True, False):
            with self.subTest(blank=blank):
                out = self.folder / ("blank.docx" if blank else "converted.docx")
                if not blank:
                    self.app._set_file(str(self.draft))
                with patch.object(gui.filedialog, "asksaveasfilename", return_value=str(out)):
                    self.app._on_blank() if blank else self.app._on_convert()
                    self._wait_for_idle()
                doc = Document(out)
                self.assertEqual(doc.paragraphs[0].text, "Jane Smith Jr.")
                self.assertEqual(doc.paragraphs[5].text,
                                 "[Start writing your essay here.]" if blank else "First paragraph.")
                self.assertEqual(self.app._last_output, str(out.resolve()))


if __name__ == "__main__":
    unittest.main()
