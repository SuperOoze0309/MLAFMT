"""Focused regressions for reference-page breaks in retained DOCX packages."""

import io
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from mla_formatter import convert_draft_to_mla, inspect_draft_paragraphs


@pytest.mark.parametrize("format_id,custom_profile,reference_title", [
    ("mla", None, "Works Cited"),
    ("apa", None, "References"),
    ("custom", {"references_title": "Course Sources"}, "Course Sources"),
])
@pytest.mark.parametrize("reference_count", [1, 2])
@pytest.mark.parametrize("preserve_heading", [False, True])
def test_docx_reference_title_starts_new_page_without_losing_protected_content(
        tmp_path, format_id, custom_profile, reference_title, reference_count, preserve_heading):
    source = Document()
    for text, font_name, font_size in (
            ("Keep this paragraph's original style.", "Courier New", 17),
            ("Keep this second paragraph's original style too.", "Arial", 11)):
        paragraph = source.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(14)
        paragraph.paragraph_format.line_spacing = 1.5
        run = paragraph.add_run(text)
        run.font.name, run.font.size, run.bold = font_name, Pt(font_size), True
    protected_xml = {paragraph.text: paragraph._p.xml for paragraph in source.paragraphs}

    source.add_paragraph("An ordinary body paragraph.")
    source.add_picture(str(Path(__file__).parent / "static" / "capybara-64.png"),
                       width=Inches(0.4))
    table = source.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Keep this table."
    original_table_xml = table._tbl.xml
    linked = source.add_paragraph("A body hyperlink: ")
    hyperlink = OxmlElement("w:hyperlink")
    relationship = source.part.relate_to("https://example.com/", RT.HYPERLINK,
                                        is_external=True)
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "example source"
    run.append(text)
    hyperlink.append(run)
    linked._p.append(hyperlink)

    # An ordinary reference heading follows the branch that rebuilds its XML.
    reference_heading = source.add_paragraph(reference_title)
    selected = [1, 2]
    if preserve_heading:
        reference_heading.runs[0].font.name = "Courier New"
        reference_heading.runs[0].font.size = Pt(16)
        protected_xml[reference_title] = reference_heading._p.xml
        selected.append(len(source.paragraphs))
    for entry in ("Zed, Z. A source.", "Alpha, A. Another source.")[:reference_count]:
        source.add_paragraph(entry)
    source_path = tmp_path / "draft.docx"
    source.save(source_path)
    fingerprint = inspect_draft_paragraphs(source_path)["source_fingerprint"]
    output = io.BytesIO()
    stats = convert_draft_to_mla(
        source_path, output, name="Jordan Example", instructor="Professor Example",
        course="ENG 101", date_str="1 January 2026", title="Example Essay",
        affiliation="Example University", format_id=format_id,
        custom_profile=custom_profile, strip_existing_heading=False,
        preserve_paragraphs=selected, source_fingerprint=fingerprint)
    output.seek(0)
    result = Document(output)

    expected_title = ("Work Cited" if format_id == "mla" and reference_count == 1 and not preserve_heading
                      else reference_title)
    heading = next(paragraph for paragraph in result.paragraphs
                   if paragraph.text == expected_title)
    if preserve_heading:
        assert heading._p.getprevious().xpath(".//w:br[@w:type='page']")
    else:
        assert heading.paragraph_format.page_break_before is True
        page_break = heading._p.pPr.find(qn("w:pageBreakBefore"))
        assert page_break is not None and page_break.get(qn("w:val"), "1") == "1"

    for text, original_xml in protected_xml.items():
        assert next(paragraph for paragraph in result.paragraphs
                    if paragraph.text == text)._p.xml == original_xml
    assert result.tables[0]._tbl.xml == original_table_xml
    retained_link = next(result.element.body.iter(qn("w:hyperlink")))
    assert result.part.rels[retained_link.get(qn("r:id"))].target_ref == "https://example.com/"
    with zipfile.ZipFile(source_path) as before, zipfile.ZipFile(io.BytesIO(output.getvalue())) as after:
        media = [name for name in before.namelist() if name.startswith("word/media/")]
        assert media
        assert all(before.read(name) == after.read(name) for name in media)
    assert stats["preserved_paragraphs"] == len(selected)
    assert stats["retained_images"] == stats["retained_tables"] == 1
    assert stats["works_cited"] == reference_count


@pytest.mark.parametrize("existing_control", [
    "direct", "inherited", "leading", "previous", "previous_paragraph", "explicit_false",
])
def test_preserved_reference_title_does_not_duplicate_existing_page_break(tmp_path, existing_control):
    source = Document()
    body = source.add_paragraph("A preserved body paragraph.")
    selected = [1]
    if existing_control == "previous":
        body.add_run().add_break(WD_BREAK.PAGE)
    if existing_control == "previous_paragraph":
        control = source.add_paragraph()
        control.paragraph_format.page_break_before = True
        selected.append(len(source.paragraphs))
    heading = source.add_paragraph()
    if existing_control == "leading":
        heading.add_run().add_break(WD_BREAK.PAGE)
    heading.add_run("Works Cited")
    if existing_control == "direct":
        heading.paragraph_format.page_break_before = True
    if existing_control in ("inherited", "explicit_false"):
        style = source.styles.add_style("Original references", WD_STYLE_TYPE.PARAGRAPH)
        style.paragraph_format.page_break_before = True
        heading.style = style
        if existing_control == "explicit_false":
            heading.paragraph_format.page_break_before = False
    selected.append(len(source.paragraphs))
    protected_xml = [source.paragraphs[index - 1]._p.xml for index in selected]
    source.add_paragraph("Alpha, A. A source.")
    path = tmp_path / "draft.docx"
    source.save(path)
    fingerprint = inspect_draft_paragraphs(path)["source_fingerprint"]
    output = io.BytesIO()
    convert_draft_to_mla(
        path, output, name="Jordan Example", instructor="Professor Example", course="ENG 101",
        date_str="1 January 2026", title="Example Essay", strip_existing_heading=False,
        preserve_paragraphs=selected, source_fingerprint=fingerprint)
    result = Document(io.BytesIO(output.getvalue()))

    # MLA adds five front-matter paragraphs. Only an explicit false override
    # requires an additional control; all other cases already start a new page.
    assert len(result.paragraphs) == len(source.paragraphs) + 5 + (existing_control == "explicit_false")
    assert all(xml in [paragraph._p.xml for paragraph in result.paragraphs] for xml in protected_xml)
    if existing_control == "explicit_false":
        result_heading = next(paragraph for paragraph in result.paragraphs if paragraph.text == "Works Cited")
        assert result_heading._p.getprevious().xpath(".//w:br[@w:type='page']")
