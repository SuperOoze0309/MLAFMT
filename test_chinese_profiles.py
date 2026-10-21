"""Focused coverage for Chinese preset layouts and retained DOCX content."""

import io
import json
import zipfile
from copy import deepcopy
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from format_profiles import (
    BUILTIN_CUSTOM_PRESETS, DEFAULT_CUSTOM_PROFILE,
    normalize_custom_profile, resolve_profile,
)
from mla_formatter import convert_draft_to_mla, inspect_draft_paragraphs


META = dict(name="张同学", instructor="王老师", course="中文写作", affiliation="示例大学",
            title="中文排版研究", date_str="2026年10月21日")
BODY = "这是中文正文，包含 English 123 与中文内容。"
REFS = ("[2] 张示例. 后列出的虚构资料. 2026.", "[1] 阿示例. 先列出的虚构资料. 2025.")
NEW_FIELDS = ("typography", "latin_font_name", "first_line_chars", "line_spacing_pt")


def _preset(preset_id):
    return deepcopy(next(item["profile"] for item in BUILTIN_CUSTOM_PRESETS
                         if item["id"] == preset_id))


def _convert(path, preset_id="zh_course", **overrides):
    settings = dict(META, format_id="custom", custom_profile=_preset(preset_id))
    settings.update(overrides)
    output = io.BytesIO()
    stats = convert_draft_to_mla(path, output, **settings)
    return Document(io.BytesIO(output.getvalue())), stats, output.getvalue()


def _find(document, text):
    return next(paragraph for paragraph in document.paragraphs if paragraph.text == text)


def _assert_font(paragraph, east_asia, size):
    runs = [element for element in paragraph._p.iter(qn("w:r"))
            if next(element.iter(qn("w:t")), None) is not None]
    assert runs
    for element in runs:
        rpr = element.find(qn("w:rPr"))
        fonts = rpr.find(qn("w:rFonts"))
        assert fonts.get(qn("w:eastAsia")) == east_asia
        assert fonts.get(qn("w:ascii")) == "Times New Roman"
        assert fonts.get(qn("w:hAnsi")) == "Times New Roman"
        assert rpr.find(qn("w:sz")).get(qn("w:val")) == str(size * 2)


def _assert_body_layout(paragraph, preset_id):
    _assert_font(paragraph, "宋体", 12)
    assert paragraph.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    indent = paragraph._p.pPr.find(qn("w:ind"))
    assert indent.get(qn("w:firstLineChars")) == "200"
    assert paragraph.paragraph_format.first_line_indent == Pt(24)
    assert not any(indent.get(qn("w:" + key)) for key in
                   ("hangingChars", "startChars", "endChars", "leftChars", "rightChars"))
    spacing = paragraph._p.pPr.find(qn("w:spacing"))
    if preset_id == "zh_thesis":
        assert paragraph.paragraph_format.line_spacing == Pt(20)
        assert paragraph.paragraph_format.line_spacing_rule == WD_LINE_SPACING.EXACTLY
        assert spacing.get(qn("w:line")) == "400"
        assert spacing.get(qn("w:lineRule")) == "exact"
    else:
        assert paragraph.paragraph_format.line_spacing == 1.5
        assert paragraph.paragraph_format.line_spacing_rule == WD_LINE_SPACING.ONE_POINT_FIVE
        assert spacing.get(qn("w:line")) == "360"
        assert spacing.get(qn("w:lineRule")) == "auto"
    assert paragraph._p.pPr.find(qn("w:snapToGrid")).get(qn("w:val")) == "0"


def test_builtin_profiles_and_resolved_copies_are_independent():
    assert {item["id"] for item in BUILTIN_CUSTOM_PRESETS} == {"zh_course", "zh_thesis"}
    first, second = BUILTIN_CUSTOM_PRESETS
    assert first["profile"] is not second["profile"]
    assert first["profile"] is not DEFAULT_CUSTOM_PROFILE
    snapshot = deepcopy(BUILTIN_CUSTOM_PRESETS)
    resolved = resolve_profile("custom", first["profile"])
    resolved["label"] = "Changed copy"
    resolved["font_size"] = 18
    assert BUILTIN_CUSTOM_PRESETS == snapshot
    assert resolve_profile("custom", first["profile"])["font_size"] == 12


def test_legacy_json_gets_neutral_new_defaults_without_mutation():
    old = {key: value for key, value in DEFAULT_CUSTOM_PROFILE.items() if key not in NEW_FIELDS}
    old["label"] = "Existing template"
    imported = json.loads(json.dumps(old))
    result = normalize_custom_profile(imported)
    assert {key: result[key] for key in NEW_FIELDS} == {
        "typography": "western", "latin_font_name": "Times New Roman",
        "first_line_chars": 0, "line_spacing_pt": 0,
    }
    assert imported == old
    for kind in ("mla", "apa"):
        profile = resolve_profile(kind)
        assert profile["typography"] == "western"
        assert profile["first_line_chars"] == profile["line_spacing_pt"] == 0
        assert profile["line_spacing"] == 2


@pytest.mark.parametrize("field,value", [
    ("typography", "unknown"), ("typography", None),
    ("latin_font_name", "unknown"), ("latin_font_name", 12),
    ("first_line_chars", -1), ("first_line_chars", 4.01),
    ("first_line_chars", True), ("first_line_chars", "2"),
    ("first_line_chars", float("nan")), ("first_line_chars", float("inf")),
    ("line_spacing_pt", -1), ("line_spacing_pt", 60.01),
    ("line_spacing_pt", False), ("line_spacing_pt", "20"),
    ("line_spacing_pt", float("nan")), ("line_spacing_pt", float("inf")),
])
def test_new_profile_fields_reject_invalid_values(field, value):
    with pytest.raises(ValueError):
        normalize_custom_profile({field: value})


@pytest.mark.parametrize("preset_id", ["zh_course", "zh_thesis"])
@pytest.mark.parametrize("extension", [".txt", ".md", ".docx"])
def test_chinese_layout_fonts_spacing_references_and_page_numbers(tmp_path, preset_id, extension):
    path = tmp_path / ("draft" + extension)
    if extension == ".docx":
        source = Document()
        source.add_paragraph("一、引言", style="Heading 1")
        source.add_paragraph(BODY)
        source.add_paragraph("参考文献")
        for reference in REFS:
            source.add_paragraph(reference)
        source.save(path)
    else:
        heading = "# 一、引言\n\n" if extension == ".md" else ""
        path.write_text(heading + BODY + "\n\n参考文献\n\n" + "\n\n".join(REFS), encoding="utf-8")
    document, stats, _ = _convert(path, preset_id)
    _assert_body_layout(_find(document, BODY), preset_id)
    for title in [paragraph for paragraph in document.paragraphs if paragraph.text == META["title"]]:
        _assert_font(title, "黑体", 16)
        assert title.alignment == WD_ALIGN_PARAGRAPH.CENTER
    if extension != ".txt":
        _assert_font(_find(document, "一、引言"), "黑体", 16)
    reference_title = _find(document, "参考文献")
    _assert_font(reference_title, "黑体", 16)
    assert reference_title.paragraph_format.page_break_before is True
    texts = [paragraph.text for paragraph in document.paragraphs]
    assert texts.index(REFS[0]) < texts.index(REFS[1])
    assert stats["works_cited"] == 2
    for reference in REFS:
        paragraph = _find(document, reference)
        assert paragraph.paragraph_format.first_line_indent == Inches(0)
        assert paragraph.paragraph_format.left_indent == Inches(0)
    for section in document.sections:
        # DOCX page dimensions are rounded to whole twips when serialized.
        assert section.page_width.cm == pytest.approx(21, abs=0.002)
        assert section.page_height.cm == pytest.approx(29.7, abs=0.002)
        assert section.left_margin.cm == pytest.approx(2.5, abs=0.002)
        assert section.top_margin.cm == pytest.approx(2.5, abs=0.002)
    if preset_id == "zh_thesis":
        assert len(document.sections) == 2
        assert document.sections[0]._sectPr.find(qn("w:footerReference")) is None
        body_section = document.sections[1]
    else:
        assert len(document.sections) == 1
        body_section = document.sections[0]
    page_type = body_section._sectPr.find(qn("w:pgNumType"))
    assert page_type.get(qn("w:start")) == "1"
    footer = body_section.footer.paragraphs[0]
    assert footer.alignment == WD_ALIGN_PARAGRAPH.CENTER
    field = footer._p.find(qn("w:fldSimple"))
    assert field is not None and field.get(qn("w:instr")).strip() == "PAGE"
    assert next(field.iter(qn("w:t"))).text == "1"


@pytest.mark.parametrize("preset_id", ["zh_course", "zh_thesis"])
def test_selected_paragraphs_and_body_objects_survive_chinese_layout(tmp_path, preset_id):
    source = Document()
    for text, font, size, align in (
            ("保留第一段 Original 123", "Courier New", 17, WD_ALIGN_PARAGRAPH.CENTER),
            ("保留第二段 Original 456", "Calibri", 10, WD_ALIGN_PARAGRAPH.RIGHT)):
        paragraph = source.add_paragraph()
        paragraph.alignment = align
        paragraph.paragraph_format.line_spacing = 1.15
        paragraph.paragraph_format.space_before = Pt(13)
        run = paragraph.add_run(text)
        run.font.name, run.font.size = font, Pt(size)
        run.font.color.rgb = RGBColor.from_string("006699")
        run.bold, run.italic = True, True
    protected = {paragraph.text: paragraph._p.xml for paragraph in source.paragraphs}
    ordinary = source.add_paragraph(BODY)
    indent = ordinary._p.get_or_add_pPr().get_or_add_ind()
    for key in ("startChars", "endChars", "leftChars", "rightChars", "hangingChars"):
        indent.set(qn("w:" + key), "300")
    source.add_picture(str(Path(__file__).parent / "static" / "capybara-64.png"), width=Inches(0.4))
    table = source.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "保留表格 Table"
    table_xml = table._tbl.xml
    linked = source.add_paragraph("超链接：")
    link = OxmlElement("w:hyperlink")
    relationship = source.part.relate_to("https://example.com/chinese", RT.HYPERLINK, is_external=True)
    link.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    italic = OxmlElement("w:i")
    rpr.append(italic)
    run.append(rpr)
    text = OxmlElement("w:t")
    text.text = "示例链接 Example"
    run.append(text)
    link.append(run)
    linked._p.append(link)
    source.add_paragraph("参考文献")
    for reference in REFS:
        source.add_paragraph(reference)
    path = tmp_path / "retained.docx"
    source.save(path)
    original_bytes = path.read_bytes()
    fingerprint = inspect_draft_paragraphs(path)["source_fingerprint"]
    document, stats, output = _convert(path, preset_id, preserve_paragraphs=[1, 2],
                                      source_fingerprint=fingerprint)
    assert path.read_bytes() == original_bytes
    for text, xml in protected.items():
        assert _find(document, text)._p.xml == xml
    _assert_body_layout(_find(document, BODY), preset_id)
    assert document.tables[0]._tbl.xml == table_xml
    retained = next(document.element.body.iter(qn("w:hyperlink")))
    assert document.part.rels[retained.get(qn("r:id"))].target_ref == "https://example.com/chinese"
    assert next(retained.iter(qn("w:i"))).get(qn("w:val"), "1") == "1"
    with zipfile.ZipFile(path) as before, zipfile.ZipFile(io.BytesIO(output)) as after:
        media = [name for name in before.namelist() if name.startswith("word/media/")]
        assert media
        assert all(before.read(name) == after.read(name) for name in media)
    assert stats["preserved_paragraphs"] == 2
    assert stats["retained_images"] == stats["retained_tables"] == 1
    assert _find(document, "参考文献").paragraph_format.page_break_before is True
    texts = [paragraph.text for paragraph in document.paragraphs]
    assert texts.index(REFS[0]) < texts.index(REFS[1])


@pytest.mark.parametrize("source_kind,target_kind,source_date,target_date", [
    ("western", "zh_course", "21 October 2026", "2026年10月21日"),
    ("zh_course", "western", "2026年10月21日", "October 21, 2026"),
    ("zh_thesis", "zh_course", "2026年10月21日", "21 October 2026"),
    ("zh_course", "zh_thesis", "21 October 2026", "2026年10月21日"),
    ("apa", "zh_course", "October 21, 2026", "2026年10月21日"),
    ("zh_course", "apa", "2026年10月21日", "October 21, 2026"),
])
def test_front_matter_deduplicates_across_layouts_and_equivalent_dates(
        tmp_path, source_kind, target_kind, source_date, target_date):
    draft = tmp_path / "body.txt"
    draft.write_text(BODY, encoding="utf-8")

    def settings(kind, date):
        return dict(META, date_str=date, format_id="apa" if kind == "apa" else "custom",
                    custom_profile=None if kind == "apa" else
                    dict(DEFAULT_CUSTOM_PROFILE) if kind == "western" else _preset(kind))

    intermediate = tmp_path / "previous.docx"
    convert_draft_to_mla(draft, intermediate, **settings(source_kind, source_date))
    output = io.BytesIO()
    stats = convert_draft_to_mla(intermediate, output, **settings(target_kind, target_date))
    result = Document(io.BytesIO(output.getvalue()))
    texts = [paragraph.text for paragraph in result.paragraphs]
    assert texts.count(BODY) == 1
    for key in ("name", "instructor", "course", "affiliation"):
        assert texts.count(META[key]) == 1
    expected_titles = 2 if target_kind in ("apa", "zh_thesis") else 1
    assert texts.count(META["title"]) == expected_titles
    assert texts.count(target_date) == 1
    assert stats["removed_heading_lines"] >= 6


def test_title_without_matching_metadata_is_not_removed(tmp_path):
    path = tmp_path / "title-body.txt"
    path.write_text(META["title"] + "\n\n" + BODY, encoding="utf-8")
    document, stats, _ = _convert(path)
    assert stats["removed_heading_lines"] == 0
    assert [paragraph.text for paragraph in document.paragraphs].count(META["title"]) == 2
