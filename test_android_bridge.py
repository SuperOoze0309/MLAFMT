"""Exercise Android's offline Python adapter on desktop as well as device."""
import importlib.util
import json
from pathlib import Path
import zipfile

from docx import Document
import pytest

spec = importlib.util.spec_from_file_location(
    "android_bridge", Path(__file__).parent / "platforms/android/app/src/main/python/android_bridge.py")
bridge = importlib.util.module_from_spec(spec)
spec.loader.exec_module(bridge)


def request(**updates):
    return json.dumps(dict(kind="convert", name="Jane Doe", title="My Essay",
                           instructor="Professor Smith", course="ENG 101",
                           date="13 October 2026", **updates))


def test_android_adapter_generates_sorted_works_cited(tmp_path):
    draft = tmp_path / "draft.md"
    draft.write_text("Body *italics*.\n\nWorks Cited\n\nZulu, Jane. Book.\n\nAlpha, Jane. Book.", encoding="utf-8")
    result = json.loads(bridge.generate(request(), str(draft), str(tmp_path / "result")))
    assert result["name"] == "Doe_MLA_Essay.docx"
    assert result["stats"]["works_cited"] == 2
    document = Document(result["path"])
    cited = next(p for p in document.paragraphs if p.text == "Works Cited")
    assert cited.paragraph_format.page_break_before is True
    assert document.paragraphs[-2].text.startswith("Alpha")
    assert document.paragraphs[-1].paragraph_format.first_line_indent.inches == -0.5


@pytest.mark.parametrize("updates,error", [
    ({"name": []}, "invalid_fields"),
    ({"name": ""}, "missing_fields"),
    ({"txt_mode": "wrong"}, "invalid_options"),
    ({"markdown": "true"}, "invalid_options"),
    ({"kind": "wrong"}, "invalid_options"),
])
def test_android_adapter_rejects_bad_requests(tmp_path, updates, error):
    data = json.loads(request())
    data.update(updates)
    assert json.loads(bridge.generate(json.dumps(data), "", str(tmp_path)))["error"] == error


def test_android_adapter_bounds_zip_expansion(tmp_path, monkeypatch):
    draft = tmp_path / "bomb.docx"
    with zipfile.ZipFile(draft, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "x" * 1000)
    monkeypatch.setattr(bridge, "MAX_DOCX_EXPANDED", 100)
    assert json.loads(bridge.generate(request(), str(draft), str(tmp_path)))["error"] == "bad_doc"


def test_android_adapter_checks_empty_and_large_files(tmp_path, monkeypatch):
    draft = tmp_path / "draft.txt"
    draft.write_bytes(b"")
    assert json.loads(bridge.generate(request(), str(draft), str(tmp_path)))["error"] == "empty_file"
    draft.write_bytes(b"12345")
    monkeypatch.setattr(bridge, "MAX_FILE_SIZE", 4)
    assert json.loads(bridge.generate(request(), str(draft), str(tmp_path)))["error"] == "too_large"


def test_android_blank_template_supports_unicode_filename(tmp_path):
    data = json.loads(request())
    data.update(kind="blank", last_name_override="张" * 200)
    result = json.loads(bridge.generate(json.dumps(data), "", str(tmp_path)))
    assert len(result["name"].encode("utf-8")) < 200
    assert Document(result["path"]).paragraphs[4].text == "My Essay"
