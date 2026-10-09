"""HTTP and generated-document regression tests; run with python -m unittest."""

import io
import os
import unittest
from unittest.mock import patch
from urllib.parse import unquote
from zipfile import ZipFile

from docx import Document
from docx.shared import Inches, Pt

import mla_web


class WebTests(unittest.TestCase):
    def setUp(self):
        self.client = mla_web.app.test_client()
        self.details = {"name": "Jane Doe", "title": "An Essay", "instructor": "Professor Smith",
                        "course": "ENG 101", "date": "8 October 2026"}
        self.responses = []

    def tearDown(self):
        for response in self.responses:
            response.close()

    def post(self, path, **kwargs):
        response = self.client.post(path, **kwargs)
        self.responses.append(response)
        return response

    def upload(self, content=b"First paragraph.\n\nSecond paragraph.", filename="draft.txt", **options):
        return self.post("/api/convert", data={**self.details, **options,
                         "file": (io.BytesIO(content), filename)}, content_type="multipart/form-data")

    def assert_error(self, response, code, status=400):
        self.assertEqual(response.status_code, status)
        self.assertEqual(response.get_json()["error"], code)

    def test_home_renders(self):
        response = self.client.get("/")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b'id="lang-zh"', response.data)
        self.assertIn(b'id="lang-en"', response.data)
        self.assertNotIn(b"{{", response.data)

    def test_json_endpoints_reject_non_objects_and_malformed_json(self):
        for path in ("/api/blank", "/api/summary"):
            for data in ("[]", "null", '"text"', "42", "true", "{"):
                with self.subTest(path=path, data=data):
                    self.assert_error(self.post(path, data=data, content_type="application/json"), "invalid_json")

    def test_json_endpoints_reject_incorrect_field_types(self):
        for path in ("/api/blank", "/api/summary"):
            for key in mla_web.TEXT_FIELDS:
                for value in (4, True, [], {}):
                    with self.subTest(path=path, key=key, value=value):
                        self.assert_error(self.post(path, json={**self.details, key: value}), "invalid_fields")

    def test_summary_keeps_existing_list_contract_and_suffix_handling(self):
        response = self.post("/api/summary", json={"name": "Jane Smith Jr."})
        self.assertEqual(response.status_code, 200)
        self.assertIsInstance(response.get_json(), list)
        self.assertIn("Header: Smith 1", response.get_json()[0])
        override = self.post("/api/summary", json={"name": "Jane Doe", "last_name_override": "Smith-Jones"})
        self.assertIn("Header: Smith-Jones 1", override.get_json()[0])

    def test_blank_requires_name_and_title(self):
        for data in ({}, {"name": "Jane"}, {"name": None, "title": "Essay"}, {"name": "  ", "title": "Essay"}):
            with self.subTest(data=data):
                self.assert_error(self.post("/api/blank", json=data), "missing_fields")

    def test_blank_generates_real_mla_docx_with_unicode_filename(self):
        response = self.post("/api/blank", json={**self.details, "name": "Jane 王", "last_name_override": "王"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.mimetype, mla_web.DOCX_MIME)
        self.assertEqual(unquote(response.headers["X-Filename"]), "王_MLA_Essay.docx")
        self.assertIn("filename*=UTF-8''", response.headers["Content-Disposition"])
        doc = Document(io.BytesIO(response.data))
        self.assertEqual(doc.sections[0].left_margin, Inches(1))
        self.assertEqual(doc.styles["Normal"].font.size, Pt(12))
        self.assertEqual(doc.paragraphs[0].text, "Jane 王")
        self.assertIn("An Essay", [p.text for p in doc.paragraphs])
        self.assertIn("王", doc.sections[0].header.paragraphs[0].text)

    def test_missing_upload_and_unsupported_extension(self):
        self.assert_error(self.post("/api/convert", data=self.details), "no_file")
        self.assert_error(self.upload(filename="draft.pdf"), "unsupported")
        self.assert_error(self.upload(filename=""), "no_file")

    def test_empty_upload(self):
        self.assert_error(self.upload(content=b""), "empty_file")

    def test_upload_file_limit_and_request_limit(self):
        with patch.object(mla_web, "MAX_FILE_SIZE", 8):
            self.assert_error(self.upload(content=b"x" * 9), "too_large", 413)
        with patch.dict(mla_web.app.config, {"MAX_CONTENT_LENGTH": 16}):
            self.assert_error(self.post("/api/convert", data=b"x" * 17,
                                       content_type="multipart/form-data; boundary=test"), "too_large", 413)

    def test_conversion_missing_fields_and_invalid_options(self):
        self.assert_error(self.upload(name="  "), "missing_fields")
        self.assert_error(self.upload(txt_mode="surprise"), "invalid_options")

    def test_corrupt_docx_is_a_client_error(self):
        self.assert_error(self.upload(content=b"not a Word file", filename="draft.docx"), "bad_doc")
        archive = io.BytesIO()
        with ZipFile(archive, "w") as zip_file:
            zip_file.writestr("notes.txt", "A ZIP is not necessarily a Word document.")
        self.assert_error(self.upload(content=archive.getvalue(), filename="draft.docx"), "bad_doc")

    def test_text_conversion_preserves_options_stats_and_citations(self):
        draft = b"First *important* paragraph.\n\nSecond paragraph.\n\nWorks Cited\n\nZebra, A. A Book.\n\nAlpha, B. Another Book."
        response = self.upload(content=draft, filename="draft.MD", markdown="true", last_name_override="Smith-Jones")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["X-Stat-body"], "2")
        self.assertEqual(response.headers["X-Stat-works-cited"], "2")
        self.assertEqual(response.headers["X-Filename"], "Smith-Jones_MLA_Essay.docx")
        doc = Document(io.BytesIO(response.data))
        texts = [paragraph.text for paragraph in doc.paragraphs]
        self.assertLess(texts.index("Alpha, B. Another Book."), texts.index("Zebra, A. A Book."))
        body = next(p for p in doc.paragraphs if p.text.startswith("First"))
        self.assertEqual(body.paragraph_format.first_line_indent, Inches(0.5))
        self.assertTrue(any(run.italic and run.text == "important" for run in body.runs))

    def test_docx_input_and_line_by_line_mode(self):
        original = Document()
        original.add_paragraph("A draft with bold text.").runs[0].bold = True
        buffer = io.BytesIO()
        original.save(buffer)
        response = self.upload(content=buffer.getvalue(), filename="draft.docx")
        self.assertEqual(response.status_code, 200)
        converted = Document(io.BytesIO(response.data))
        paragraph = next(p for p in converted.paragraphs if p.text == "A draft with bold text.")
        self.assertTrue(paragraph.runs[0].bold)
        lines = self.upload(content=b"First line.\nSecond line.", txt_mode="line_by_line")
        self.assertEqual(lines.headers["X-Stat-body"], "2")

    def test_temporary_upload_is_removed_after_failure(self):
        captured = []

        def fail(path, *_args, **_kwargs):
            captured.append(path)
            raise RuntimeError("test conversion failure")

        with patch.object(mla_web, "convert_draft_to_mla", side_effect=fail):
            self.assert_error(self.upload(), "failed", 500)
        self.assertEqual(len(captured), 1)
        self.assertFalse(os.path.exists(captured[0]))


if __name__ == "__main__":
    unittest.main()
