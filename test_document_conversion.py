"""Conversion tests use real generated PDF/DOCX bytes, never disguised DOC files."""

from io import BytesIO
import os
from pathlib import Path
import subprocess
import threading
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
import pytest

import document_conversion as conversion


def docx_bytes(text="Fictional conversion sample"):
    document = Document()
    document.add_paragraph(text)
    document.add_table(rows=1, cols=2).rows[0].cells[0].text = "Table cell"
    result = BytesIO()
    document.save(result)
    return result.getvalue()


def pdf_bytes(pages=("First page sample",), *, encrypted=False, image=False):
    pypdf = pytest.importorskip("pypdf")
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
    writer = pypdf.PdfWriter()
    for text in pages:
        page = writer.add_blank_page(width=612, height=792)
        resources = DictionaryObject()
        resources[NameObject("/Font")] = DictionaryObject({
            NameObject("/F1"): DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica")})})
        if image:
            resources[NameObject("/XObject")] = DictionaryObject({
                NameObject("/UnusedForm"): DictionaryObject({NameObject("/Subtype"): NameObject("/Form")})})
        page[NameObject("/Resources")] = resources
        if text:
            escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            content = DecodedStreamObject()
            content.set_data(f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("latin-1"))
            page[NameObject("/Contents")] = writer._add_object(content)
    if encrypted:
        writer.encrypt("fixture-password")
    result = BytesIO()
    writer.write(result)
    return result.getvalue()


def assert_error(code, function, *args, **kwargs):
    with pytest.raises(conversion.ConversionError) as caught:
        function(*args, **kwargs)
    assert caught.value.code == code


def test_pdf_to_editable_docx_retains_page_text_order_and_input_bytes():
    original = pdf_bytes(("First page (test) and punctuation.", "Second page sample"))
    before = bytes(original)
    result = conversion.convert_document(original, "C:/private/fictional.PDF", "docx")
    assert original == before
    assert result.filename == "fictional.docx"
    assert result.mime_type == conversion.MIME_TYPES["docx"]
    assert result.engine == "pypdf"
    assert result.source_format == "pdf" and result.target_format == "docx"
    assert result.warnings == ("pdf_text_reconstruction",)
    document = Document(BytesIO(result.data))
    text = "\n".join(p.text for p in document.paragraphs)
    assert text.index("First page (test)") < text.index("Second page")
    assert document.sections[0].page_width.inches == 8.5
    with ZipFile(BytesIO(result.data)) as archive:
        assert archive.testzip() is None
        assert b'w:type="page"' in archive.read("word/document.xml")


def test_pdf_blank_page_is_preserved_and_disclosed():
    result, warnings = conversion.pdf_to_docx(pdf_bytes(("Page one", "", "Page three")))
    assert "pdf_blank_pages" in warnings
    with ZipFile(BytesIO(result)) as archive:
        assert archive.read("word/document.xml").count(b'w:type="page"') == 2


def test_scanned_or_empty_pdf_does_not_fabricate_editable_content():
    assert_error("pdf_needs_ocr", conversion.pdf_to_docx, pdf_bytes(("",)))


def test_pdf_encryption_is_reported_without_attempting_password_bypass():
    assert_error("encrypted_document", conversion.pdf_to_docx, pdf_bytes(encrypted=True))


def test_pdf_page_limit_precedes_text_extraction():
    with patch.object(conversion, "MAX_PDF_PAGES", 1):
        assert_error("pdf_page_limit", conversion.pdf_to_docx, pdf_bytes(("A", "B")))


def test_pdf_text_and_stream_limits_are_enforced():
    data = pdf_bytes(("This is a text content fixture",))
    with patch.object(conversion, "MAX_PDF_TEXT", 4):
        assert_error("pdf_content_limit", conversion.pdf_to_docx, data)
    with patch.object(conversion, "MAX_PDF_PAGE_CONTENT", 4):
        assert_error("pdf_content_limit", conversion.pdf_to_docx, data)


def test_pdf_parser_errors_are_stable_and_do_not_include_draft_content():
    assert_error("bad_pdf", conversion.pdf_to_docx, b"%PDF-1.7\nPRIVATE broken fixture\n%%EOF")


@pytest.mark.parametrize("declared,data,code", [
    ("pdf", b"not a pdf", "bad_pdf"),
    ("docx", b"<html>Fake Word</html>", "bad_doc"),
    ("doc", b"{\\rtf1 A renamed RTF is not a binary DOC}", "bad_doc"),
    ("doc", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 504, "bad_doc"),
])
def test_renamed_or_truncated_files_are_rejected(declared, data, code):
    assert_error(code, conversion.validate_document, data, declared)


def rewrite_docx(changes):
    out = BytesIO()
    with ZipFile(BytesIO(docx_bytes())) as source, ZipFile(out, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            target.writestr(item.filename, source.read(item.filename))
        for name, contents in changes.items():
            target.writestr(name, contents)
    return out.getvalue()


@pytest.mark.parametrize("name,content,code", [
    ("../outside.txt", "blocked", "bad_doc"),
    ("word/vbaProject.bin", "blocked", "unsupported_document"),
    ("word/_rels/foreign.xml.rels", '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="x" Type="https://example.org/image" Target="https://example.org/private.png" TargetMode="External"/></Relationships>', "external_content"),
])
def test_docx_package_resource_boundaries(name, content, code):
    assert_error(code, conversion.validate_document, rewrite_docx({name: content}), "docx")


def test_docx_zip_expansion_limit_prevents_processing_compression_bombs():
    data = rewrite_docx({"word/large.txt": "0" * 50000})
    with patch.object(conversion, "MAX_OUTPUT_SIZE", 40000):
        assert_error("too_large", conversion.validate_document, data, "docx")


def test_docx_hyperlinks_do_not_require_network_or_get_rejected():
    extra = '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="x" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" Target="https://example.org/" TargetMode="External"/></Relationships>'
    data = rewrite_docx({"word/_rels/link.xml.rels": extra})
    assert conversion.validate_document(data, "docx") == data


def test_cancellation_before_pdf_conversion_does_not_process_input():
    cancelled = threading.Event()
    cancelled.set()
    assert_error("cancelled", conversion.pdf_to_docx, b"not even parsed", cancel_event=cancelled)


@pytest.mark.parametrize("timeout", [True, 0, -1, float("inf"), float("nan"), 601, "120"])
def test_invalid_timeout_options_are_rejected(timeout):
    assert_error("invalid_options", conversion.convert_document, b"x", "a.pdf", "docx", timeout=timeout)


def test_same_format_and_unknown_format_are_rejected():
    assert_error("same_format", conversion.convert_document, docx_bytes(), "a.docx", "docx")
    assert_error("unsupported", conversion.convert_document, b"text", "a.txt", "pdf")
    assert_error("invalid_options", conversion.convert_document, b"x", "a.pdf", "rtf")


def test_file_size_type_and_empty_boundaries():
    assert_error("empty_file", conversion.validate_document, b"", "pdf")
    assert_error("invalid_file", conversion.validate_document, "not bytes", "pdf")
    with patch.object(conversion, "MAX_FILE_SIZE", 4):
        assert_error("too_large", conversion.validate_document, b"12345", "pdf")


def test_missing_office_dependency_is_not_a_fake_success():
    with patch.object(conversion, "find_soffice", return_value=None):
        assert_error("office_engine_missing", conversion.convert_document,
                     docx_bytes(), "a.docx", "doc")


def test_capability_matrix_has_exactly_six_directional_routes():
    with patch.object(conversion, "find_soffice", return_value=None), patch.object(conversion, "_pypdf_available", return_value=True):
        state = conversion.conversion_capabilities()
    assert state["local"] is True
    assert len(state["routes"]) == 6
    assert {(r["source"], r["target"]) for r in state["routes"] if r["available"]} == {("pdf", "docx")}
    with patch.object(conversion, "find_soffice", return_value="soffice"), patch.object(conversion, "_pypdf_available", return_value=True):
        assert all(r["available"] for r in conversion.conversion_capabilities()["routes"])


def test_explicit_office_setting_is_one_executable_not_a_shell_command(tmp_path):
    assert conversion.find_soffice(str(tmp_path / "missing.exe") + " --headless") is None
    with patch.dict(os.environ, {"FORMCAPY_SOFFICE": str(tmp_path / "missing.exe")}):
        assert conversion.find_soffice() is None


class HangingProcess:
    pid = 1234567
    returncode = None

    def poll(self):
        return self.returncode

    def wait(self, timeout):
        raise subprocess.TimeoutExpired("fixture", timeout)


def test_office_timeout_stops_only_the_job_and_removes_its_temporary_files(tmp_path):
    process = HangingProcess()
    stopped = []

    def stop(value):
        stopped.append(value)
        value.returncode = -9

    with patch.object(conversion, "find_soffice", return_value="fixture-soffice"), patch("subprocess.Popen", return_value=process), patch.object(conversion, "_stop_process", side_effect=stop):
        assert_error("conversion_timeout", conversion.convert_document,
                     docx_bytes(), "a.docx", "pdf", timeout=0.01, temp_root=tmp_path)
    assert stopped == [process]
    assert not list(tmp_path.iterdir())


def test_office_command_uses_isolated_profile_real_doc_filter_and_no_shell(tmp_path):
    captured = {}

    def fail(command, **options):
        captured.update(command=command, options=options)
        raise OSError("test: executable unavailable")

    with patch.object(conversion, "find_soffice", return_value="fixture-soffice"), patch("subprocess.Popen", side_effect=fail):
        assert_error("conversion_failed", conversion.convert_document,
                     docx_bytes(), "../../private.docx", "doc", temp_root=tmp_path)
    command = captured["command"]
    assert "doc:MS Word 97" in command
    assert any(item.startswith("-env:UserInstallation=file:") for item in command)
    assert command[-1].endswith("document.docx")
    assert "shell" not in captured["options"]
    assert not list(tmp_path.iterdir())


@pytest.mark.skipif(not conversion.find_soffice(), reason="Native LibreOffice is not installed")
def test_real_libreoffice_six_routes_and_real_binary_doc(tmp_path):
    """When installed, inspect actual PDF text and DOC round trips independently."""
    pypdf = pytest.importorskip("pypdf")
    original = docx_bytes("Fictional six-route sample 42")
    native_doc = conversion.convert_document(original, "sample.docx", "doc", temp_root=tmp_path)
    assert native_doc.data.startswith(conversion.OLE_SIGNATURE)
    native_pdf = conversion.convert_document(original, "sample.docx", "pdf", temp_root=tmp_path)
    assert "six-route sample 42" in pypdf.PdfReader(BytesIO(native_pdf.data)).pages[0].extract_text()
    returned_docx = conversion.convert_document(native_doc.data, "sample.doc", "docx", temp_root=tmp_path)
    assert "six-route sample 42" in "\n".join(p.text for p in Document(BytesIO(returned_docx.data)).paragraphs)
    assert conversion.convert_document(native_doc.data, "sample.doc", "pdf", temp_root=tmp_path).data.startswith(b"%PDF-")
    pdf_docx = conversion.convert_document(native_pdf.data, "sample.pdf", "docx", temp_root=tmp_path)
    assert "six-route sample 42" in "\n".join(p.text for p in Document(BytesIO(pdf_docx.data)).paragraphs)
    pdf_doc = conversion.convert_document(native_pdf.data, "sample.pdf", "doc", temp_root=tmp_path)
    assert pdf_doc.data.startswith(conversion.OLE_SIGNATURE)
    assert not list(tmp_path.iterdir())
