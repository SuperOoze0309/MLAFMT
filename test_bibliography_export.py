"""Fictional CSL list fixtures: exporter and browser bridge boundaries."""

from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from lxml import etree
import pytest

import browser_bridge
from bibliography_export import BibliographyExportError, render_bibliography_docx


NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}


def bibliography():
    return {
        "style": "mla", "heading": "Works Cited",
        "html": "", "text": "",
        "entries": [
            {"id": "z", "html": '<div class="csl-entry">Zed, Jo. <i>River &amp; Stone</i>. Test Press, 2024.</div>\n',
             "text": "Zed, Jo. River & Stone. Test Press, 2024."},
            {"id": "a", "html": '<div class="csl-entry">Alfa Research. <a href="https://example.org/a?b=1&amp;c=2"><b>Signals</b></a>. 2023.</div>\n',
             "text": "Alfa Research. Signals. 2023."},
        ],
        "layout": {"hangingindent": True, "linespacing": 2, "entryspacing": 0,
                   "second-field-align": False, "maxoffset": 0},
    }


def package(data):
    result = render_bibliography_docx(data)
    archive = ZipFile(BytesIO(result))
    assert archive.testzip() is None
    return result, archive, etree.fromstring(archive.read("word/document.xml"))


def test_export_order_formatting_links_and_page_geometry():
    result, archive, xml = package(bibliography())
    assert len(result) > 1000
    texts = xml.xpath("//w:p", namespaces=NS)
    assert ["".join(p.xpath(".//w:t/text()", namespaces=NS)) for p in texts] == [
        "Works Cited", "Zed, Jo. River & Stone. Test Press, 2024.", "Alfa Research. Signals. 2023."]
    # Input order is deliberately reverse alphabetical; exporter never sorts.
    assert xml.xpath("//w:r[w:t='River & Stone']/w:rPr/w:i", namespaces=NS)
    assert xml.xpath("//w:hyperlink/w:r[w:t='Signals']/w:rPr/w:b", namespaces=NS)
    assert b"https://example.org/a?b=1&amp;c=2" in archive.read("word/_rels/document.xml.rels")
    assert not any(name.startswith("word/header") for name in archive.namelist())
    assert not any(name.startswith("word/media") for name in archive.namelist())
    assert xml.xpath("//w:sectPr/w:pgSz/@w:w", namespaces=NS) == ["12240"]
    assert xml.xpath("//w:sectPr/w:pgSz/@w:h", namespaces=NS) == ["15840"]
    for edge in ("top", "bottom", "left", "right"):
        assert xml.xpath(f"//w:sectPr/w:pgMar/@w:{edge}", namespaces=NS) == ["1440"]
    assert xml.xpath("//w:ind/@w:hanging", namespaces=NS) == ["720", "720"]


def test_apa_heading_bold_and_csl_layout_spacing():
    data = bibliography()
    data.update(style="apa", heading="References")
    data["layout"].update(linespacing=1.5, entryspacing=1, hangingindent=False)
    _, _, xml = package(data)
    assert xml.xpath("//w:p[1]/w:r/w:rPr/w:b", namespaces=NS)
    assert xml.xpath("//w:p[2]/w:pPr/w:spacing/@w:line", namespaces=NS) == ["360"]
    assert xml.xpath("//w:p[2]/w:pPr/w:spacing/@w:after", namespaces=NS) == ["360"]
    assert not xml.xpath("//w:p[2]/w:pPr/w:ind", namespaces=NS)


def test_nested_font_semantics_line_breaks_and_unicode():
    data = bibliography()
    data["entries"] = [{"id": "unicode", "text": "中文 🦫 Café Test x2y\nnext",
        "html": '<div class="csl-entry"><span style="font-variant:small-caps;">中文 🦫 Café</span> <i>Test <span style="font-style:normal;font-weight:700;text-decoration:underline;">x</span></i><sup>2</sup><sub>y</sub><br/>next</div>'}]
    _, _, xml = package(data)
    assert xml.xpath("//w:smallCaps", namespaces=NS)
    assert xml.xpath("//w:r[w:t='x']/w:rPr/w:i/@w:val", namespaces=NS) == ["0"]
    assert xml.xpath("//w:r[w:t='x']/w:rPr/w:b", namespaces=NS)
    assert xml.xpath("//w:r[w:t='x']/w:rPr/w:u/@w:val", namespaces=NS) == ["single"]
    assert xml.xpath("//w:vertAlign/@w:val", namespaces=NS) == ["superscript", "subscript"]
    assert xml.xpath("//w:br", namespaces=NS)
    assert "中文 🦫 Café" in "".join(xml.xpath("//w:t/text()", namespaces=NS))


@pytest.mark.parametrize("html,text", [
    ('<div class="csl-entry">Shortened citation.</div>', "Full original citation."),
    ('<div class="csl-entry">A&#0;B</div>', "A\x00B"),
    ('<div class="csl-entry">Missing <sup>2</sup></div>', "Missing"),
    ('<div class="csl-entry">Different &amp; title</div>', "Different title"),
])
def test_reject_html_that_does_not_match_the_reviewed_plain_text(html, text):
    data = bibliography()
    data["entries"] = [{"id": "mismatch", "html": html, "text": text}]
    with pytest.raises(BibliographyExportError):
        render_bibliography_docx(data)


def test_matching_nbsp_line_breaks_and_literal_entities_preserve_all_text():
    data = bibliography()
    data["entries"] = [{"id": "representations",
        "html": '<div class="csl-entry">A&nbsp;  title &amp;lt;XML&amp;gt;<br/>Next 🦫</div>',
        "text": "A   title &lt;XML&gt;\r\nNext 🦫"}]
    content = render_bibliography_docx(data)
    text = Document(BytesIO(content)).paragraphs[1].text
    assert text == "A\u00a0  title &lt;XML&gt;\nNext 🦫"


@pytest.mark.parametrize("html", [
    '<script>alert(1)</script>', '<img src="https://example.org/x">',
    '<div onclick="evil()">Text</div>', '<span style="color:red">Text</span>',
    '<span style="font-size:20pt">Text</span>', '<small>Text</small>',
    '<i>Unclosed', '<i><b>Mismatch</i></b>', '<br></br>',
    '<div class="csl-entry"><div>Nested block</div></div>',
    '<div class="csl-left-margin">1</div>', '<!-- hidden -->Text', '<!DOCTYPE html>Text',
    '<a href="javascript:alert(1)">Text</a>', '<a href="data:text/html,x">Text</a>',
    '<a href="file:///c:/x">Text</a>', '<a href="//example.org/x">Text</a>',
    '<a href="https://u:p@example.org/">Text</a>', '<a href="https://example.org:bad/">Text</a>',
    '<a href="https://example.org/ x">Text</a>', '<a href="https://example.org/\\x">Text</a>',
    '<a href="https://example.org/x" href="https://example.org/y">Text</a>',
    '<span style="font-weight:bold;font-weight:normal">Text</span>',
    '<div class="csl-entry"></div>', '', '   ', '<div>One</div><div>Two</div>',
])
def test_reject_dangerous_or_unrepresentable_html(html):
    data = bibliography()
    data["entries"][0]["html"] = html
    with pytest.raises(BibliographyExportError, match="invalid_options"):
        render_bibliography_docx(data)


@pytest.mark.parametrize("key,value", [
    ("linespacing", True), ("linespacing", float("inf")), ("linespacing", float("nan")),
    ("linespacing", -1), ("entryspacing", 21), ("entryspacing", "1"),
    ("hangingindent", 1), ("second-field-align", "flush"), ("second-field-align", None),
    ("maxoffset", 1001),
])
def test_reject_invalid_layout(key, value):
    data = bibliography()
    data["layout"][key] = value
    with pytest.raises(BibliographyExportError):
        render_bibliography_docx(data)


@pytest.mark.parametrize("mutation", [
    lambda data: data.update(style="custom"),
    lambda data: data.update(heading="Wrong heading"),
    lambda data: data.update(entries=[]),
    lambda data: data.update(entries=data["entries"] * 101),
    lambda data: data["entries"][0].update(id="../unsafe"),
    lambda data: data["entries"][0].update(id=data["entries"][1]["id"]),
    lambda data: data["entries"][0].update(html="a" * (32 * 1024 + 1)),
    lambda data: data["entries"][0].update(html="A\x00B"),
    lambda data: data["entries"][0].update(text="A\ud800B"),
    lambda data: data["entries"][0].update(extra="unknown"),
    lambda data: data["layout"].update(extra="unknown"),
    lambda data: data.update(draft="must not be in bibliography payload"),
])
def test_reject_unknown_fields_duplicates_and_limits(mutation):
    data = bibliography()
    mutation(data)
    with pytest.raises(BibliographyExportError):
        render_bibliography_docx(data)


def test_default_line_spacing_and_metadata_do_not_contain_paper_identity():
    data = bibliography()
    data["layout"]["linespacing"] = 0
    result = render_bibliography_docx(data)
    document = Document(BytesIO(result))
    assert document.paragraphs[1].paragraph_format.line_spacing == 2
    assert document.core_properties.author == ""
    assert document.core_properties.last_modified_by == ""


def test_full_entry_limit_is_supported_without_sorting():
    data = bibliography()
    data["entries"] = [{"id": f"item-{200-index}", "html": f'<div class="csl-entry">Entry {200-index}.</div>',
                        "text": f"Entry {200-index}."} for index in range(200)]
    result = render_bibliography_docx(data)
    document = Document(BytesIO(result))
    assert len(document.paragraphs) == 201
    assert document.paragraphs[1].text == "Entry 200."
    assert document.paragraphs[-1].text == "Entry 1."


def test_total_html_limit_is_enforced_across_separate_entries():
    data = bibliography()
    data["entries"] = [{"id": f"item-{index}", "html": "x" * (32 * 1024), "text": "x" * (32 * 1024)}
                       for index in range(65)]
    with pytest.raises(BibliographyExportError):
        render_bibliography_docx(data)


def test_bridge_recovers_after_bad_rich_text_and_cleans_every_request(tmp_path):
    with patch.object(browser_bridge.tempfile, "tempdir", str(tmp_path)):
        for index in range(4):
            bad = bibliography()
            bad["entries"][0]["html"] = '<div class="csl-entry">Silently shortened.</div>'
            for payload, expected_ok in ((bad, False), (bibliography(), True)):
                started = json.loads(browser_bridge.begin())
                metadata, content = browser_bridge.finish(json.dumps({
                    "kind": "bibliography", "bibliography": payload,
                }), started["workdir"])
                meta = json.loads(metadata)
                assert meta["ok"] is expected_ok
                if expected_ok:
                    assert content.startswith(b"PK")
                    assert meta["stats"]["works_cited"] == 2
                else:
                    assert meta["error"] == "invalid_options"
                    assert content == b""
                assert not Path(started["workdir"]).exists()
        assert not list(tmp_path.glob("mlafmt-*"))


def test_bridge_independent_request_and_cleanup_on_success(tmp_path):
    # A bridge-created temp location is bounded by its existing cleanup guard.
    with patch.object(browser_bridge.tempfile, "tempdir", str(tmp_path)):
        started = json.loads(browser_bridge.begin())
        assert Path(started["workdir"]).is_dir()
        request = {"kind": "bibliography", "bibliography": bibliography()}
        metadata, content = browser_bridge.finish(json.dumps(request), started["workdir"])
        meta = json.loads(metadata)
        assert meta["ok"] is True
        assert meta["name"] == "FormCapy-MLA-References.docx"
        assert meta["stats"]["works_cited"] == 2
        assert content.startswith(b"PK")
        assert not Path(started["workdir"]).exists()


@pytest.mark.parametrize("payload", [None, {}, {"style": "mla"}, "invalid"])
def test_bridge_invalid_export_cleanup(tmp_path, payload):
    with patch.object(browser_bridge.tempfile, "tempdir", str(tmp_path)):
        started = json.loads(browser_bridge.begin())
        request = {"kind": "bibliography", "bibliography": payload}
        metadata, content = browser_bridge.finish(json.dumps(request), started["workdir"])
        assert json.loads(metadata)["error"] == "invalid_options"
        assert content == b""
        assert not Path(started["workdir"]).exists()


def test_bridge_legacy_blank_still_requires_paper_information(tmp_path):
    with patch.object(browser_bridge.tempfile, "tempdir", str(tmp_path)):
        started = json.loads(browser_bridge.begin())
        metadata, content = browser_bridge.finish('{"kind":"blank"}', started["workdir"])
        assert json.loads(metadata)["error"] == "missing_fields"
        assert content == b""
        assert not Path(started["workdir"]).exists()
